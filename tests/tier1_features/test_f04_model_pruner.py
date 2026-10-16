"""Feature F4: Model Decomposition & Pruning Isolation Tests.

Verifies:
1. Rule 1: Level 6 node with has_children=True is deleted.
2. Rule 2: Matched assembly keeps the unit node itself but prunes internal children.
3. Rule 3: Node with has_children=True but no children expanded (level < 6) is kept.
4. Rule 4: Matched leaf component (has_children=False) is deleted.
5. Rule execution across machine models (Virgo, Libra2, Iris2024, Sirius2, Mebius, Polaris).
"""

import pytest

from src.core.model_pruner import ModelPruner
from src.core.models import BOMNode, BOMTree, ModelRule


class TestF04ModelPruner:
    """Test suite for Feature F4: Model Decomposition & Pruning."""

    def test_f04_rule1_level6_with_children_deleted(self):
        """Test 1: Rule 1 - Prune phantom branch where Level == 6 and has_children=True."""
        rule = ModelRule(item_name="PWB MAIN", match_mode="Part_name")
        pruner = ModelPruner()

        root = BOMNode(level=1, item_id="R1", item_name="ROOT")
        # Build chain down to level 6
        l5 = BOMNode(level=5, item_id="L5", item_name="LV5")
        l6 = BOMNode(level=6, item_id="L6", item_name="PWB MAIN ASSY", has_children=True)
        l6.add_child(BOMNode(level=7, item_id="L7", item_name="PHANTOM"))
        l5.add_child(l6)
        root.add_child(l5)

        tree = BOMTree(roots=[root])
        pruned_tree = pruner.prune_tree(tree, rules=[rule])

        # Level 6 node and its phantom child must be deleted
        assert len(pruned_tree.roots[0].children[0].children) == 0

    def test_f04_rule2_prune_internal_children_keep_unit(self):
        """Test 2: Rule 2 - Matched assembly keeps the assembly row for inventory, prunes sub-components."""
        rule = ModelRule(item_name="BOTTLE WASTE", match_mode="Full_name")
        pruner = ModelPruner()

        root = BOMNode(level=1, item_id="BW01", item_name="BOTTLE WASTE", has_children=True)
        sub1 = BOMNode(level=2, item_id="P1", item_name="INTERNAL BRACKET")
        sub2 = BOMNode(level=2, item_id="P2", item_name="INTERNAL SPONGE")
        root.add_child(sub1)
        root.add_child(sub2)

        tree = BOMTree(roots=[root])
        pruned = pruner.prune_tree(tree, rules=[rule])

        # BOTTLE WASTE remains, but internal children are pruned
        assert pruned.size() == 1
        assert pruned.roots[0].item_name == "BOTTLE WASTE"
        assert len(pruned.roots[0].children) == 0

    def test_f04_rule3_leaf_with_children_flag_kept(self):
        """Test 3: Rule 3 - When matched node has has_children=True but no children in tree, it is kept."""
        rule = ModelRule(item_name="UNEXPANDED ASSY", match_mode="Full_name")
        pruner = ModelPruner()

        # Node level < 6 with has_children=True and len(children) == 0
        node = BOMNode(level=2, item_id="UA01", item_name="UNEXPANDED ASSY", has_children=True)
        root = BOMNode(level=1, item_id="R1", item_name="ROOT")
        root.add_child(node)

        tree = BOMTree(roots=[root])
        pruned = pruner.prune_tree(tree, rules=[rule])

        assert pruned.size() == 2
        assert pruned.roots[0].children[0].item_id == "UA01"

    def test_f04_rule4_matched_leaf_deleted(self):
        """Test 4: Rule 4 - Matched leaf part (has_children=False) is deleted completely."""
        rule = ModelRule(item_name="OPTIONAL SCREW", match_mode="Full_name")
        pruner = ModelPruner()

        root = BOMNode(level=1, item_id="R1", item_name="ROOT")
        leaf1 = BOMNode(level=2, item_id="S01", item_name="OPTIONAL SCREW", has_children=False)
        leaf2 = BOMNode(level=2, item_id="S02", item_name="MANDATORY SCREW", has_children=False)
        root.add_child(leaf1)
        root.add_child(leaf2)

        tree = BOMTree(roots=[root])
        pruned = pruner.prune_tree(tree, rules=[rule])

        assert pruned.size() == 2
        assert len(pruned.roots[0].children) == 1
        assert pruned.roots[0].children[0].item_id == "S02"

    def test_f04_six_models_rule_execution(self):
        """Test 5: Verify model pruner initializes and executes rules across 6 machine models."""
        pruner = ModelPruner()

        for model in ["Virgo", "Libra2", "Iris2024", "Sirius2", "Mebius", "Polaris"]:
            rules = pruner.get_rules_for_model(model)
            assert len(rules) > 0

            # Test specific part code rule in Mebius (302FP20250)
            if model == "Mebius":
                root = BOMNode(level=1, item_id="M_ROOT", item_name="MEBIUS ROOT")
                target_part = BOMNode(level=2, item_id="302FP20250", item_name="SPECIAL PART", has_children=False)
                other_part = BOMNode(level=2, item_id="OTHER123", item_name="OTHER PART", has_children=False)
                root.add_child(target_part)
                root.add_child(other_part)

                pruned = pruner.prune_tree(BOMTree(roots=[root]), model_name="Mebius")
                assert len(pruned.roots[0].children) == 1
                assert pruned.roots[0].children[0].item_id == "OTHER123"

    def test_f04_prune_excel_file_preserves_formatting(self, tmp_path):
        """Test 6: Verify prune_excel_file preserves 100% of formatting, headers, and column styles."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        excel_path = tmp_path / "PLM_test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Headers with custom style
        headers = ["Home", "Level", "Item Type", "Item Id", "Has Children", "Quantity", "Item Name"]
        ws.append(headers)
        ws.column_dimensions["D"].width = 25.0

        header_font = Font(name="Arial", size=12, bold=True)
        header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(1, col_idx)
            cell.font = header_font
            cell.fill = header_fill

        # Data rows
        # Row 2: Root (Keep)
        ws.append([1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"])
        # Row 3: Assembly (Keep)
        ws.append([2, 1, "Parts", "ASSY01", "True", 1, "BOTTLE WASTE"])
        # Row 4: Child of BOTTLE WASTE (Delete under Rule 2)
        ws.append([3, 2, "Parts", "CHILD01", "False", 1, "INTERNAL PART"])
        # Row 5: Another Level 1 part (Keep)
        ws.append([4, 1, "Parts", "PART02", "False", 1, "EXTERNAL COVER"])

        wb.save(str(excel_path))
        wb.close()

        rule = ModelRule(item_name="BOTTLE WASTE", match_mode="Full_name")
        pruner = ModelPruner()
        out_path = tmp_path / "PLM_filtered.xlsx"

        pruner.prune_excel_file(excel_path, out_path, rules=[rule])

        # Verify output
        wb_out = openpyxl.load_workbook(str(out_path))
        ws_out = wb_out.active

        # Check total surviving rows: Header + Root + ASSY01 + PART02 = 4 rows
        assert ws_out.max_row == 4
        # Check header preserved
        out_headers = [ws_out.cell(1, c).value for c in range(1, len(headers) + 1)]
        assert out_headers == headers
        # Check formatting preserved
        assert ws_out.cell(1, 1).font.name == "Arial"
        assert ws_out.cell(1, 1).font.bold is True
        assert ws_out.cell(1, 1).fill.fill_type == "solid"
        assert ws_out.column_dimensions["D"].width == 25.0

        # Check child row was pruned, PART02 now at row 4
        assert ws_out.cell(3, 4).value == "ASSY01"
        assert ws_out.cell(4, 4).value == "PART02"
        wb_out.close()

    def test_f04_normalize_model_name_variants(self):
        """Test 7: Verify normalize_model_name correctly maps real-world model names."""
        from src.core.model_pruner import normalize_model_name

        assert normalize_model_name("Libra235ppm") == "Libra2"
        assert normalize_model_name("libra2") == "Libra2"
        assert normalize_model_name("Sirius2(21ppm)") == "Sirius2"
        assert normalize_model_name("Sirius2") == "Sirius2"
        assert normalize_model_name("Polaris Next") == "Polaris"
        assert normalize_model_name("Polaris") == "Polaris"
        assert normalize_model_name("Virgo") == "Virgo"
        assert normalize_model_name("virgo") == "Virgo"
        assert normalize_model_name("Iris2024") == "Iris2024"
        assert normalize_model_name("Mebius") == "Mebius"
        assert normalize_model_name("UnknownModelXYZ") is None
        assert normalize_model_name("") is None
        assert normalize_model_name(None) is None

    def test_spica_pruning_rules_execution(self):
        """Test 8: Verify Spica decomposition rules (prune_node for 3V2YJ01090 and branch-free Level 2 rules)."""
        rule_pwb = ModelRule(
            item_name="PWB SWITCH ASSY",
            match_mode="Full_name",
            part_code="3V2YJ01090",
            action="prune_node",
            parent_part_code=None,
        )
        rule_cover = ModelRule(
            item_name="COVER LEFT HIGH M",
            match_mode="Full_name",
            part_code="31C1T04120",
            action="prune_children",
            parent_part_code=None,
        )
        rule_drum_part = ModelRule(
            item_name="DRUM ASSY WITH PACKING",
            match_mode="Full_name",
            part_code="31C0L11010",
            action="prune_children",
            parent_part_code="3VC0LM1110",
        )

        pruner = ModelPruner()

        # Build two model trees representing two variants (e.g. EU with TVC1TZ0030, US with TVC0LZ0020)
        for phantom_code in ["TVC1TZ0030", "TVC0LZ0020"]:
            root = BOMNode(level=0, item_id="MACHINE_ROOT", item_name="ROOT")
            phantom = BOMNode(level=1, item_id=phantom_code, item_name="PHANTOM")
            
            # PWB node with children
            pwb = BOMNode(level=2, item_id="3V2YJ01090", item_name="PWB SWITCH ASSY", has_children=True)
            pwb.add_child(BOMNode(level=3, item_id="CHILD1", item_name="RESISTOR"))
            
            # Cover node with children
            cover = BOMNode(level=2, item_id="31C1T04120", item_name="COVER LEFT HIGH M", has_children=True)
            cover.add_child(BOMNode(level=3, item_id="CHILD2", item_name="SCREW"))
            
            # Drum node with children
            drum = BOMNode(level=2, item_id="3VC0LM1110", item_name="DRUM UNIT", has_children=True)
            drum_part = BOMNode(level=3, item_id="31C0L11010", item_name="DRUM ASSY WITH PACKING", has_children=True)
            drum_part.add_child(BOMNode(level=4, item_id="SUB_PART", item_name="BEARING"))
            drum.add_child(drum_part)

            phantom.children.extend([pwb, cover, drum])
            root.add_child(phantom)

            tree = BOMTree(roots=[root])
            pruned = pruner.prune_tree(tree, rules=[rule_pwb, rule_cover, rule_drum_part])

            # 1. 3V2YJ01090 must be completely eliminated (prune_node)
            phantom_out = pruned.roots[0].children[0]
            remaining_ids = [c.item_id for c in phantom_out.children]
            assert "3V2YJ01090" not in remaining_ids, f"3V2YJ01090 must be deleted in variant {phantom_code}"
            assert "CHILD1" not in remaining_ids

            # 2. COVER LEFT HIGH M must be kept, but its children pruned
            assert "31C1T04120" in remaining_ids
            cover_out = next(c for c in phantom_out.children if c.item_id == "31C1T04120")
            assert len(cover_out.children) == 0, f"Children of 31C1T04120 must be pruned in variant {phantom_code}"

            # 3. Drum part under 3VC0LM1110 must have its children pruned
            assert "3VC0LM1110" in remaining_ids
            drum_out = next(c for c in phantom_out.children if c.item_id == "3VC0LM1110")
            drum_part_out = drum_out.children[0]
            assert drum_part_out.item_id == "31C0L11010"
            assert len(drum_part_out.children) == 0, "Children of 31C0L11010 must be pruned"


