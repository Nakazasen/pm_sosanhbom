"""Feature F1: 14-Column PLM Parser Isolation Tests.

Verifies:
1. Parsing authentic 14-column TC14 Excel exports into structured BOMTree.
2. Parsing 13-column legacy format gracefully.
3. Resilient column alias normalization (Vietnamese & English).
4. Positional fallback mapping when headers are unlabelled.
5. Handling nonexistent files and empty workbooks cleanly.
"""

from pathlib import Path
import openpyxl
import pytest

from src.core.models import BOMTree
from src.core.tree_parser import PLMTreeParser, parse_plm_excel


class TestF01PLMParser:
    """Test suite for Feature F1: 14-Column PLM Parser."""

    def test_f01_parse_valid_14col_excel(self, sample_plm_excel_file: Path):
        """Test 1: Parse authentic 14-column TC14 Excel file."""
        parser = PLMTreeParser()
        tree = parser.parse_excel(sample_plm_excel_file)

        assert isinstance(tree, BOMTree)
        assert tree.size() == 11
        assert len(tree.roots) == 3

        # Check root 1 (LSU UNIT)
        lsu = tree.roots[0]
        assert lsu.level == 1
        assert lsu.item_id == "302FP93010"
        assert lsu.item_name == "LSU UNIT"
        assert lsu.has_children is True
        assert len(lsu.children) == 2

        # Check polygon motor assy and its children
        polygon = lsu.children[0]
        assert polygon.item_id == "302FP94010"
        assert len(polygon.children) == 2
        assert polygon.children[0].item_id == "302FP02010"
        assert polygon.children[0].quantity == 1.0
        assert polygon.children[1].item_id == "B1303060"
        assert polygon.children[1].quantity == 4.0

    def test_f01_parse_13col_legacy_fallback(self, tmp_path: Path):
        """Test 2: Parse legacy 13-column format without Notice No column."""
        file_path = tmp_path / "PLM_13COL.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active

        headers = [
            "Row Index", "Level", "Item Type", "Item Id", "Has Children", "Quantity",
            "1st Parts", "2nd BOM Flag", "Occurrence Effectivities",
            "Item Revision Projects List", "Item Name", "Revision", "Item Rev Status"
        ]
        ws.append(headers)
        ws.append([1, 1, "Assembly", "UNIT01", "False", 1.0, "", "", "01-Jan-2024 UP", "PRJ", "MAIN UNIT", "A", "Released"])
        wb.save(file_path)
        wb.close()

        parser = PLMTreeParser()
        tree = parser.parse_excel(file_path)

        assert tree.size() == 1
        assert tree.roots[0].item_id == "UNIT01"
        assert tree.roots[0].item_name == "MAIN UNIT"

    def test_f01_column_alias_resilience(self, tmp_path: Path):
        """Test 3: Normalize Vietnamese headers commonly used in local factories."""
        file_path = tmp_path / "PLM_VIETNAMESE.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active

        headers = [
            "STT", "Cấp", "Loại", "Mã linh kiện", "Có con", "Số lượng",
            "1st Parts", "2nd BOM Flag", "Hiệu lực", "Dự án", "Tên linh kiện", "ECN", "Phiên bản", "Trạng thái"
        ]
        ws.append(headers)
        ws.append([1, 1, "Part", "302FP01010", "False", 2.0, "", "", "01-Jan-2024 UP", "PRJ", "LENS HOLDER", "N1", "01", "OK"])
        wb.save(file_path)
        wb.close()

        tree = parse_plm_excel(file_path)
        assert tree.size() == 1
        node = tree.roots[0]
        assert node.item_id == "302FP01010"
        assert node.quantity == 2.0
        assert node.item_name == "LENS HOLDER"

    def test_f01_positional_fallback_when_headers_missing(self, tmp_path: Path):
        """Test 4: Use positional mapping fallback when columns have generic or unlabelled headers."""
        file_path = tmp_path / "UNLABELLED_COLUMNS.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Col1", "Col2", "Col3", "Col4", "Col5", "Col6", "Col7", "Col8", "Col9", "Col10", "Col11", "Col12", "Col13", "Col14"])
        ws.append([1, 1, "Part", "FALLBACK_PART", "False", 1.0, "", "", "01-Jan-2024 UP", "PRJ", "NAME", "N", "01", "OK"])
        wb.save(file_path)
        wb.close()

        parser = PLMTreeParser()
        tree = parser.parse_excel(file_path)
        assert tree.size() >= 1

    def test_f01_nonexistent_and_empty_file_handling(self, tmp_path: Path):
        """Test 5: Raise FileNotFoundError for nonexistent file, return empty tree for empty workbook."""
        parser = PLMTreeParser()

        # Non-existent file
        with pytest.raises(FileNotFoundError):
            parser.parse_excel(tmp_path / "DOES_NOT_EXIST.xlsx")

        # Empty workbook
        file_path = tmp_path / "EMPTY.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.delete_rows(1, 10)
        wb.save(file_path)
        wb.close()

        tree = parser.parse_excel(file_path)
        assert tree.size() == 0
