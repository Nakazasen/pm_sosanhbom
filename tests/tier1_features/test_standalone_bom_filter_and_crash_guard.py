"""Unit and Integration Tests for Standalone BOM Filtering and Crash Guard Subsystems."""

from __future__ import annotations

import faulthandler
import logging
import os
from pathlib import Path
import tempfile
import threading
from typing import Any

import openpyxl
import pytest
from PyQt6.QtWidgets import QApplication, QMessageBox

from src.core.bom_filter_manager import BOMFilterManager
from src.core.crash_handler import (
    _collect_system_diagnostics,
    _write_crash_report,
    get_crash_logs_dir,
    open_logs_folder,
    resolve_logs_directory,
    setup_crash_protection,
    ssbom_excepthook,
    ssbom_threading_excepthook,
)
from src.core.model_pruner import ModelPruner, auto_detect_model_from_path
from src.core.models import ModelRule
from src.gui.bom_filter_dialog import BOMFilterConfigDialog
from src.gui.standalone_bom_filter_dialog import StandaloneBOMFilterDialog, StandalonePruneWorker


@pytest.fixture(scope="session")
def qapp():
    """Ensure a single QApplication instance exists for GUI tests."""
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


def create_sample_plm_file(file_path: Path, parts: list[tuple[int, int, str, str, str, int, str]]) -> Path:
    """Helper to create standard 14-column PLM test file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = [
        "Home", "Level", "Item Type", "Item Id", "Has Children", "Quantity",
        "1st Parts", "2nd BOM Flag", "Effectivity", "Projects List", "Item Name",
        "Notice No", "Revision", "Item Rev Status"
    ]
    ws.append(headers)
    for p in parts:
        ws.append([p[0], p[1], p[2], p[3], p[4], p[5], "", "", "", "", p[6], "", "A", "Released"])
    wb.save(str(file_path))
    wb.close()
    return file_path


# =============================================================================
# PART 1: Core ModelPruner Standalone & Re-filtering Tests
# =============================================================================

class TestModelPrunerStandalone:
    """Tests for ModelPruner.prune_file, prune_directory, and auto-detection."""

    def test_auto_detect_model_from_path(self):
        """Verify model auto-detection from filename and directory components."""
        # 1. Exact model in filename
        assert auto_detect_model_from_path("D:/boms/PLM_Spica_01.xlsx") == "Spica"
        assert auto_detect_model_from_path("PLM_Virgo_Machine.xlsx") == "Virgo"
        assert auto_detect_model_from_path("PLM_Libra2_Export.xlsm") == "Libra2"
        assert auto_detect_model_from_path("C:/Production/Iris2024/PLM_data.xlsx") == "Iris2024"

        # 2. Machine code in filename (C0V4 -> Spica / known dictionary mapping)
        # Even if dictionary is missing, known model in folder or stem works
        assert auto_detect_model_from_path("Spica/110C0L2US0/PLM_110C0L2US0.xlsx") == "Spica"
        assert auto_detect_model_from_path("unknown_file_xyz.xlsx") is None

    def test_prune_file_with_backup_and_stats(self, tmp_path: Path):
        """Verify prune_file executes in-place, creates a backup, and returns accurate stats."""
        plm_path = tmp_path / "PLM_Spica_test.xlsx"
        parts = [
            (1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"),
            (2, 1, "Parts", "ASSY01", "True", 1, "BOTTLE WASTE"),
            (3, 2, "Parts", "CHILD01", "False", 1, "SUB COMPONENT 1"),
            (4, 1, "Parts", "PART02", "False", 1, "NORMAL FRAME"),
        ]
        create_sample_plm_file(plm_path, parts)

        pruner = ModelPruner()
        rule = ModelRule(item_name="BOTTLE WASTE", match_mode="Full_name")

        res = pruner.prune_file(
            input_path=plm_path,
            rules=[rule],
            create_backup=True,
        )

        assert res["success"] is True
        assert res["initial_rows"] == 5  # header + 4 rows
        assert res["final_rows"] == 4    # header + ROOT01 + ASSY01 + PART02 (CHILD01 deleted)
        assert res["deleted_rows"] == 1
        assert res["backup_path"] is not None

        # Verify backup file exists and contains all 5 original rows
        backup_p = Path(res["backup_path"])
        assert backup_p.is_file()
        wb_b = openpyxl.load_workbook(str(backup_p))
        assert wb_b.active.max_row == 5
        wb_b.close()

        # Verify output file has 4 rows
        wb_out = openpyxl.load_workbook(str(plm_path))
        assert wb_out.active.max_row == 4
        wb_out.close()

    def test_prune_file_xoa_them_sequential_filtering(self, tmp_path: Path):
        """Verify business scenario 'xóa thêm (lọc sót)': re-pruning on an already pruned file.

        Simulates Spica scenario:
        Pass 1: Prune BOTTLE WASTE (CHILD01 removed).
        Pass 2: Later, add new rule for 3V2YJ01090 (NORMAL FRAME) and re-prune directly on disk.
        """
        plm_path = tmp_path / "PLM_Spica_rescan.xlsx"
        parts = [
            (1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"),
            (2, 1, "Parts", "ASSY01", "True", 1, "BOTTLE WASTE"),
            (3, 2, "Parts", "CHILD01", "False", 1, "SUB COMPONENT 1"),
            (4, 1, "Parts", "3V2YJ01090", "False", 1, "NORMAL FRAME"),
            (5, 1, "Parts", "KEEP01", "False", 1, "VALID MOTOR"),
        ]
        create_sample_plm_file(plm_path, parts)

        pruner = ModelPruner()

        # --- First Pass: Filter BOTTLE WASTE ---
        rule1 = ModelRule(item_name="BOTTLE WASTE", match_mode="Full_name")
        res1 = pruner.prune_file(plm_path, rules=[rule1], create_backup=True)
        assert res1["initial_rows"] == 6
        assert res1["final_rows"] == 5
        assert res1["deleted_rows"] == 1

        # --- Second Pass: Sequential re-filtering for 3V2YJ01090 on the SAME file ---
        rule2 = ModelRule(part_code="3V2YJ01090", match_mode="Full_name")
        res2 = pruner.prune_file(plm_path, rules=[rule2], create_backup=True)
        assert res2["initial_rows"] == 5
        assert res2["final_rows"] == 4  # 3V2YJ01090 row deleted!
        assert res2["deleted_rows"] == 1

        # Check surviving rows in final file
        wb_final = openpyxl.load_workbook(str(plm_path))
        ws_final = wb_final.active
        surviving_ids = [ws_final.cell(r, 4).value for r in range(2, ws_final.max_row + 1)]
        wb_final.close()

        assert "3V2YJ01090" not in surviving_ids
        assert "CHILD01" not in surviving_ids
        assert surviving_ids == ["ROOT01", "ASSY01", "KEEP01"]

    def test_prune_directory_batch(self, tmp_path: Path):
        """Verify batch directory pruning processes multiple files."""
        sub_dir = tmp_path / "boms_batch"
        sub_dir.mkdir()

        f1 = sub_dir / "PLM_Spica_01.xlsx"
        f2 = sub_dir / "PLM_Spica_02.xlsx"
        parts = [
            (1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"),
            (2, 1, "Parts", "ASSY01", "True", 1, "BOTTLE WASTE"),
            (3, 2, "Parts", "CHILD01", "False", 1, "SUB 1"),
        ]
        create_sample_plm_file(f1, parts)
        create_sample_plm_file(f2, parts)

        pruner = ModelPruner()
        rule = ModelRule(item_name="BOTTLE WASTE")

        results = pruner.prune_directory(sub_dir, rules=[rule], recursive=False, file_pattern="PLM_*.xlsx")
        assert len(results) == 2
        for r in results:
            assert r["success"] is True
            assert r["deleted_rows"] == 1


# =============================================================================
# PART 2: Crash Handler & Stability Tests
# =============================================================================

class TestCrashProtection:
    """Tests for crash prevention, faulthandler, and exception logging."""

    def test_setup_crash_protection(self, tmp_path: Path):
        """Verify crash protection sets up log directory, faulthandler, and file logging."""
        log_dir = setup_crash_protection(base_dir=tmp_path)
        assert log_dir.is_dir()
        assert (log_dir / "crash_fault.log").exists()

        # faulthandler should be active
        assert faulthandler.is_enabled()

    def test_unhandled_main_thread_exception_hook(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        """Verify ssbom_excepthook generates timestamped diagnostic report and does not crash silently."""
        log_dir = setup_crash_protection(base_dir=tmp_path)
        monkeypatch.setattr("src.core.crash_handler._show_gui_crash_dialog", lambda *args, **kwargs: None)

        try:
            raise RuntimeError("Test simulated crash error in SSBOM")
        except RuntimeError as exc:
            import sys
            ssbom_excepthook(type(exc), exc, exc.__traceback__)

        # Check that crash report was generated
        crash_reports = list(log_dir.glob("crash_*.log"))
        assert len(crash_reports) >= 1

        report_content = crash_reports[0].read_text(encoding="utf-8")
        assert "SSBOM MANAGER CRASH DIAGNOSTIC REPORT" in report_content
        assert "Test simulated crash error in SSBOM" in report_content
        assert "RuntimeError" in report_content
        assert "STACK TRACE:" in report_content

    def test_unhandled_threading_excepthook(self, tmp_path: Path):
        """Verify threading exception hook catches background thread failures."""
        log_dir = setup_crash_protection(base_dir=tmp_path)

        class ThreadArgs:
            def __init__(self, exc):
                self.exc_type = type(exc)
                self.exc_value = exc
                self.exc_traceback = exc.__traceback__
                self.thread = threading.Thread(name="WorkerSimulationThread")

        try:
            raise ValueError("Simulation background thread explosion")
        except ValueError as exc:
            args = ThreadArgs(exc)
            ssbom_threading_excepthook(args)

        thread_reports = list(log_dir.glob("crash_WorkerSimulationThread_*.log"))
        assert len(thread_reports) >= 1
        content = thread_reports[0].read_text(encoding="utf-8")
        assert "Simulation background thread explosion" in content

    def test_diagnostics_collection(self):
        """Verify system diagnostics payload contains required fields."""
        diag = _collect_system_diagnostics()
        assert "timestamp" in diag
        assert "platform" in diag
        assert "python_version" in diag
        assert "active_threads" in diag


# =============================================================================
# PART 3: StandaloneBOMFilterDialog UI Tests
# =============================================================================

class TestStandaloneBOMFilterDialogUI:
    """UI Tests for StandaloneBOMFilterDialog."""

    def test_dialog_init_and_model_loading(self, qapp: QApplication, tmp_path: Path):
        """Verify standalone filter dialog initializes cleanly with model list."""
        db_path = tmp_path / "test_master.db"
        mgr = BOMFilterManager(base_dir=tmp_path, remote_db_path=db_path)

        dlg = StandaloneBOMFilterDialog(
            initial_model="Spica",
            filter_manager=mgr,
        )
        assert dlg.combo_model.count() >= 1
        assert "Spica" in [dlg.combo_model.itemText(i) for i in range(dlg.combo_model.count())]
        assert dlg.chk_backup.isChecked() is True
        dlg.close()

    def test_load_files_and_execute_worker(self, qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        """Verify loading files into table and executing background prune worker."""
        monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: None)
        db_path = tmp_path / "test_master.db"
        mgr = BOMFilterManager(base_dir=tmp_path, remote_db_path=db_path)

        f1 = tmp_path / "PLM_Spica_Unit1.xlsx"
        parts = [
            (1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"),
            (2, 1, "Parts", "ASSY01", "True", 1, "BOTTLE WASTE"),
            (3, 2, "Parts", "CHILD01", "False", 1, "SUB PART"),
        ]
        create_sample_plm_file(f1, parts)

        dlg = StandaloneBOMFilterDialog(
            initial_model="Spica",
            filter_manager=mgr,
        )
        dlg.rb_mode_files.setChecked(True)
        dlg._load_selected_files([f1])

        assert dlg.table_files.rowCount() == 1
        assert dlg.table_files.item(0, 1).text() == "PLM_Spica_Unit1.xlsx"
        assert dlg.table_files.item(0, 2).text() == "Spica"

        # Run worker directly
        files_payload = [(0, f1, "Spica")]
        worker = StandalonePruneWorker(
            files=files_payload,
            filter_manager=mgr,
            create_backup=True,
        )

        completed_events = []
        worker.file_finished_signal.connect(lambda row, res: completed_events.append((row, res)))
        worker.run()

        assert len(completed_events) == 1
        assert completed_events[0][1]["success"] is True
        dlg.close()

    def test_open_standalone_button_in_bom_filter_dialog(self, qapp: QApplication, tmp_path: Path):
        """Verify BOMFilterConfigDialog has the '⚡ Áp Dụng Cho File Có Sẵn...' button wired up."""
        db_path = tmp_path / "test_master.db"
        mgr = BOMFilterManager(base_dir=tmp_path, remote_db_path=db_path)

        config_dlg = BOMFilterConfigDialog(filter_manager=mgr, initial_model="Virgo")
        assert hasattr(config_dlg, "btn_apply_existing")
        assert config_dlg.btn_apply_existing.text() == "⚡ Áp Dụng Cho File Có Sẵn..."
        config_dlg.close()

    def test_bom_scan_dialog_has_apply_filter_button(self, qapp: QApplication, tmp_path: Path):
        """Verify BOMScanDialog (Quét trạng thái BOM) has the '⚡ Lọc lại file có sẵn...' button."""
        from src.gui.plm_download_dialog import BOMScanDialog
        dlg = BOMScanDialog(dest_dir=tmp_path)
        assert hasattr(dlg, "btn_apply_filter")
        assert dlg.btn_apply_filter.text() == "⚡ Lọc lại file có sẵn..."
        dlg.close()

    def test_prune_file_rejects_old_xls_format(self, tmp_path: Path):
        """Verify prune_file raises a clear ValueError when given legacy .xls binary format."""
        xls_file = tmp_path / "legacy_bom.xls"
        xls_file.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")  # OLE CFB header
        pruner = ModelPruner()
        with pytest.raises(ValueError, match="Định dạng tệp Excel cũ"):
            pruner.prune_file(xls_file)

    def test_prune_file_permission_error_on_save(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        """Verify prune_file catches save PermissionError and provides a clear Excel lock explanation."""
        plm_path = tmp_path / "PLM_Locked.xlsx"
        parts = [
            (1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"),
            (2, 1, "Parts", "ASSY01", "True", 1, "BOTTLE WASTE"),
            (3, 2, "Parts", "CHILD01", "False", 1, "SUB PART"),
        ]
        create_sample_plm_file(plm_path, parts)

        pruner = ModelPruner()
        rule = ModelRule(item_name="BOTTLE WASTE")

        # Mock openpyxl save to simulate live Windows file lock by Excel
        def mock_save(self, filename):
            raise PermissionError("[Errno 13] Permission denied")

        monkeypatch.setattr("openpyxl.Workbook.save", mock_save)

        with pytest.raises(PermissionError, match="đang mở trong Microsoft Excel"):
            pruner.prune_file(plm_path, rules=[rule])

    def test_prune_file_zero_rows_deleted_skips_unnecessary_save(self, tmp_path: Path):
        """Verify prune_file does not touch or re-save the file if 0 rows are pruned."""
        plm_path = tmp_path / "PLM_Untouched.xlsx"
        parts = [
            (1, 0, "Parts", "ROOT01", "True", 1, "MAIN UNIT"),
            (2, 1, "Parts", "KEEP01", "False", 1, "GOOD PART"),
        ]
        create_sample_plm_file(plm_path, parts)
        mtime_before = plm_path.stat().st_mtime_ns

        pruner = ModelPruner()
        rule = ModelRule(item_name="NON_EXISTENT_PART")
        res = pruner.prune_file(plm_path, rules=[rule])

        assert res["deleted_rows"] == 0
        mtime_after = plm_path.stat().st_mtime_ns
        assert mtime_before == mtime_after

    def test_show_gui_crash_dialog_safely_ignored_on_background_thread(self, tmp_path: Path):
        """Verify _show_gui_crash_dialog avoids opening QMessageBox when called from a worker thread."""
        from src.core.crash_handler import _show_gui_crash_dialog
        report = tmp_path / "crash_test.log"
        report.write_text("crash dump", encoding="utf-8")

        result = []
        def _bg():
            try:
                _show_gui_crash_dialog(report, RuntimeError("Test"))
                result.append("ok")
            except Exception as e:
                result.append(f"err: {e}")

        t = threading.Thread(target=_bg)
        t.start()
        t.join(timeout=2.0)

        assert result == ["ok"]

    def test_resolve_logs_directory_respects_custom_base_dir(self, tmp_path: Path):
        """Verify resolve_logs_directory respects an explicit base_dir even if previously cached."""
        custom_base = tmp_path / "custom_app_dir"
        resolved = resolve_logs_directory(base_dir=custom_base)
        assert resolved == (custom_base / "logs").resolve()
        assert resolved.is_dir()

    def test_dialog_close_event_with_running_worker(self, qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        """Verify closeEvent and reject prompt user and gracefully stop worker thread without crashing Qt."""
        from PyQt6.QtGui import QCloseEvent
        from unittest.mock import MagicMock

        dlg = StandaloneBOMFilterDialog(initial_model="Spica")
        dlg.worker = MagicMock()
        dlg.worker.isRunning.return_value = True

        # When user confirms close
        monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes)
        close_evt = QCloseEvent()
        dlg.closeEvent(close_evt)

        assert close_evt.isAccepted() is True
        dlg.worker.cancel.assert_called_once()
        dlg.worker.wait.assert_called_once()
        dlg.close()
