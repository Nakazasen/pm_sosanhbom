"""Feature F14: TC14 BOM Export Pipeline Isolation Tests.

Verifies:
1. Native Excel export via 'Awp0ExportToExcel' and file download capture.
2. Handling export timeout with structured TC14ExportError.
3. Interactive occurrence tree scraper as a secondary fallback.
4. OpenXML workbook synthesis from scraped occurrence rows.
5. Strict fulfillment of the Application Layer Interface Contract:
   download_bom_full(part_number: str, output_dir: Path) -> Path
"""

from pathlib import Path
from unittest.mock import MagicMock, patch
import openpyxl
import pytest
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.remote.webelement import WebElement

from src.automation.tc14.client import (
    BOMItem,
    SearchResult,
    TC14AutomationClient,
    TC14ExportError,
)
from src.automation.tc14.session import TC14SessionManager


class TestF14TC14ExportPipeline:
    """Test suite for Feature F14: TC14 BOM Export Pipeline."""

    def test_f14_native_export_to_excel_success(self, tmp_path: Path, mock_tc14_driver):
        """Test 1: Native Excel export triggers dialog, confirms, and captures file."""
        export_btn = MagicMock(spec=WebElement)
        dialog_btn = MagicMock(spec=WebElement)

        download_dir = tmp_path / "downloads"
        download_dir.mkdir(parents=True, exist_ok=True)
        # Pre-create downloaded file to simulate browser saving file
        simulated_file = download_dir / "PLM_110C103NL0.xlsx"
        simulated_file.write_bytes(b"PK\x03\x04")

        session_mgr = TC14SessionManager(driver_factory=lambda: mock_tc14_driver)
        session_mgr.start_session()
        session_mgr.mark_authenticated("vn_pe03")

        client = TC14AutomationClient(session_manager=session_mgr, download_dir=download_dir)
        client._poll_for_downloaded_file = MagicMock(return_value=simulated_file)

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_inst = MagicMock()
            mock_wait_cls.return_value = wait_inst
            # 1: export_cmd clickable, 2: dialog present, 3: confirm_btn clickable
            wait_inst.until.side_effect = [export_btn, MagicMock(), dialog_btn]

            out_file = client.export_bom_excel(timeout=5)
            assert out_file == simulated_file
            export_btn.click.assert_called_once()
            dialog_btn.click.assert_called_once()

    def test_f14_export_wait_timeout_raises_error(self, tmp_path: Path, mock_tc14_driver):
        """Test 2: Raise TC14ExportError when export dialog or file download fails."""
        session_mgr = TC14SessionManager(driver_factory=lambda: mock_tc14_driver)
        session_mgr.start_session()
        session_mgr.mark_authenticated("vn_pe03")

        client = TC14AutomationClient(session_manager=session_mgr, download_dir=tmp_path)

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_inst = MagicMock()
            mock_wait_cls.return_value = wait_inst
            wait_inst.until.side_effect = TimeoutException("Export button timed out")

            with pytest.raises(TC14ExportError):
                client.export_bom_excel(timeout=1)

    def test_f14_fallback_interactive_tree_scraper(self, mock_tc14_driver):
        """Test 3: Extract occurrence items from DOM virtual table rows."""
        c1 = MagicMock(spec=WebElement)
        c1.text = "302FP93010"
        c2 = MagicMock(spec=WebElement)
        c2.text = "LSU UNIT"

        row1 = MagicMock(spec=WebElement)
        row1.get_attribute.side_effect = lambda attr: "1" if attr == "aria-level" else "true"
        row1.text = "302FP93010\tLSU UNIT"
        row1.find_elements.return_value = [c1, c2]

        c3 = MagicMock(spec=WebElement)
        c3.text = "302FP02010"
        c4 = MagicMock(spec=WebElement)
        c4.text = "MOTOR BRACKET"

        row2 = MagicMock(spec=WebElement)
        row2.get_attribute.side_effect = lambda attr: "2" if attr == "aria-level" else "false"
        row2.text = "302FP02010\tMOTOR BRACKET"
        row2.find_elements.return_value = [c3, c4]

        mock_tc14_driver.find_elements.return_value = [row1, row2]

        session_mgr = TC14SessionManager(driver_factory=lambda: mock_tc14_driver)
        session_mgr.start_session()
        session_mgr.mark_authenticated("vn_pe03")

        client = TC14AutomationClient(session_manager=session_mgr)
        items = client.scrape_bom_tree_interactive(max_depth=3, expand_children=False)

        assert len(items) == 2
        assert items[0].level == 1
        assert items[0].item_id == "302FP93010"
        assert items[0].item_name == "LSU UNIT"
        assert items[1].level == 2
        assert items[1].item_id == "302FP02010"
        assert items[1].item_name == "MOTOR BRACKET"

    def test_f14_synthesize_openxml_from_scraped_items(self, tmp_path: Path, mock_tc14_driver):
        """Test 4: Synthesize standard 14-column PLM workbook from scraped BOM items."""
        session_mgr = TC14SessionManager(driver_factory=lambda: mock_tc14_driver)
        client = TC14AutomationClient(session_manager=session_mgr)

        items = [
            BOMItem(level=1, item_id="UNIT1", item_name="UNIT 1", has_children=True, quantity=1.0, effectivity="01-Jan-2024 UP", revision="A"),
            BOMItem(level=2, item_id="PART1", item_name="PART 1", has_children=False, quantity=2.0, effectivity="01-Jan-2024 UP", revision="01"),
        ]

        out_xlsx = tmp_path / "SYNTHESIZED_PLM.xlsx"
        created_path = client.synthesize_plm_excel_from_tree(items, out_xlsx)

        assert created_path.exists()
        wb = openpyxl.load_workbook(created_path)
        ws = wb.active
        assert ws.max_row == 3  # Header + 2 items
        assert ws.cell(2, 2).value == 1  # Level
        assert ws.cell(2, 4).value == "UNIT1"  # Item Id
        wb.close()

    def test_f14_download_bom_full_contract_fulfillment(self, tmp_path: Path, mock_tc14_driver):
        """Test 5: Fulfill Application Layer Contract download_bom_full."""
        target_dir = tmp_path / "output_bom"
        target_dir.mkdir()

        session_mgr = TC14SessionManager(driver_factory=lambda: mock_tc14_driver)
        session_mgr.start_session()
        session_mgr.mark_authenticated("vn_pe03")

        client = TC14AutomationClient(session_manager=session_mgr)
        client.login = MagicMock(return_value=True)
        client.search_part = MagicMock(return_value=[SearchResult(item_id="110C103NL0", title="MA4500", uid="U1")])
        client.open_item = MagicMock()
        client.navigate_to_content_tab = MagicMock(return_value=True)

        expected_dest = target_dir / "PLM_110C103NL0.xlsx"
        expected_dest.write_text("DUMMY PLM DATA")
        client.export_bom_excel = MagicMock(return_value=expected_dest)

        result_path = client.download_bom_full("110C103NL0", target_dir)
        assert result_path.exists()
        assert result_path.name == "PLM_110C103NL0.xlsx"
