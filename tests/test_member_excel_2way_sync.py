"""Unit tests for Bi-directional (2-way) Member Sync between SQLite Database and Danhsachthanhvien.xlsm.

Tests cover:
1. Importing member roster from Excel (Columns C, D, E, F, G).
2. Handling missing account_id and generating deterministic ID.
3. Exporting member roster from DB back to Excel (updating in-place, appending new rows, .xlsm.bak backup).
4. GUI MemberRosterDialog 2-way sync buttons and reload actions.
5. LeaderView Table 1.3 display format 'Họ và tên (Mã thành viên)' and model-based auto tick-checking.
"""

from __future__ import annotations

import os
import shutil
from pathlib import Path
import pytest
import openpyxl
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication, QCheckBox, QComboBox

from src.core.member_database import (
    DEFAULT_MEMBER_EXCEL_PATH,
    MemberDatabaseManager,
    MemberRecord,
)
from src.gui.leader_view import LeaderWorkspaceView
from src.gui.member_roster_dialog import MemberRosterDialog


@pytest.fixture
def mock_excel_xlsm(tmp_path: Path) -> Path:
    """Create a mock Danhsachthanhvien.xlsm with valid header and sample rows."""
    excel_path = tmp_path / "Danhsachthanhvien.xlsm"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DS KDTVN"

    # Header at Row 4
    headers = [
        "STT",
        "Mã \nnhân viên\n社員番号",
        "Mã thành viên",
        "Họ và tên\n氏名",
        "Phòng ban\n課",
        "Công đoạn mặc định",
        "Dòng máy phụ trách",
        "Bộ phận\n部門",
        "Mã code",
    ]
    for col_idx, h in enumerate(headers, 1):
        ws.cell(row=4, column=col_idx, value=h)

    # Sample rows
    sample_data = [
        (1, "V1001", "KhiemA_mecha1.1", "Nguyễn Ích Khiêm", "Cơ 1.1", "LSU", "Virgo"),
        (2, "V1002", "Ly_mecha1.2", "Đỗ Thị Ly", "Cơ 1.2", "DRUM", "Virgo, Iris2024"),
        (3, "V1003", "Son_mecha1", "Nguyễn Văn Sơn", "Cơ 1", "FUSER", ""),
        (4, "V1004", "", "Phạm Hải Đăng", "KTCT Điện", "", "Corvus"),  # Missing account_id
    ]

    for row_idx, r_data in enumerate(sample_data, 5):
        ws.cell(row=row_idx, column=1, value=r_data[0])
        ws.cell(row=row_idx, column=2, value=r_data[1])
        ws.cell(row=row_idx, column=3, value=r_data[2])
        ws.cell(row=row_idx, column=4, value=r_data[3])
        ws.cell(row=row_idx, column=5, value=r_data[4])
        ws.cell(row=row_idx, column=6, value=r_data[5])
        ws.cell(row=row_idx, column=7, value=r_data[6])
        ws.cell(row=row_idx, column=8, value="Bộ phận kỹ thuật chế tạo")

    wb.save(str(excel_path))
    wb.close()
    return excel_path


class TestMemberExcelTwoWaySync:
    """Test suite for bi-directional Excel sync service."""

    def test_import_from_excel_maps_columns_correctly(
        self,
        tmp_path: Path,
        mock_excel_xlsm: Path,
    ) -> None:
        """Verify import reads Cols C, D, E, F, G and upserts into database."""
        db_path = tmp_path / "ssbom_master.db"
        db = MemberDatabaseManager(
            base_dir=tmp_path,
            remote_db_path=db_path,
            excel_path=mock_excel_xlsm,
        )

        total, updated, msg = db.import_from_excel(mock_excel_xlsm)
        assert total == 4
        assert "thành công" in msg.lower()

        members = {m.account_id: m for m in db.get_members()}

        # 1. KhiemA
        assert "KhiemA_mecha1.1" in members
        khiem = members["KhiemA_mecha1.1"]
        assert khiem.full_name == "Nguyễn Ích Khiêm"
        assert khiem.department == "Cơ 1.1"
        assert khiem.default_sub_unit == "LSU"
        assert khiem.machine_names == "Virgo"

        # 2. Ly
        assert "Ly_mecha1.2" in members
        ly = members["Ly_mecha1.2"]
        assert ly.full_name == "Đỗ Thị Ly"
        assert ly.department == "Cơ 1.2"
        assert ly.default_sub_unit == "DRUM"
        assert "Iris2024" in ly.machine_names

        # 3. Missing account_id auto-generation
        # Should generate a deterministic ID containing HaiDang and dien
        haidang = next((m for m in members.values() if "Hải Đăng" in m.full_name), None)
        assert haidang is not None
        assert haidang.account_id != ""
        assert "haidang" in haidang.account_id.lower() or "dien" in haidang.account_id.lower()
        assert haidang.department == "KTCT Điện"
        assert haidang.machine_names == "Corvus"

    def test_export_to_excel_updates_in_place_and_creates_backup(
        self,
        tmp_path: Path,
        mock_excel_xlsm: Path,
    ) -> None:
        """Verify export writes back changes in-place, creates .xlsm.bak, and appends new members."""
        db_path = tmp_path / "ssbom_master.db"
        db = MemberDatabaseManager(
            base_dir=tmp_path,
            remote_db_path=db_path,
            excel_path=mock_excel_xlsm,
        )
        db.import_from_excel(mock_excel_xlsm)

        # Update Khiem's machine and subunit in DB
        khiem = next(m for m in db.get_members() if m.account_id == "KhiemA_mecha1.1")
        khiem.default_sub_unit = "FUSER"
        khiem.machine_names = "Virgo, Corvus"
        db.update_member(khiem, sync_to_excel=False)

        # Add brand new member in DB
        new_m = MemberRecord(
            account_id="Tung_mecha1.3",
            full_name="Nguyễn Bá Tùng",
            department="Cơ 1.3",
            default_sub_unit="LASER",
            machine_names="Iris2024",
        )
        db.add_member(new_m, sync_to_excel=False)

        # Export to Excel
        ok, msg = db.export_to_excel(mock_excel_xlsm)
        assert ok is True
        assert "thành công" in msg

        # Verify .xlsm.bak exists
        bak_file = mock_excel_xlsm.with_suffix(".xlsm.bak")
        assert bak_file.exists()

        # Load exported Excel and verify
        wb = openpyxl.load_workbook(mock_excel_xlsm, data_only=True)
        ws = wb["DS KDTVN"]

        # Row 5 (KhiemA) should be updated in place
        assert ws.cell(row=5, column=3).value == "KhiemA_mecha1.1"
        assert ws.cell(row=5, column=6).value == "FUSER"
        assert ws.cell(row=5, column=7).value == "Virgo, Corvus"

        # New member should be appended
        found_new = False
        for r in range(5, ws.max_row + 1):
            if ws.cell(row=r, column=3).value == "Tung_mecha1.3":
                found_new = True
                assert ws.cell(row=r, column=4).value == "Nguyễn Bá Tùng"
                assert ws.cell(row=r, column=5).value == "Cơ 1.3"
                assert ws.cell(row=r, column=6).value == "LASER"
                assert ws.cell(row=r, column=7).value == "Iris2024"
                break
        assert found_new is True
        wb.close()

    def test_roster_dialog_excel_buttons(
        self,
        qapp: QApplication,
        tmp_path: Path,
        mock_excel_xlsm: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Verify MemberRosterDialog has Excel buttons and import action updates table."""
        from PyQt6.QtWidgets import QMessageBox
        monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
        monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
        monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
        monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes)

        dlg = MemberRosterDialog(base_dir=tmp_path)
        dlg.db_manager.excel_path = mock_excel_xlsm

        # Check buttons exist
        assert hasattr(dlg, "btn_import_excel")
        assert hasattr(dlg, "btn_export_excel")
        assert hasattr(dlg, "btn_open_excel")

        # Click import
        dlg._on_import_excel_clicked()

        # Check table has loaded members
        records = dlg._all_records
        acc_ids = [r.account_id for r in records]
        assert "KhiemA_mecha1.1" in acc_ids
        assert "Ly_mecha1.2" in acc_ids
        dlg.close()

    def test_leader_view_table_1_3_display_and_model_auto_check(
        self,
        qapp: QApplication,
        tmp_path: Path,
        mock_excel_xlsm: Path,
    ) -> None:
        """Verify LeaderView Table 1.3 displays 'Họ và tên (Mã)' and auto-ticks based on project model."""
        db_path = tmp_path / "ssbom_master.db"
        db = MemberDatabaseManager(
            base_dir=tmp_path,
            remote_db_path=db_path,
            excel_path=mock_excel_xlsm,
        )
        db.import_from_excel(mock_excel_xlsm)

        view = LeaderWorkspaceView(base_dir=tmp_path)
        step1 = view.step1_widget

        # Set project model to 'Virgo'
        step1.model_combo.setCurrentText("Virgo")
        step1._reload_roster_from_db(force_model_match=True)

        table = step1.staff_table
        assert table.rowCount() >= 4

        # Find row for KhiemA (Virgo), Ly (Virgo, Iris2024), Son (empty machine), HaiDang (Corvus)
        rows_by_acc: dict[str, int] = {}
        for r in range(table.rowCount()):
            acc_id = table.item(r, 1).data(Qt.ItemDataRole.UserRole)
            if acc_id:
                rows_by_acc[acc_id] = r

        assert "KhiemA_mecha1.1" in rows_by_acc
        khiem_row = rows_by_acc["KhiemA_mecha1.1"]

        # Check Col 1 format: 'Nguyễn Ích Khiêm (KhiemA_mecha1.1)'
        khiem_text = table.item(khiem_row, 1).text()
        assert "Nguyễn Ích Khiêm" in khiem_text
        assert "(KhiemA_mecha1.1)" in khiem_text

        # Khiem has model Virgo -> Checkbox should be checked (True)
        chk_khiem = table.cellWidget(khiem_row, 0).findChild(QCheckBox)
        assert chk_khiem.isChecked() is True

        # Ly has model Virgo, Iris2024 -> Checkbox should be checked (True)
        ly_row = rows_by_acc["Ly_mecha1.2"]
        chk_ly = table.cellWidget(ly_row, 0).findChild(QCheckBox)
        assert chk_ly.isChecked() is True

        # Son has empty machine -> Checkbox should be checked (True - applies to all models)
        son_row = rows_by_acc["Son_mecha1"]
        chk_son = table.cellWidget(son_row, 0).findChild(QCheckBox)
        assert chk_son.isChecked() is True

        # Hai Dang (KTCT Điện) has model Corvus -> For project Virgo, Checkbox should NOT be checked (False)
        assert "HaiDang_dien" in rows_by_acc
        haidang_row = rows_by_acc["HaiDang_dien"]
        chk_haidang = table.cellWidget(haidang_row, 0).findChild(QCheckBox)
        assert chk_haidang.isChecked() is False

    def test_import_from_xlsx_with_empty_subunit_and_spaces_in_machines(
        self,
        tmp_path: Path,
    ) -> None:
        """Verify .xlsx file import preserves spaces in model names and handles empty subunit (Col 6)."""
        xlsx_path = tmp_path / "Danhsachthanhvien.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DS KDTVN"

        # Headers at Row 4
        headers = [
            "STT",
            "Mã \nnhân viên\n社員番号",
            "Mã thành viên",
            "Họ và tên\n氏名",
            "Phòng ban\n課",
            "Công đoạn mặc định",
            "Dòng máy phụ trách(Tên máy)",
            "Bộ phận\n部門",
            "Mã code",
        ]
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=4, column=col_idx, value=h)

        # Real test data matching production file
        rows = [
            (1, "VN120204", "Huy_mecha2.2", "Nguyễn Quang Huy", "Cơ 2.2", None, "Spica, Virgo, VIRGO2, Kairos"),
            (2, "VN120205", "Hai_mecha1.2", "Bùi Văn Hải", "Cơ 1.2", None, "Iris2020, Iris 2024, Sirius 2 (21ppm),Sirius 2 (26ppm)"),
            (3, "VN120206", "Phan_mecha1.1", "Trần Đình Phan", "Cơ 1.1", None, "Libra 1, Libra 2, 6th A3, 6th A3 KDC"),
            (4, "VN260269", None, "Phạm Hải Đăng", "KTCT Điện", None, None),
        ]
        for r_idx, r_val in enumerate(rows, 5):
            for c_idx, val in enumerate(r_val, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)

        wb.save(str(xlsx_path))
        wb.close()

        db_path = tmp_path / "ssbom_master.db"
        mgr = MemberDatabaseManager(base_dir=tmp_path, remote_db_path=db_path, excel_path=xlsx_path)

        total, updated, msg = mgr.import_from_excel(xlsx_path, clean_legacy_seeds=True)
        assert total == 4

        members = {m.account_id: m for m in mgr.get_members()}
        assert "Hai_mecha1.2" in members
        hai = members["Hai_mecha1.2"]
        assert hai.default_sub_unit == ""
        # Check preserved spaces in machine names
        assert "Sirius 2 (21ppm)" in hai.machine_names
        assert "Sirius 2 (26ppm)" in hai.machine_names
        assert "Iris 2024" in hai.machine_names

        phan = members["Phan_mecha1.1"]
        assert "6th A3" in phan.machine_names
        assert "6th A3 KDC" in phan.machine_names
        assert "Libra 1" in phan.machine_names

        # Test model query
        sirius_members = mgr.get_members_for_machine("Sirius 2 (21ppm)")
        assert any(m.account_id == "Hai_mecha1.2" for m in sirius_members)

        a3_members = mgr.get_members_for_machine("6th A3")
        assert any(m.account_id == "Phan_mecha1.1" for m in a3_members)
