"""Regression tests: .xlsm BOM stays a valid macro-enabled workbook after PLM/R3 update,
and the update dialog stays readable in dark theme."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest

from src.core import inheritance_engine as ie
from src.core.inheritance_engine import InheritanceEngine

FAKE_VBA = b"\xd0\xcf\x11\xe0fake-vba-project"


def _make_xlsm(path: Path) -> None:
    """Build a minimal macro-enabled workbook with a PLM sheet and a vbaProject part."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PLM"
    ws.append(["Level", "Type", "Item Id", "Has Children", "Qty"] + [""] * 9 + ["Giai thich", "PIC", "Check"])
    ws.append([1, "Part", "P-001", "False", 1] + [""] * 9 + ["old note", "An", "OK"])
    tmp = path.with_suffix(".xlsx")
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                ).replace(
                    b"</Types>",
                    b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/'
                    b'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
                )
            dst.writestr(item, data)
        dst.writestr("xl/vbaProject.bin", FAKE_VBA)
    tmp.unlink()


def _workbook_content_type(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        ct = z.read("[Content_Types].xml").decode()
    m = re.search(r'PartName="/xl/workbook\.xml"\s+ContentType="([^"]+)"', ct) or re.search(
        r'ContentType="([^"]+)"\s+PartName="/xl/workbook\.xml"', ct
    )
    return m.group(1) if m else ""


@pytest.fixture
def xlsm(tmp_path: Path) -> Path:
    p = tmp_path / "BOM_TEST.xlsm"
    _make_xlsm(p)
    return p


def test_plm_update_fallback_keeps_xlsm_macro_enabled(xlsm: Path, tmp_path: Path) -> None:
    # Arrange: force the openpyxl fallback (no Excel available)
    new_plm = pd.DataFrame([{"level": 1, "item_type": "Part", "item_id": "P-001", "quantity": 2},
                            {"level": 1, "item_type": "Part", "item_id": "P-002", "quantity": 1}])
    with patch.object(ie, "excel_com_update_sheet", return_value=False), \
         patch.object(InheritanceEngine, "refresh_workbook_pivots", return_value=False):
        # Act
        summary = InheritanceEngine().update_workbook_with_inheritance(
            xlsm, new_plm, archive_dir=tmp_path / "old"
        )

    # Assert: still a macro-enabled package with VBA, data written, annotation inherited
    assert _workbook_content_type(xlsm) == "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
    with zipfile.ZipFile(xlsm) as z:
        assert "xl/vbaProject.bin" in z.namelist()
    wb = openpyxl.load_workbook(xlsm, keep_vba=True)
    assert wb["PLM"]["C3"].value == "P-002"
    assert wb["PLM"]["O2"].value == "old note"
    assert summary.archived_sheet_name in wb.sheetnames
    assert summary.inherited_count == 1


def test_plm_update_creates_backup_before_writing(xlsm: Path, tmp_path: Path) -> None:
    original = xlsm.read_bytes()
    with patch.object(ie, "excel_com_update_sheet", return_value=False), \
         patch.object(InheritanceEngine, "refresh_workbook_pivots", return_value=False):
        InheritanceEngine().update_workbook_with_inheritance(
            xlsm, pd.DataFrame([{"item_id": "X"}]), archive_dir=tmp_path / "old"
        )
    backups = list((tmp_path / "old").glob("BOM_TEST_backup_*.xlsm"))
    assert len(backups) == 1
    assert backups[0].read_bytes() == original


def test_plm_update_prefers_excel_com_and_skips_openpyxl_save(xlsm: Path, tmp_path: Path) -> None:
    with patch.object(ie, "excel_com_update_sheet", return_value=True) as com, \
         patch.object(InheritanceEngine, "_write_rows_openpyxl") as fallback:
        InheritanceEngine().update_workbook_with_inheritance(
            xlsm, pd.DataFrame([{"item_id": "P-001"}]), archive_dir=tmp_path / "old"
        )
    com.assert_called_once()
    _, sheet, _, rows = com.call_args.args[:4]
    assert sheet == "PLM"
    assert rows[0][2] == "P-001" and rows[0][17] == '=IF(C2="","",C2)'
    fallback.assert_not_called()


def test_r3_update_fallback_keeps_xlsm_macro_enabled(xlsm: Path, tmp_path: Path) -> None:
    df = pd.DataFrame([{"level": 1, "material": "M-1", "quantity": 3}])
    with patch.object(ie, "excel_com_update_sheet", return_value=False), \
         patch.object(InheritanceEngine, "refresh_workbook_pivots", return_value=False):
        InheritanceEngine().update_r3_in_workbook(xlsm, df, archive_dir=tmp_path / "old")
    assert _workbook_content_type(xlsm) == "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
    wb = openpyxl.load_workbook(xlsm, keep_vba=True)
    assert wb["R3"]["B2"].value == "M-1"
    assert wb["R3"]["T2"].value == '=IF(B2="","",B2)'


def test_update_dialog_palette_has_readable_text_in_both_themes() -> None:
    from src.gui.styles.tokens import calculate_contrast_ratio
    from src.gui.update_dialog import UpdateDialog

    for dark in (True, False):
        c = UpdateDialog._palette(dark)
        assert calculate_contrast_ratio(c["text"], c["card_bg"]) >= 4.5
        assert calculate_contrast_ratio(c["text"], c["edit_bg"]) >= 4.5
        assert calculate_contrast_ratio(c["muted"], c["card_bg"]) >= 4.5
