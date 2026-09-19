"""Unit test suite for Member Workspace View (Requirement R6 & formnguoidung).

Covers:
1. Auto-loading engineer assignment (Engineer Name, Machine Code, Department) without manual text boxes.
2. 3 Input Tables:
   - CTTT (15 columns, cols A:N and Q, backward compatibility with cols 1 & 7, clipboard paste).
   - MSI (35 canonical fixed unit rows, 12 columns, quick-edit synchronization, contrast styling).
   - Label 7980/7990 (12 columns, add/remove/clear operations).
3. Preliminary self-check against local BOMs (auto-discovery of PLM_*.xlsx and R3_*.xls).
4. Submission seal stamping CTTT!Q2 = "OK" (green fill) and unlock/unsubmit functionality.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch
import openpyxl
import pandas as pd
import pytest
from PyQt6.QtWidgets import QApplication, QMessageBox

from src.gui.member_view import CANONICAL_MSI_UNITS, MemberWorkspaceView
from src.reporting.excel_generator import COLOR_GREEN_FILL_HEX, COLOR_RED_FILL_HEX


@pytest.fixture
def qapp() -> QApplication:
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


class TestMemberViewR6:
    """Requirement R6 Verification Test Suite for MemberWorkspaceView."""

    def test_assignment_auto_loading_and_department_resolution(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify automatic extraction of engineer name, machine code, and department from path."""
        machine_dir = tmp_path / "110C103NL1"
        machine_dir.mkdir(parents=True, exist_ok=True)
        assignment_file = machine_dir / "BuiVanHai_mecha1.xlsm"

        # Create a valid formnguoidung template workbook
        wb = openpyxl.Workbook()
        ws_cttt = wb.active
        ws_cttt.title = "CTTT"
        ws_cttt.append(["Title", None, None])
        ws_cttt.append(["Unit", "Trang", "Mã linh kiện", "Tên linh kiện", "Số lượng", "Phụ trách"])
        ws_cttt.append(["LSU", "01", "302FP02010", "MOTOR BRACKET", 2.0, "BuiVanHai_mecha1"])
        wb.save(assignment_file)
        wb.close()

        view = MemberWorkspaceView(base_dir=tmp_path)
        ok = view.open_assignment_package(assignment_file)

        assert ok is True
        assert view.engineer_name == "BuiVanHai_mecha1"
        assert view.machine_code == "110C103NL1"
        assert view.department == "Phòng Cơ 1"
        assert view.author_edit.text() == "BuiVanHai_mecha1"
        assert view.lbl_engineer_name.text() == "BuiVanHai_mecha1"
        assert view.lbl_machine_code.text() == "110C103NL1"
        assert view.lbl_department.text() == "Phòng Cơ 1"
        assert view.lbl_assignment_path.text() == str(assignment_file)
        assert view.is_submitted_ok is False
        assert "CHƯA NỘP" in view.lbl_submission_seal.text()

        # Check that CTTT row was auto-loaded
        assert view.cttt_table.rowCount() == 1
        assert view.cttt_table.item(0, 1).text() == "302FP02010"
        assert view.cttt_table.item(0, 2).text() == "MOTOR BRACKET"

        # Test alternative department suffixes
        file_mecha2 = machine_dir / "TranVanB_mecha2.xlsm"
        file_mecha2.write_text("dummy")
        view.open_assignment_package(file_mecha2)
        assert view.department == "Phòng Cơ 2"

        file_mecha3 = machine_dir / "LeVanC_mecha3.xlsm"
        file_mecha3.write_text("dummy")
        view.open_assignment_package(file_mecha3)
        assert view.department == "Phòng Cơ 3"

        file_pthtct = machine_dir / "PhamVanD_PTHTCT.xlsm"
        file_pthtct.write_text("dummy")
        view.open_assignment_package(file_pthtct)
        assert view.department == "Phụ Trách Hệ Thống Cấu Trúc"

    def test_auto_discover_machine_boms(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify automatic discovery and parsing of PLM_*.xlsx and R3_*.xls in machine folder."""
        machine_dir = tmp_path / "110C103NL1"
        machine_dir.mkdir(parents=True, exist_ok=True)

        # Create PLM Excel (TC2412 style: Name, Parts Text, Release Status)
        plm_path = machine_dir / "PLM_110C103NL1.xlsx"
        df_plm = pd.DataFrame([
            {"Name": "302FP02010", "Parts Text": "MOTOR BRACKET", "Quantity": 2.0, "Release Status": "A"},
            {"Name": "302FP02020", "Parts Text": "SCREW M3", "Quantity": 4.0, "Release Status": "A"},
        ])
        df_plm.to_excel(plm_path, index=False)

        # Create R3 Excel
        r3_path = machine_dir / "R3_110C103NL1.xls"
        df_r3 = pd.DataFrame([
            {"material": "302FP02010", "quantity": 2.0, "revlev": "A"},
            {"material": "302FP02020", "quantity": 4.0, "revlev": "A"},
        ])
        df_r3.to_excel(r3_path, index=False)

        view = MemberWorkspaceView(base_dir=tmp_path)
        view.load_machine_directory(machine_dir)

        assert view.plm_data is not None
        assert not view.plm_data.empty
        assert view.r3_data is not None
        assert not view.r3_data.empty
        assert "PLM_110C103NL1.xlsx" in view.lbl_ref_status.text()
        assert "R3_110C103NL1.xls" in view.lbl_ref_status.text()

    def test_canonical_35_msi_units_initialization_and_quick_edit(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify MSI table has 35 fixed units matching formnguoidung and syncs with quick-edit."""
        view = MemberWorkspaceView(base_dir=tmp_path)

        assert view.msi_table.rowCount() == 35
        assert view.msi_table.columnCount() == 12

        # Check first and last canonical units
        assert view.msi_table.item(0, 2).text() == "IMAGE UNIT"
        assert view.msi_table.item(1, 2).text() == "FUSER UNIT"
        assert view.msi_table.item(34, 2).text() == "HONTAI(ký tự cố định hontai)"

        # Check row selection updates quick-edit fields
        view.msi_table.item(1, 0).setText("302FP93020")
        view.msi_table.item(1, 3).setText("2NL")
        view.msi_table.item(1, 4).setText("SERVICE NOTE")
        view._on_msi_table_row_selected(1, 0)

        assert view.msi_barcode_edit.text() == "302FP93020"
        assert view.msi_unit_name_edit.text() == "FUSER UNIT"
        assert view.msi_code_edit.text() == "2NL"
        assert view.msi_service_edit.text() == "SERVICE NOTE"

        # Check quick-edit application back to table
        view.msi_code_edit.setText("3MM")
        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            view._apply_msi_quick_edit()
        assert view.msi_table.item(1, 3).text() == "3MM"

    def test_label_7980_7990_operations(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify Label 7980/7990 table manipulation (12 columns)."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_label_table()
        assert view.label_table.rowCount() == 0
        assert view.label_table.columnCount() == 12

        row_idx = view.add_label_row(
            sub_7980="LSU",
            page_7980="01",
            code_7980="7980-VN-001",
            name_7980="LCP LABEL",
            qty_7980=1,
            author_7980="Hai_mecha1",
            sub_7990="HIGH VOLTAGE",
            page_7990="02",
            code_7990="7990-VN-002",
            name_7990="CAUTION LABEL",
            qty_7990=2,
            author_7990="Hai_mecha1",
        )
        assert row_idx == 0
        assert view.label_table.rowCount() == 1
        assert view.label_table.item(0, 2).text() == "7980-VN-001"
        assert view.label_table.item(0, 8).text() == "7990-VN-002"

        # Remove row
        view.label_table.setCurrentCell(0, 0)
        view.remove_selected_label_row()
        assert view.label_table.rowCount() == 0

    def test_cttt_clipboard_paste(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify pasting tabular data from system clipboard into CTTT table."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()

        tsv_data = "01\t302FP02010\tMOTOR BRACKET\t2.0\tOK\n02\t302FP02020\tSCREW M3\t4.0\tOK\n"
        clipboard = QApplication.clipboard()
        clipboard.setText(tsv_data)

        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            view.paste_cttt_from_clipboard()

        assert view.cttt_table.rowCount() == 2
        assert view.cttt_table.item(0, 1).text() == "302FP02010"
        assert view.cttt_table.item(0, 3).text() == "2.0"
        assert view.cttt_table.item(1, 1).text() == "302FP02020"
        assert view.cttt_table.item(1, 3).text() == "4.0"

    def test_preliminary_self_check_with_tc2412_and_msi_evaluation(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify preliminary self-check against TC2412 BOM and MSI evaluation with soft green/red colors."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()

        # Supply reference data
        plm_df = pd.DataFrame([
            {"Name": "302FP02010", "Quantity": 2.0, "Release Status": "A"},
            {"Name": "302FP02020", "Quantity": 4.0, "Release Status": "A"},  # Mismatch: CTTT will have 6.0
        ])
        r3_df = pd.DataFrame([
            {"material": "302FP02010", "quantity": 2.0, "revlev": "A"},
            {"material": "302FP02020", "quantity": 4.0, "revlev": "A"},
        ])
        view.set_reference_data(plm_data=plm_df, r3_data=r3_df)

        view.add_cttt_row(page="01", part_code="302FP02010", quantity=2.0)
        view.add_cttt_row(page="02", part_code="302FP02020", quantity=6.0)

        # Set MSI row
        view.msi_table.item(0, 1).setText("302FP93010")
        view.msi_table.item(0, 3).setText("1HN")
        view.fix_serial_master.add_subunit("302FP93010", "1HN", "-")

        res = view.run_preliminary_self_check()
        assert res["ok_count"] == 1
        assert res["ng_count"] == 1
        assert res["status"] == "NG"

        # Check visual color coding in CTTT table
        item_ok = view.cttt_table.item(0, 7)
        assert item_ok.text() == "OK"
        assert item_ok.background().color().name().upper().endswith(COLOR_GREEN_FILL_HEX)

        item_ng = view.cttt_table.item(1, 7)
        assert item_ng.text() == "NG"
        assert item_ng.background().color().name().upper().endswith(COLOR_RED_FILL_HEX)

        # Check MSI table evaluation
        msi_status = view.msi_table.item(0, 10)
        assert msi_status.text() == "OK"
        assert msi_status.background().color().name().upper().endswith(COLOR_GREEN_FILL_HEX)

    def test_submission_seal_stamping_q2_ok_and_unlock(
        self,
        qapp: QApplication,
        tmp_path: Path,
    ) -> None:
        """Verify 'Xác nhận Nộp' stamps CTTT!Q2='OK' in green and 'Hủy nộp' clears it."""
        machine_dir = tmp_path / "110C103NL1"
        machine_dir.mkdir(parents=True, exist_ok=True)
        assignment_file = machine_dir / "BuiVanHai_mecha1.xlsm"

        # Create blank assignment file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CTTT"
        ws.append(["Title"])
        ws.append(["Header"])
        wb.save(assignment_file)
        wb.close()

        view = MemberWorkspaceView(base_dir=tmp_path)
        view.open_assignment_package(assignment_file)
        view.clear_cttt_table()

        view.add_cttt_row(page="01", part_code="302FP02010", part_name="MOTOR", quantity=2.0, status="OK")

        signal_payload = []
        view.submission_completed.connect(lambda info: signal_payload.append(info))

        # 1. Submit
        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            res_file = view.submit_data()

        assert res_file is not None
        assert res_file == assignment_file
        assert view.is_submitted_ok is True
        assert "ĐÃ NỘP BÀI (Q2 = OK)" in view.lbl_submission_seal.text()
        assert len(signal_payload) == 1
        assert signal_payload[0]["q2_status"] == "OK"

        # Check physical workbook
        wb_check = openpyxl.load_workbook(assignment_file)
        ws_check = wb_check["CTTT"]
        assert ws_check["Q2"].value == "OK"
        assert ws_check["Q2"].fill.start_color.rgb.endswith(COLOR_GREEN_FILL_HEX)
        wb_check.close()

        # 2. Unlock / Unsubmit
        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            view.unlock_submission()

        assert view.is_submitted_ok is False
        assert "CHƯA NỘP" in view.lbl_submission_seal.text()

        # Check physical workbook after unlock
        wb_unlocked = openpyxl.load_workbook(assignment_file)
        ws_unlocked = wb_unlocked["CTTT"]
        assert ws_unlocked["Q2"].value is None
        wb_unlocked.close()
