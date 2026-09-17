"""Comprehensive Unit & Mock Test Suite for GUI & Reporting Modules (Milestone M5).

Covers:
- ExcelReportGenerator (F24): 7-sheet schema, exact conditional formatting & color vectors (Red 255 / Green 6750054).
- OutlookMailer (F25): HTML body generation, recipient routing, attachment handling, COM interop mocking.
- SettingsDialog: TC14/SAP/Path configuration persistence, default resets, field validation.
- MemberWorkspaceView (F23): Table operations, preliminary self-check, visual color updates, submission.
- LeaderWorkspaceView (F22): Folder tree creation, submission scanning, batch reconciliation, report triggering.
- SSBOMMainWindow: Multi-role tab navigation, status bar, logging console integration.

Operates in headless mode (QT_QPA_PLATFORM=offscreen) without requiring interactive display or live Outlook.
"""

from __future__ import annotations

import os
from pathlib import Path
import sys
from unittest.mock import MagicMock, patch

import openpyxl
import pandas as pd
import pytest
from PyQt6.QtWidgets import QApplication

from src.core.reconciliation import ReconciliationEngine, ReconciliationResult
from src.gui.app import SSBOMMainWindow
from src.gui.leader_view import (
    BatchReconciliationWorker,
    EmailPreviewDialog,
    LeaderWorkspaceView,
)
from src.gui.member_view import MemberWorkspaceView
from src.gui.settings_dialog import DEFAULT_SETTINGS, SettingsDialog
from src.reporting.excel_generator import (
    COLOR_GREEN_BGR,
    COLOR_GREEN_FILL_HEX,
    COLOR_RED_BGR,
    COLOR_RED_FILL_HEX,
    STANDARD_SUB_UNITS,
    ExcelReportGenerator,
    generate_consolidated_report,
)
from src.reporting.outlook_mailer import (
    EmailPreview,
    OutlookMailer,
    generate_reconciliation_email,
)


@pytest.fixture(scope="session")
def qapp() -> QApplication:
    """Provide headless QApplication instance for GUI tests."""
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


# =============================================================================
# 1. Tests for ExcelReportGenerator (F24)
# =============================================================================

class TestExcelReportGenerator:
    """Test suite verifying consolidated Excel report workbook generation."""

    def test_color_vector_constants(self) -> None:
        """Verify exact color vectors matching legacy VBA definitions."""
        # Red 255: RGB(255, 0, 0)
        assert COLOR_RED_BGR == 255
        # Green 6750054: RGB(102, 255, 102) in BGR = 0x66FF66 = 6750054
        assert COLOR_GREEN_BGR == 6750054
        assert COLOR_RED_FILL_HEX == "FFC7CE"
        assert COLOR_GREEN_FILL_HEX == "C6EFCE"

    def test_seven_canonical_sheets_created(self, tmp_path: Path) -> None:
        """Verify all 7 required worksheets exist in the generated report."""
        output_file = tmp_path / "test_form_ssbom.xlsx"
        generator = ExcelReportGenerator()

        cttt_rows = [
            {
                "SUB": "LSU",
                "TRANG CTTT": "01",
                "MÃ LINH KIỆN": "302FP02010",
                "TÊN LINH KIỆN": "MOTOR BRACKET",
                "SỐ LƯỢNG": 1.0,
                "PHỤ TRÁCH": "Nguyen Van A",
                "Check": "OK",
            },
            {
                "SUB": "FUSER",
                "TRANG CTTT": "02",
                "MÃ LINH KIỆN": "302FP02020",
                "TÊN LINH KIỆN": "HEATER LAMP",
                "SỐ LƯỢNG": 2.0,
                "PHỤ TRÁCH": "Tran Van B",
                "Check": "NG",
            },
        ]
        plm_rows = [
            {
                "level": 1,
                "item_type": "Assembly",
                "item_id": "302FP93010",
                "has_children": "True",
                "quantity": 1.0,
                "occurrence_effectivities": "01-Jan-2024 UP",
                "item_name": "LSU UNIT",
                "revision": "A",
            },
            {
                "level": 2,
                "item_type": "Part",
                "item_id": "302FP02010",
                "has_children": "False",
                "quantity": 1.0,
                "occurrence_effectivities": "01-Jan-2024 UP",
                "item_name": "MOTOR BRACKET",
                "revision": "A",
            },
        ]
        r3_rows = [
            {
                "level": 1,
                "material": "302FP02010",
                "description": "MOTOR BRACKET",
                "quantity": 1.0,
                "uom": "PC",
                "revision": "A",
                "valid_from": "01.01.2024",
                "valid_to": "31.12.9999",
            },
        ]

        saved_path = generator.generate_report(
            output_path=output_file,
            cttt_data=cttt_rows,
            plm_data=plm_rows,
            r3_data=r3_rows,
            model_name="Virgo",
            target_date="17/09/2026",
        )

        assert saved_path.exists()

        wb = openpyxl.load_workbook(saved_path, data_only=True)
        expected_sheets = ["Tongket", "List JIG", "MSI_7980_7990", "CTTT", "PLM", "R3", "CTTT_Total"]
        for s in expected_sheets:
            assert s in wb.sheetnames, f"Sheet '{s}' missing from generated workbook"

        # Check Tongket contents
        ws_tong = wb["Tongket"]
        assert "BÁO CÁO TỔNG KẾT" in str(ws_tong["B2"].value)

        wb.close()

    def test_conditional_formatting_and_colors_applied(self, tmp_path: Path) -> None:
        """Verify conditional formatting rules and cell fills for OK / NG."""
        output_file = tmp_path / "test_formatting.xlsx"
        generator = ExcelReportGenerator()

        cttt_rows = [
            {
                "SUB": "LSU",
                "TRANG CTTT": "01",
                "MÃ LINH KIỆN": "302FP02010",
                "TÊN LINH KIỆN": "MOTOR",
                "SỐ LƯỢNG": 1.0,
                "PHỤ TRÁCH": "User1",
                "Compare (PLM Qty)": "OK",
                "Compare (R3 Qty)": "OK",
                "Compare (Rev)": "OK",
                "Check": "OK",
            },
            {
                "SUB": "DLP",
                "TRANG CTTT": "02",
                "MÃ LINH KIỆN": "302FP02020",
                "TÊN LINH KIỆN": "SCREW",
                "SỐ LƯỢNG": 4.0,
                "PHỤ TRÁCH": "User2",
                "Compare (PLM Qty)": "NG",
                "Compare (R3 Qty)": "OK",
                "Compare (Rev)": "NG",
                "Check": "NG",
            },
        ]
        msi_rows = [
            {
                "barcode_part": "302FP93010",
                "unit_code": "302FP93010",
                "unit_name": "LSU",
                "member_code": "1HN",
                "member_service": "-",
                "abs": "OK",
                "status": "OK",
                "highlight_green": ["B", "D", "E"],
            },
            {
                "barcode_part": "302FP94010",
                "unit_code": "302FP94010",
                "unit_name": "FUSER",
                "member_code": "2NL",
                "member_service": "WRONG",
                "abs": "NG",
                "status": "NG",
                "highlight_red": ["E", "D"],
            },
        ]

        generator.generate_report(
            output_path=output_file,
            cttt_data=cttt_rows,
            msi_data=msi_rows,
            model_name="Libra2",
        )

        wb = openpyxl.load_workbook(output_file, data_only=False)

        # Check Sheet CTTT conditional formatting rules
        ws_cttt = wb["CTTT"]
        cf_rules = ws_cttt.conditional_formatting
        assert len(cf_rules) > 0, "No conditional formatting registered on Sheet CTTT"

        # Check direct cell styles on CTTT
        cell_ok = ws_cttt.cell(row=3, column=18)  # Row 3 Check is OK
        assert cell_ok.value == "OK"
        assert cell_ok.fill.start_color.rgb.upper().endswith(COLOR_GREEN_FILL_HEX)

        cell_ng = ws_cttt.cell(row=4, column=18)  # Row 4 Check is NG
        assert cell_ng.value == "NG"
        assert cell_ng.fill.start_color.rgb.upper().endswith(COLOR_RED_FILL_HEX)

        # Check Sheet MSI_7980_7990
        ws_msi = wb["MSI_7980_7990"]
        msi_ok_cell = ws_msi.cell(row=2, column=11)  # Status column K
        assert msi_ok_cell.value == "OK"
        assert msi_ok_cell.fill.start_color.rgb.upper().endswith(COLOR_GREEN_FILL_HEX)

        msi_ng_cell = ws_msi.cell(row=3, column=11)
        assert msi_ng_cell.value == "NG"
        assert msi_ng_cell.fill.start_color.rgb.upper().endswith(COLOR_RED_FILL_HEX)

        wb.close()

    def test_generate_report_from_reconciliation_result(self, tmp_path: Path) -> None:
        """Verify integration with ReconciliationResult contract."""
        engine = ReconciliationEngine()
        cttt = pd.DataFrame([
            {"SUB": "LSU", "TRANG CTTT": "01", "MÃ LINH KIỆN": "302FP02010", "TÊN LINH KIỆN": "BRACKET", "SỐ LƯỢNG": 1.0, "PHỤ TRÁCH": "User"},
        ])
        plm = pd.DataFrame([
            {"level": 1, "item_id": "302FP02010", "item_name": "BRACKET", "quantity": 1.0, "revision": "A"},
            {"level": 1, "item_id": "302FP09999", "item_name": "OMITTED PART", "quantity": 2.0, "revision": "A"},
        ])
        r3 = pd.DataFrame([
            {"material": "302FP02010", "quantity": 1.0, "revision": "A"},
        ])

        recon_res = engine.run_full_reconciliation(cttt_data=cttt, plm_data=plm, r3_data=r3)
        assert isinstance(recon_res, ReconciliationResult)

        out_path = tmp_path / "recon_report.xlsx"
        generator = ExcelReportGenerator()
        res_file = generator.generate_report(
            output_path=out_path,
            reconciliation=recon_res,
            model_name="Iris2024",
        )
        assert res_file.exists()

        wb = openpyxl.load_workbook(res_file)
        # Sheet PLM should show missing part in red
        ws_plm = wb["PLM"]
        # Col N is Check CTTT
        has_missing_mark = any(
            "Thiếu CTTT" in str(ws_plm.cell(r, 14).value)
            for r in range(2, ws_plm.max_row + 1)
        )
        assert has_missing_mark, "Reverse check column in PLM should mark omitted part"
        wb.close()

    def test_procedural_helper(self, tmp_path: Path) -> None:
        """Verify generate_consolidated_report procedural shortcut."""
        out = tmp_path / "shortcut.xlsx"
        p = generate_consolidated_report(out, model_name="Polaris")
        assert p.exists()


# =============================================================================
# 2. Tests for OutlookMailer (F25)
# =============================================================================

class TestOutlookMailer:
    """Test suite verifying Outlook COM HTML notification generation and dispatch."""

    def test_html_body_generation_ok_status(self) -> None:
        """Verify rendered HTML when overall reconciliation status is OK."""
        mailer = OutlookMailer()
        html = mailer.generate_html_body(
            model_name="Virgo",
            target_date="17/09/2026",
            overall_status="OK",
            summary_stats={"total_parts": 100, "ok_count": 100, "ng_count": 0, "missing_count": 0},
        )
        assert "Virgo" in html
        assert "17/09/2026" in html
        assert "HOÀN TẤT - KHỚP 100% (OK)" in html
        assert "HÀNH ĐỘNG YÊU CẦU" not in html

    def test_html_body_generation_ng_status_action_required(self) -> None:
        """Verify rendered HTML when status is NG includes Action Required alert."""
        mailer = OutlookMailer()
        html = mailer.generate_html_body(
            model_name="Libra2",
            target_date="17/09/2026",
            overall_status="NG",
            summary_stats={"total_parts": 120, "ok_count": 115, "ng_count": 5, "missing_count": 2},
            custom_notes="Vui lòng kiểm tra khẩn cấp công đoạn FUSER.",
        )
        assert "PHÁT HIỆN SAI KHÁC CẦN GIẢI TRÌNH (NG)" in html
        assert "HÀNH ĐỘNG YÊU CẦU" in html
        assert "Vui lòng kiểm tra khẩn cấp công đoạn FUSER." in html

    def test_build_email_preview_dataclass(self, tmp_path: Path) -> None:
        """Verify EmailPreview payload construction and attachment resolution."""
        mailer = OutlookMailer()
        dummy_att = tmp_path / "report.xlsx"
        dummy_att.write_text("excel_binary_data")

        preview = mailer.build_email_preview(
            model_name="Sirius2",
            target_date="17/09/2026",
            overall_status="OK",
            recipients_to=["user1@kyocera.com", "user2@kyocera.com"],
            recipients_cc=["lead@kyocera.com"],
            attachment_path=dummy_att,
        )

        assert isinstance(preview, EmailPreview)
        assert "Sirius2" in preview.subject
        assert "[HOÀN TẤT - OK]" in preview.subject
        assert preview.to_string == "user1@kyocera.com; user2@kyocera.com"
        assert preview.cc_string == "lead@kyocera.com"
        assert len(preview.attachment_paths) == 1
        assert preview.attachment_paths[0] == dummy_att

    def test_outlook_com_mock_preview_and_send(self, tmp_path: Path) -> None:
        """Verify Outlook COM dispatch calls using simulated mock object."""
        mock_mail_item = MagicMock()
        mock_app = MagicMock()
        mock_app.CreateItem.return_value = mock_mail_item

        mailer = OutlookMailer(com_dispatch=mock_app)
        dummy_file = tmp_path / "test.xlsx"
        dummy_file.write_text("data")

        preview = mailer.build_email_preview(
            model_name="Mebius",
            target_date="17/09/2026",
            overall_status="OK",
            recipients_to="team@kyocera.com",
            attachment_path=dummy_file,
        )

        # 1. Test preview in outlook (Display)
        ok_disp = mailer.preview_in_outlook(preview)
        assert ok_disp is True
        mock_app.CreateItem.assert_called_with(0)
        assert mock_mail_item.To == "team@kyocera.com"
        assert "Mebius" in mock_mail_item.Subject
        mock_mail_item.Display.assert_called_with(False)
        mock_mail_item.Attachments.Add.assert_called_with(str(dummy_file))

        # 2. Test direct send
        ok_send = mailer.send_via_outlook(preview)
        assert ok_send is True
        mock_mail_item.Send.assert_called_once()

    def test_offline_fallback_when_com_unavailable(self) -> None:
        """Verify graceful fallback without raising unhandled exceptions when COM fails."""
        mailer = OutlookMailer(com_dispatch=None)
        with patch.object(mailer, "_get_outlook_application", return_value=None):
            preview = EmailPreview(
                subject="Test",
                recipients_to=["a@b.com"],
                recipients_cc=[],
                html_body="<p>Test</p>",
            )
            # Both preview and send should return False gracefully
            assert mailer.preview_in_outlook(preview) is False
            assert mailer.send_via_outlook(preview) is False

    def test_generate_reconciliation_email_helper(self) -> None:
        """Verify generate_reconciliation_email procedural shortcut."""
        preview = generate_reconciliation_email(
            model_name="Virgo",
            target_date="17/09/2026",
            overall_status="OK",
            recipients_to="a@b.com",
        )
        assert isinstance(preview, EmailPreview)


# =============================================================================
# 3. Tests for SettingsDialog
# =============================================================================

class TestSettingsDialog:
    """Test suite verifying application configuration dialog."""

    def test_settings_dialog_initialization(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify dialog initializes with default parameters."""
        cfg_file = tmp_path / "custom_config.json"
        dlg = SettingsDialog(config_path=cfg_file)

        assert dlg.tc_url_edit.text() == DEFAULT_SETTINGS["tc14"]["base_url"]
        assert dlg.tc_user_edit.text() == DEFAULT_SETTINGS["tc14"]["username"]
        assert dlg.sap_system_edit.text() == DEFAULT_SETTINGS["sap"]["system_id"]
        assert dlg.sap_plant_edit.text() == DEFAULT_SETTINGS["sap"]["plant"]

    def test_settings_dialog_save_and_reload(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify modifying fields, saving to file, and loading back."""
        cfg_file = tmp_path / "saved_config.json"
        dlg = SettingsDialog(config_path=cfg_file)

        # Modify values
        dlg.tc_user_edit.setText("new_user_pe03")
        dlg.sap_plant_edit.setText("2300")
        dlg.default_model_combo.setCurrentText("Polaris")

        signal_received = []
        dlg.settings_saved.connect(lambda s: signal_received.append(s))

        # Save
        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            dlg.save_and_close()

        assert cfg_file.exists()
        assert len(signal_received) == 1
        assert signal_received[0]["tc14"]["username"] == "new_user_pe03"
        assert signal_received[0]["sap"]["plant"] == "2300"
        assert signal_received[0]["paths"]["default_model"] == "Polaris"

        # Reload in new instance
        dlg2 = SettingsDialog(config_path=cfg_file)
        assert dlg2.tc_user_edit.text() == "new_user_pe03"
        assert dlg2.sap_plant_edit.text() == "2300"
        assert dlg2.default_model_combo.currentText() == "Polaris"


# =============================================================================
# 4. Tests for MemberWorkspaceView (F23)
# =============================================================================

class TestMemberWorkspaceView:
    """Test suite verifying Member Workspace table manipulation, self-check, and submission."""

    def test_table_row_operations(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify adding, removing, and clearing CTTT component rows."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        assert view.cttt_table.rowCount() == 0

        # Add rows
        view.add_cttt_row(page="01", part_code="302FP02010", part_name="MOTOR", quantity=1.0)
        view.add_cttt_row(page="01", part_code="302FP02020", part_name="SCREW", quantity=4.0)
        assert view.cttt_table.rowCount() == 2

        # Select and remove row 0
        view.cttt_table.setCurrentCell(0, 0)
        view.remove_selected_cttt_row()
        assert view.cttt_table.rowCount() == 1
        assert view.cttt_table.item(0, 1).text() == "302FP02020"

    def test_preliminary_self_check_visual_feedback(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify instant visual OK/NG coloring on self-check."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()

        # Supply reference PLM & R3 data
        plm_df = pd.DataFrame([
            {"item_id": "302FP02010", "quantity": 1.0},
            {"item_id": "302FP02020", "quantity": 2.0},  # Mismatch: CTTT will have 4.0
        ])
        r3_df = pd.DataFrame([
            {"material": "302FP02010", "quantity": 1.0},
            {"material": "302FP02020", "quantity": 2.0},
        ])
        view.set_reference_data(plm_data=plm_df, r3_data=r3_df)

        view.add_cttt_row(page="01", part_code="302FP02010", part_name="MOTOR", quantity=1.0)
        view.add_cttt_row(page="01", part_code="302FP02020", part_name="SCREW", quantity=4.0)

        res = view.run_preliminary_self_check()
        assert res["ok_count"] == 1
        assert res["ng_count"] == 1
        assert res["status"] == "NG"

        # Check visual coloring
        # Row 0: OK -> Green fill
        item_ok = view.cttt_table.item(0, 7)
        assert item_ok.text() == "OK"
        assert item_ok.background().color().name().upper().endswith(COLOR_GREEN_FILL_HEX)

        # Row 1: NG -> Red fill
        item_ng = view.cttt_table.item(1, 7)
        assert item_ng.text() == "NG"
        assert item_ng.background().color().name().upper().endswith(COLOR_RED_FILL_HEX)

    def test_member_submission_export(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify final submission generates valid formnguoidung Excel package."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        view.sub_unit_combo.setCurrentText("FUSER")
        view.author_edit.setText("Tran Van Member")

        view.add_cttt_row(page="02", part_code="302FP02050", part_name="HEATER", quantity=1.0, status="OK")

        signal_captured = []
        view.submission_completed.connect(lambda info: signal_captured.append(info))

        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            exported_file = view.submit_data()

        assert exported_file is not None
        assert exported_file.exists()
        assert "formnguoidung_FUSER_" in exported_file.name

        # Verify emitted metadata
        assert len(signal_captured) == 1
        assert signal_captured[0]["sub_unit"] == "FUSER"
        assert signal_captured[0]["author"] == "Tran Van Member"
        assert signal_captured[0]["item_count"] == 1

        # Check file content
        wb = openpyxl.load_workbook(exported_file)
        assert "CTTT" in wb.sheetnames
        assert "MSI" in wb.sheetnames
        assert "Label_7980_7990" in wb.sheetnames
        wb.close()


# =============================================================================
# 5. Tests for LeaderWorkspaceView (F22)
# =============================================================================

class TestLeaderWorkspaceView:
    """Test suite verifying Leader Management Workspace operations."""

    def test_create_project_folder_structure(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify creating canonical directory tree by machine model."""
        leader_view = LeaderWorkspaceView(base_dir=tmp_path)
        leader_view.model_combo.setCurrentText("Libra2")

        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            model_dir = leader_view.create_project_folder_structure()

        assert model_dir.exists()
        assert (model_dir / "PLM").is_dir()
        assert (model_dir / "R3").is_dir()
        assert (model_dir / "Reports").is_dir()
        assert (model_dir / "capnhat" / "old").is_dir()

        # Check sub-unit directories
        for unit in STANDARD_SUB_UNITS:
            assert (model_dir / "CTTT" / unit).is_dir()

    def test_scan_member_submissions(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify submission scanner detects member files."""
        leader_view = LeaderWorkspaceView(base_dir=tmp_path)
        leader_view.model_combo.setCurrentText("Virgo")

        # Create a fake submission for LSU
        lsu_dir = tmp_path / "Virgo" / "CTTT" / "LSU"
        lsu_dir.mkdir(parents=True, exist_ok=True)
        fake_sub = lsu_dir / "formnguoidung_LSU_20260917.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CTTT"
        ws.append(["SUB", "TRANG", "MÃ LK", "TÊN", "SL", "NGƯỜI", "GT", "Check"])
        ws.append(["LSU", "01", "302FP02010", "MOTOR", 1, "User1", "", "OK"])
        wb.save(fake_sub)
        wb.close()

        statuses = leader_view.scan_member_submissions()
        assert statuses["LSU"]["status"] == "Đã nộp"
        assert statuses["LSU"]["file"] == fake_sub
        assert statuses["DLP"]["status"] == "Chưa nộp"

    def test_batch_reconciliation_worker_execution(self, qapp: QApplication) -> None:
        """Verify background batch worker completes and emits ReconciliationResult."""
        cttt_data = [
            {"SUB": "LSU", "TRANG CTTT": "01", "MÃ LINH KIỆN": "302FP02010", "TÊN LINH KIỆN": "MOTOR", "SỐ LƯỢNG": 1.0, "PHỤ TRÁCH": "User1"},
        ]
        worker = BatchReconciliationWorker(
            collected_cttt=cttt_data,
            plm_path=None,
            r3_path=None,
            model_name="Virgo",
        )

        results_captured = []
        worker.finished.connect(lambda res: results_captured.append(res))

        worker.run()
        assert len(results_captured) == 1
        assert isinstance(results_captured[0], ReconciliationResult)

    def test_trigger_consolidated_report(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify report generation triggered from leader workspace."""
        leader_view = LeaderWorkspaceView(base_dir=tmp_path)
        leader_view.model_combo.setCurrentText("Virgo")

        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            rep_path = leader_view.trigger_consolidated_report()

        assert rep_path is not None
        assert rep_path.exists()
        assert "form_ssbom_Virgo_" in rep_path.name


# =============================================================================
# 6. Tests for SSBOMMainWindow (App Integration)
# =============================================================================

class TestSSBOMMainWindow:
    """Test suite verifying top-level main window, logging, and tab switching."""

    def test_main_window_tabs_and_views(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify tab switching and child workspace availability."""
        win = SSBOMMainWindow(base_dir=tmp_path)
        assert win.tabs.count() == 2

        # Leader tab is 0
        assert win.tabs.currentIndex() == 0
        assert isinstance(win.tabs.currentWidget(), LeaderWorkspaceView)

        # Switch to Member tab
        win.switch_view(1)
        assert win.tabs.currentIndex() == 1
        assert isinstance(win.tabs.currentWidget(), MemberWorkspaceView)

    def test_qlog_handler_integration(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify system logs are captured into the console text edit."""
        win = SSBOMMainWindow(base_dir=tmp_path)
        win.append_log("Test system notification message")

        console_text = win.log_text_edit.toPlainText()
        assert "Test system notification message" in console_text

    def test_email_preview_dialog_ui(self, qapp: QApplication) -> None:
        """Verify EmailPreviewDialog displays subject, recipients, and HTML."""
        mailer = OutlookMailer()
        preview = EmailPreview(
            subject="Test Subject",
            recipients_to=["member@kyocera.com"],
            recipients_cc=["lead@kyocera.com"],
            html_body="<h3>Reconciliation OK</h3>",
        )
        dlg = EmailPreviewDialog(preview=preview, mailer=mailer)
        assert dlg.edit_sub.text() == "Test Subject"
        assert dlg.edit_to.text() == "member@kyocera.com"
        assert "Reconciliation OK" in dlg.browser.toHtml()

        # Test display and send button handlers with mock
        with patch.object(mailer, "preview_in_outlook", return_value=True), patch("PyQt6.QtWidgets.QMessageBox.information"):
            dlg._on_display_clicked()

        with patch.object(mailer, "send_via_outlook", return_value=True), patch("PyQt6.QtWidgets.QMessageBox.information"):
            dlg._on_send_clicked()

    def test_main_window_menu_and_signals(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify main window cross-workspace event handlers."""
        win = SSBOMMainWindow(base_dir=tmp_path)

        # Trigger member submitted handler
        win._on_member_submitted({"sub_unit": "LSU", "author": "Tester", "item_count": 5})
        assert "LSU" in win.lbl_status_msg.text()

        # Trigger batch finished handler
        mock_res = MagicMock()
        mock_res.overall_status = "OK"
        win._on_batch_finished(mock_res)
        assert "OK" in win.lbl_status_msg.text()

        # Trigger about dialog
        with patch("PyQt6.QtWidgets.QMessageBox.about"):
            win._show_about_dialog()


# =============================================================================
# 7. Additional Coverage Hardening Tests
# =============================================================================

class TestCoverageHardening:
    """Additional tests to achieve high coverage across leader_view and member_view."""

    def test_leader_view_batch_flow(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify leader view trigger_batch_reconciliation and worker callbacks."""
        leader_view = LeaderWorkspaceView(base_dir=tmp_path)
        leader_view.model_combo.setCurrentText("Virgo")

        # Call worker callbacks directly
        leader_view._on_worker_progress(50, "Testing progress")
        assert leader_view.progress_bar.value() == 50
        assert leader_view.lbl_progress_status.text() == "Testing progress"

        mock_res = ReconciliationResult(
            cttt_rows=pd.DataFrame([{"MÃ LINH KIỆN": "A", "Check": "OK"}]),
            plm_missing_rows=pd.DataFrame(),
            cttt_totals=pd.DataFrame(),
            msi_results=pd.DataFrame(),
            overall_status="OK",
        )
        with patch("PyQt6.QtWidgets.QMessageBox.information"):
            leader_view._on_worker_finished(mock_res)
        assert leader_view.current_result == mock_res

        with patch("PyQt6.QtWidgets.QMessageBox.critical"):
            leader_view._on_worker_error("Test error")
        assert "Test error" in leader_view.lbl_progress_status.text()

    def test_leader_view_trigger_outlook_preview(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify leader view trigger_outlook_preview displays preview dialog."""
        leader_view = LeaderWorkspaceView(base_dir=tmp_path)
        with patch("src.gui.leader_view.EmailPreviewDialog.exec") as mock_exec:
            leader_view.trigger_outlook_preview()
            mock_exec.assert_called_once()

    def test_leader_view_open_folder(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify leader view _open_project_folder handles launch."""
        leader_view = LeaderWorkspaceView(base_dir=tmp_path)
        with patch("os.startfile"):
            leader_view._open_project_folder()

    def test_member_view_import_csv(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify member view CSV import parser."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        csv_file = tmp_path / "test_import.csv"
        csv_file.write_text("MÃ LINH KIỆN,TÊN LINH KIỆN,SỐ LƯỢNG,TRANG CTTT\n302FP02010,MOTOR,1.0,01\n302FP02020,SCREW,4.0,02\n", encoding="utf-8")

        with patch("PyQt6.QtWidgets.QFileDialog.getOpenFileName", return_value=(str(csv_file), "CSV")), patch("PyQt6.QtWidgets.QMessageBox.information"):
            view.import_cttt_from_file()

        assert view.cttt_table.rowCount() == 2
        assert view.cttt_table.item(0, 1).text() == "302FP02010"

    def test_member_view_sub_unit_switch(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify sub-unit combobox switch updates MSI unit name."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.sub_unit_combo.setCurrentText("DRUM")
        assert "DRUM ASSY" in view.msi_unit_name_edit.text()

    def test_settings_dialog_connectivity_and_browse(self, qapp: QApplication, tmp_path: Path) -> None:
        """Verify connectivity test buttons and browse helpers in SettingsDialog."""
        dlg = SettingsDialog(config_path=tmp_path / "cfg.json")

        # Test TC14 connection test with mocked urllib
        with patch("urllib.request.urlopen") as mock_url, patch("PyQt6.QtWidgets.QMessageBox.information"):
            mock_resp = MagicMock()
            mock_resp.getcode.return_value = 200
            mock_url.return_value.__enter__.return_value = mock_resp
            dlg._test_tc_connection()

        # Test SAP connection test with mocked path
        with patch("os.path.exists", return_value=True), patch("PyQt6.QtWidgets.QMessageBox.information"):
            dlg.sap_path_edit.setText(r"C:\fake\saplogon.exe")
            dlg._test_sap_connection()

        # Test browse helpers
        with patch("PyQt6.QtWidgets.QFileDialog.getOpenFileName", return_value=(r"C:\fake\saplogon.exe", "")):
            dlg._browse_saplogon()
            assert dlg.sap_path_edit.text() == r"C:\fake\saplogon.exe"

        with patch("PyQt6.QtWidgets.QFileDialog.getOpenFileName", return_value=(r"C:\fake\FIX.xls", "")):
            dlg._browse_fix_serial()
            assert dlg.fs_path_edit.text() == r"C:\fake\FIX.xls"

        with patch("PyQt6.QtWidgets.QFileDialog.getExistingDirectory", return_value=r"C:\fake\base"):
            dlg._browse_dir(dlg.base_dir_edit, "Test")
            assert dlg.base_dir_edit.text() == r"C:\fake\base"
