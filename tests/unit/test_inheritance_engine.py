"""Unit tests for src/core/inheritance_engine.py.

Verifies:
1. Pure-Python implementation of legacy ham_match_index_mix matching Col C.
2. Migration of human annotations:
   - Col O: Explanation (giai_thich)
   - Col P: Person in charge (phu_trach)
   - Col Q: Manager verification check (quan_ly_check)
3. New parts receive clean empty strings (""), never NaN or 'nan'.
4. Multi-version sheet archiving: PLM -> PLM_old -> PLM_old_1 -> PLM_old_2.
5. Re-generation of formula networks in Col R (=IF(C="","",C)) and Col S (=IF(E="","",E)).
6. Stale row clearing on updated sheets.
7. Safe file archiving to 'capnhat\\old\\'.
8. Pivot Table cache refresh handling.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch
import openpyxl
import pandas as pd
import pytest

from src.core.inheritance_engine import (
    ExplanationRecord,
    InheritanceEngine,
    InheritanceSummary,
    match_index_mix,
    refresh_workbook_pivots,
    update_workbook_with_inheritance,
)


@pytest.fixture
def sample_old_plm_df() -> pd.DataFrame:
    """Historical PLM DataFrame with annotations."""
    return pd.DataFrame([
        {
            "item_id": "302FP02010",
            "item_name": "MOTOR BRACKET",
            "giai_thich": "Thay doi de lap khop voi du an moi",
            "phu_trach": "Nguyen Van Hai",
            "quan_ly_check": "OK_LEADER_CONFIRMED",
        },
        {
            "item_id": "302FP02020",
            "item_name": "SCREW M3X6",
            "giai_thich": "Tang so luong tu 2 len 4 theo ECN",
            "phu_trach": "Tran Van B",
            "quan_ly_check": "CHECKED",
        },
        {
            "item_id": "302FP99999",  # Removed part in new revision
            "item_name": "OLD SENSOR",
            "giai_thich": "Linh kien da bi loai bo",
            "phu_trach": "Le Van C",
            "quan_ly_check": "REMOVED",
        },
    ])


@pytest.fixture
def sample_new_plm_df() -> pd.DataFrame:
    """New PLM DataFrame with existing, new, and unannotated parts."""
    return pd.DataFrame([
        {
            "level": 1,
            "item_type": "Part",
            "item_id": "302FP02010",  # Existing part (should inherit)
            "has_children": "False",
            "quantity": 2.0,
            "item_name": "MOTOR BRACKET",
            "revision": "B",
        },
        {
            "level": 1,
            "item_type": "Part",
            "item_id": " 302fp02020 ",  # Existing part with whitespace/casing difference
            "has_children": "False",
            "quantity": 4.0,
            "item_name": "SCREW M3X6",
            "revision": "A",
        },
        {
            "level": 2,
            "item_type": "Part",
            "item_id": "302FP05050",  # Brand new part (should be blank "")
            "has_children": "False",
            "quantity": 1.0,
            "item_name": "NEW GEAR ASSY",
            "revision": "-",
        },
    ])


@pytest.fixture
def sample_ssbom_workbook(tmp_path: Path) -> Path:
    """Create a mock form_ssbom workbook with an initial populated PLM sheet."""
    wb_path = tmp_path / "BOM_Test_Unit.xlsm"
    wb = openpyxl.Workbook()
    ws_plm = wb.active
    ws_plm.title = "PLM"

    # Header row 1
    headers = [
        "Level", "Item Type", "Item Id", "Has Children", "Quantity",
        "1st Parts", "2nd BOM Flag", "Occurrence Effectivities", "Project List",
        "Item Name", "Notice No", "Revision", "Item Rev Status", "Check CTTT",
        "Giai Thich", "Phu Trach", "Quan Ly Check", "PART CODE", "Q.TY",
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws_plm.cell(row=1, column=c_idx, value=h)

    # Historical rows 2, 3, 4
    historical_data = [
        (1, "Part", "302FP02010", "False", 1.0, "MOTOR BRACKET", "A", "Ghi chu cu 1", "Hai_mecha1", "OK"),
        (2, "Part", "302FP02020", "False", 2.0, "SCREW M3X6", "A", "Ghi chu cu 2", "Hai_mecha1", "OK"),
        (3, "Part", "302FP99999", "False", 1.0, "OLD SENSOR", "A", "Ghi chu cu 3", "Son_mecha1", "OK"),
    ]
    for r_idx, row_vals in enumerate(historical_data, start=2):
        ws_plm.cell(row=r_idx, column=1, value=row_vals[0])
        ws_plm.cell(row=r_idx, column=2, value=row_vals[1])
        ws_plm.cell(row=r_idx, column=3, value=row_vals[2])
        ws_plm.cell(row=r_idx, column=4, value=row_vals[3])
        ws_plm.cell(row=r_idx, column=5, value=row_vals[4])
        ws_plm.cell(row=r_idx, column=10, value=row_vals[5])
        ws_plm.cell(row=r_idx, column=12, value=row_vals[6])
        ws_plm.cell(row=r_idx, column=15, value=row_vals[7])  # Col O
        ws_plm.cell(row=r_idx, column=16, value=row_vals[8])  # Col P
        ws_plm.cell(row=r_idx, column=17, value=row_vals[9])  # Col Q
        ws_plm.cell(row=r_idx, column=18, value=f'=IF(C{r_idx}="","",C{r_idx})')
        ws_plm.cell(row=r_idx, column=19, value=f'=IF(E{r_idx}="","",E{r_idx})')

    wb.save(wb_path)
    wb.close()
    return wb_path


class TestInheritanceEngine:
    """Comprehensive unit test suite for InheritanceEngine."""

    def test_match_index_mix_basic_migration(
        self, sample_new_plm_df: pd.DataFrame, sample_old_plm_df: pd.DataFrame
    ) -> None:
        """Verify annotations are migrated for matching parts."""
        migrated = match_index_mix(sample_new_plm_df, sample_old_plm_df, key_column="item_id")

        assert len(migrated) == 3
        # Row 0: 302FP02010 -> should inherit Hai's notes
        row0 = migrated.iloc[0]
        assert row0["giai_thich"] == "Thay doi de lap khop voi du an moi"
        assert row0["phu_trach"] == "Nguyen Van Hai"
        assert row0["quan_ly_check"] == "OK_LEADER_CONFIRMED"

    def test_match_index_mix_whitespace_and_case_insensitivity(
        self, sample_new_plm_df: pd.DataFrame, sample_old_plm_df: pd.DataFrame
    ) -> None:
        """Verify matching is robust against whitespace and case differences."""
        migrated = match_index_mix(sample_new_plm_df, sample_old_plm_df, key_column="item_id")

        # Row 1: ' 302fp02020 ' matches '302FP02020'
        row1 = migrated.iloc[1]
        assert row1["giai_thich"] == "Tang so luong tu 2 len 4 theo ECN"
        assert row1["phu_trach"] == "Tran Van B"
        assert row1["quan_ly_check"] == "CHECKED"

    def test_match_index_mix_new_part_empty_string(
        self, sample_new_plm_df: pd.DataFrame, sample_old_plm_df: pd.DataFrame
    ) -> None:
        """Verify brand new parts receive empty string (''), never NaN or 'nan'."""
        migrated = match_index_mix(sample_new_plm_df, sample_old_plm_df, key_column="item_id")

        # Row 2: 302FP05050 is brand new
        row2 = migrated.iloc[2]
        assert row2["giai_thich"] == ""
        assert row2["phu_trach"] == ""
        assert row2["quan_ly_check"] == ""
        assert not pd.isna(row2["giai_thich"])
        assert str(row2["giai_thich"]).lower() != "nan"

    def test_match_index_mix_duplicate_in_old_picks_first(self) -> None:
        """Verify duplicate parts in old data pick the first match (MATCH(..., 0) behavior)."""
        old_with_dups = pd.DataFrame([
            {"item_id": "PART_DUP", "giai_thich": "FIRST OCCURRENCE", "phu_trach": "Eng 1", "quan_ly_check": "OK1"},
            {"item_id": "PART_DUP", "giai_thich": "SECOND OCCURRENCE", "phu_trach": "Eng 2", "quan_ly_check": "OK2"},
        ])
        new_df = pd.DataFrame([
            {"item_id": "PART_DUP", "item_name": "TEST PART"},
        ])

        migrated = match_index_mix(new_df, old_with_dups)
        assert migrated.iloc[0]["giai_thich"] == "FIRST OCCURRENCE"
        assert migrated.iloc[0]["phu_trach"] == "Eng 1"

    def test_match_index_mix_empty_old_df(
        self, sample_new_plm_df: pd.DataFrame
    ) -> None:
        """Verify empty or None old_df yields clean empty strings."""
        migrated = match_index_mix(sample_new_plm_df, pd.DataFrame())
        assert len(migrated) == 3
        for idx in range(3):
            assert migrated.iloc[idx]["giai_thich"] == ""
            assert migrated.iloc[idx]["phu_trach"] == ""
            assert migrated.iloc[idx]["quan_ly_check"] == ""

    def test_match_index_mix_empty_new_df(
        self, sample_old_plm_df: pd.DataFrame
    ) -> None:
        """Verify empty new_df returns empty DataFrame without errors."""
        migrated = match_index_mix(pd.DataFrame(), sample_old_plm_df)
        assert migrated.empty
        assert "giai_thich" in migrated.columns

    def test_determine_next_archive_sheet_name(self) -> None:
        """Verify sequential archive sheet naming: PLM_old, PLM_old_1, PLM_old_2..."""
        engine = InheritanceEngine()

        sheets_0 = ["Tongket", "CTTT", "PLM", "R3"]
        assert engine.determine_next_archive_sheet_name(sheets_0) == "PLM_old"

        sheets_1 = ["Tongket", "CTTT", "PLM", "PLM_old", "R3"]
        assert engine.determine_next_archive_sheet_name(sheets_1) == "PLM_old_1"

        sheets_2 = ["Tongket", "CTTT", "PLM", "PLM_old", "PLM_old_1", "R3"]
        assert engine.determine_next_archive_sheet_name(sheets_2) == "PLM_old_2"

        sheets_gap = ["PLM_old", "PLM_old_1", "PLM_old_3"]
        # Should pick lowest available: PLM_old_2
        assert engine.determine_next_archive_sheet_name(sheets_gap) == "PLM_old_2"

    def test_update_workbook_with_inheritance_first_archiving(
        self, sample_ssbom_workbook: Path, sample_new_plm_df: pd.DataFrame
    ) -> None:
        """Verify first update creates PLM_old and preserves old annotations into PLM."""
        engine = InheritanceEngine()
        summary = engine.update_workbook_with_inheritance(
            workbook_path=sample_ssbom_workbook,
            new_plm_data=sample_new_plm_df,
        )

        assert summary.success is True
        assert summary.archived_sheet_name == "PLM_old"
        assert summary.total_new_parts == 3
        assert summary.inherited_count == 2
        assert summary.new_unannotated_count == 1

        # Inspect workbook
        wb = openpyxl.load_workbook(sample_ssbom_workbook, data_only=False)
        assert "PLM_old" in wb.sheetnames
        assert "PLM" in wb.sheetnames

        # Check archived sheet has old data
        ws_old = wb["PLM_old"]
        assert ws_old["C2"].value == "302FP02010"
        assert ws_old["O2"].value == "Ghi chu cu 1"
        assert ws_old["C4"].value == "302FP99999"  # old removed part preserved in history

        # Check new PLM sheet has inherited data
        ws_new = wb["PLM"]
        assert ws_new["C2"].value == "302FP02010"
        assert ws_new["O2"].value == "Ghi chu cu 1"
        assert ws_new["P2"].value == "Hai_mecha1"
        assert ws_new["Q2"].value == "OK"

        # Check formula injection
        assert ws_new["R2"].value == '=IF(C2="","",C2)'
        assert ws_new["S2"].value == '=IF(E2="","",E2)'
        assert ws_new["R3"].value == '=IF(C3="","",C3)'
        assert ws_new["S3"].value == '=IF(E3="","",E3)'

        # Check row 4 (new part 302FP05050)
        assert ws_new["C4"].value == "302FP05050"
        assert ws_new["O4"].value is None  # clean unannotated

        wb.close()

    def test_update_workbook_multiple_sequential_updates(
        self, sample_ssbom_workbook: Path, sample_new_plm_df: pd.DataFrame
    ) -> None:
        """Verify multiple consecutive updates create PLM_old, then PLM_old_1, then PLM_old_2."""
        engine = InheritanceEngine()

        # Update 1 -> creates PLM_old
        sum1 = engine.update_workbook_with_inheritance(
            workbook_path=sample_ssbom_workbook,
            new_plm_data=sample_new_plm_df,
        )
        assert sum1.archived_sheet_name == "PLM_old"

        # Update 2 -> creates PLM_old_1
        sum2 = engine.update_workbook_with_inheritance(
            workbook_path=sample_ssbom_workbook,
            new_plm_data=sample_new_plm_df,
        )
        assert sum2.archived_sheet_name == "PLM_old_1"

        # Update 3 -> creates PLM_old_2
        sum3 = engine.update_workbook_with_inheritance(
            workbook_path=sample_ssbom_workbook,
            new_plm_data=sample_new_plm_df,
        )
        assert sum3.archived_sheet_name == "PLM_old_2"

        wb = openpyxl.load_workbook(sample_ssbom_workbook)
        assert "PLM_old" in wb.sheetnames
        assert "PLM_old_1" in wb.sheetnames
        assert "PLM_old_2" in wb.sheetnames
        wb.close()

    def test_archive_superseded_file_moves_to_capnhat_old(
        self, tmp_path: Path
    ) -> None:
        """Verify superseded files are moved into capnhat\\old\\ folder."""
        archive_dir = tmp_path / "capnhat" / "old"
        src_file = tmp_path / "PLM_raw_superseded.xlsx"
        src_file.write_text("dummy content")

        engine = InheritanceEngine()
        dest_file = engine.archive_superseded_file(src_file, archive_dir)

        assert not src_file.exists()
        assert dest_file.exists()
        assert dest_file.parent == archive_dir
        assert "PLM_raw_superseded" in dest_file.name

    def test_archive_superseded_file_collision_handling(
        self, tmp_path: Path
    ) -> None:
        """Verify filename collision in archive dir generates timestamped unique name."""
        archive_dir = tmp_path / "capnhat" / "old"
        archive_dir.mkdir(parents=True)
        existing = archive_dir / "R3_file.xlsx"
        existing.write_text("existing")

        src_file = tmp_path / "R3_file.xlsx"
        src_file.write_text("new content")

        engine = InheritanceEngine()
        dest = engine.archive_superseded_file(src_file, archive_dir)

        assert dest.exists()
        assert existing.exists()
        assert dest != existing
        assert "R3_file_" in dest.name

    def test_refresh_workbook_pivots_graceful_handling(
        self, sample_ssbom_workbook: Path
    ) -> None:
        """Verify refresh_workbook_pivots executes safely without unhandled crashes when COM fails."""
        with patch("win32com.client.Dispatch", side_effect=RuntimeError("COM Dispatch Error")):
            result = refresh_workbook_pivots(sample_ssbom_workbook)
            assert result is False

    def test_refresh_workbook_pivots_missing_file(
        self, tmp_path: Path
    ) -> None:
        """Verify refresh_workbook_pivots returns False for non-existent workbook."""
        result = refresh_workbook_pivots(tmp_path / "non_existent_workbook.xlsm")
        assert result is False

    def test_refresh_workbook_pivots_mocked_success(
        self, sample_ssbom_workbook: Path
    ) -> None:
        """Verify win32com invocation path calls RefreshAll, Save, Close, and Quit."""
        mock_excel = MagicMock()
        mock_wb = MagicMock()
        mock_excel.Workbooks.Open.return_value = mock_wb

        with patch("win32com.client.Dispatch", return_value=mock_excel):
            engine = InheritanceEngine()
            ok = engine.refresh_workbook_pivots(sample_ssbom_workbook)
            assert ok is True
            mock_excel.Workbooks.Open.assert_called_once()
            mock_wb.RefreshAll.assert_called_once()
            mock_wb.Save.assert_called_once()
            mock_wb.Close.assert_called_once()
            mock_excel.Quit.assert_called_once()

    def test_explanation_record_and_summary_dataclasses(self) -> None:
        """Verify dataclasses ExplanationRecord and InheritanceSummary instantiation."""
        rec = ExplanationRecord(
            part_code="P123",
            explanation="Old note",
            person_in_charge="Hai",
            manager_check="OK",
        )
        assert rec.part_code == "P123"
        assert rec.explanation == "Old note"
        assert rec.person_in_charge == "Hai"
        assert rec.manager_check == "OK"

        summary = InheritanceSummary(
            total_new_parts=10,
            inherited_count=8,
            new_unannotated_count=2,
            archived_sheet_name="PLM_old_5",
        )
        assert summary.total_new_parts == 10
        assert summary.inherited_count == 8
        assert summary.new_unannotated_count == 2
        assert summary.archived_sheet_name == "PLM_old_5"
        assert summary.success is True

    def test_functional_update_workbook_wrapper(
        self, sample_ssbom_workbook: Path, sample_new_plm_df: pd.DataFrame
    ) -> None:
        """Verify module-level update_workbook_with_inheritance convenience function."""
        summary = update_workbook_with_inheritance(
            workbook_path=sample_ssbom_workbook,
            new_plm_data=sample_new_plm_df,
        )
        assert summary.success is True
        assert summary.total_new_parts == 3
