"""Comprehensive unit and mock tests for Teamcenter TC14 Web Automation Module.

Tests cover:
- Selectors and URL route builders (TC14Selectors, TC14URLs)
- Headless driver factory, options, CDP download behavior, and fallback mechanisms
- Session lifecycle manager, cookie persistence, and timeout detection
- High-level TC14AutomationClient and TC14Client interface contracts
- Native Excel export polling, file routing, and interactive tree scraper fallback
- 14-column PLM OpenXML Excel generation
"""

import json
import time
from unittest.mock import MagicMock, patch

import openpyxl
import pytest
from selenium.common.exceptions import (
    TimeoutException,
    WebDriverException,
)

from src.automation.tc14 import (
    BOMItem,
    BrowserConfig,
    SearchResult,
    TC14AuthenticationError,
    TC14AutomationClient,
    TC14Client,
    TC14Error,
    TC14ExportError,
    TC14NavigationError,
    TC14SearchError,
    TC14Selectors,
    TC14SessionManager,
    TC14TimeoutError,
    TC14URLs,
    create_driver,
)
from src.automation.tc14.session import (
    _build_download_prefs,
    _enable_cdp_download_behavior,
)


# =============================================================================
# 1. Selectors and URL Routes
# =============================================================================


class TestTC14SelectorsAndURLs:
    """Validate URL construction, route patterns, and selector definitions."""

    def test_search_url_construction(self):
        base = "http://tcmp3gwb:3000/"
        url = TC14URLs.get_search_url(base, "110C103NL0")
        expected = (
            "http://tcmp3gwb:3000/#/teamcenter.search.search?"
            "searchCriteria=110C103NL0&secondaryCriteria=*&isGlobalSearch=true"
        )
        assert url == expected

    def test_search_url_encoding_special_characters(self):
        base = "http://tcmp3gwb:3000"
        url = TC14URLs.get_search_url(base, "PART 123/45#A")
        assert "PART%20123/45%23A" in url
        assert url.startswith("http://tcmp3gwb:3000/#/")

    def test_object_url_construction(self):
        base = "http://tcmp3gwb:3000/"
        url = TC14URLs.get_object_url(base, "w$cJfYnMo9v21")
        assert url == (
            "http://tcmp3gwb:3000/#/com.siemens.splm.clientapi.tcui.xrt.showObject?"
            "uid=w%24cJfYnMo9v21"
        )

    def test_home_url_construction(self):
        base = "http://tcmp3gwb:3000"
        assert TC14URLs.get_home_url(base) == "http://tcmp3gwb:3000/#/showHome"

    def test_selectors_format(self):
        """Ensure all selectors are valid (By, locator_string) pairs."""
        selectors_to_check = [
            TC14Selectors.USERNAME_INPUT,
            TC14Selectors.PASSWORD_INPUT,
            TC14Selectors.LOGIN_BUTTON,
            TC14Selectors.LOGIN_ERROR,
            TC14Selectors.BANNER_HEADER,
            TC14Selectors.BANNER_USER,
            TC14Selectors.SEARCH_INPUT,
            TC14Selectors.SEARCH_CONTAINER,
            TC14Selectors.SEARCH_RESULT_ITEMS,
            TC14Selectors.SEARCH_ITEM_TITLE,
            TC14Selectors.SEARCH_ITEM_OPEN_BTN,
            TC14Selectors.NAVIGATION_TABS,
            TC14Selectors.CONTENT_TAB,
            TC14Selectors.EXPORT_EXCEL_COMMAND,
            TC14Selectors.EXPORT_DIALOG,
            TC14Selectors.EXPORT_CONFIRM_BTN,
            TC14Selectors.TREE_TABLE_CONTAINER,
            TC14Selectors.TREE_PINNED_ROWS,
            TC14Selectors.TREE_EXPAND_BTN,
        ]
        for by_type, locator in selectors_to_check:
            assert isinstance(by_type, str)
            assert isinstance(locator, str)
            assert len(locator) > 0


# =============================================================================
# 2. Driver Factory, Preferences, and CDP
# =============================================================================


class TestDriverFactory:
    """Validate headless browser options, preferences, and CDP setup."""

    def test_browser_config_defaults(self):
        cfg = BrowserConfig()
        assert cfg.browser_type == "edge"
        assert cfg.headless is True
        assert cfg.window_size == "1920,1080"
        assert cfg.page_load_timeout == 45
        assert cfg.session_timeout_seconds == 1800

    def test_build_download_prefs(self, tmp_path):
        prefs = _build_download_prefs(tmp_path)
        assert prefs["download.default_directory"] == str(tmp_path.resolve())
        assert prefs["download.prompt_for_download"] is False
        assert prefs["download.directory_upgrade"] is True
        assert prefs["safebrowsing.enabled"] is True

    def test_enable_cdp_download_behavior(self, tmp_path):
        mock_driver = MagicMock()
        _enable_cdp_download_behavior(mock_driver, tmp_path)
        mock_driver.execute_cdp_cmd.assert_called_once_with(
            "Page.setDownloadBehavior",
            {
                "behavior": "allow",
                "downloadPath": str(tmp_path.resolve()),
            },
        )

    @patch("src.automation.tc14.session.webdriver.Edge")
    def test_create_driver_edge_success(self, mock_edge, tmp_path):
        mock_instance = MagicMock()
        mock_edge.return_value = mock_instance

        driver = create_driver(browser_type="edge", download_dir=tmp_path, headless=True)
        assert driver == mock_instance
        mock_edge.assert_called_once()
        mock_instance.execute_cdp_cmd.assert_called_once()

    @patch("src.automation.tc14.session.webdriver.Chrome")
    @patch("src.automation.tc14.session.webdriver.Edge")
    def test_create_driver_fallback_to_chrome(self, mock_edge, mock_chrome, tmp_path):
        """When Edge fails to start, create_driver should fallback to Chrome."""
        mock_edge.side_effect = WebDriverException("Edge binary not found")
        mock_chrome_instance = MagicMock()
        mock_chrome.return_value = mock_chrome_instance

        driver = create_driver(browser_type="edge", download_dir=tmp_path)
        assert driver == mock_chrome_instance
        mock_edge.assert_called_once()
        mock_chrome.assert_called_once()

    @patch("src.automation.tc14.session.webdriver.Chrome")
    @patch("src.automation.tc14.session.webdriver.Edge")
    def test_create_driver_all_fail_raises(self, mock_edge, mock_chrome, tmp_path):
        mock_edge.side_effect = WebDriverException("Edge fail")
        mock_chrome.side_effect = WebDriverException("Chrome fail")

        with pytest.raises(WebDriverException) as exc_info:
            create_driver(browser_type="edge", download_dir=tmp_path)
        assert "Failed to initialize any WebDriver" in str(exc_info.value)


# =============================================================================
# 3. Session Lifecycle & Timeout Detection
# =============================================================================


class TestTC14SessionManager:
    """Validate session tracking, timeout detection, and cookie persistence."""

    def test_initial_state(self):
        mgr = TC14SessionManager()
        assert mgr._driver is None
        assert mgr.is_authenticated is False
        assert mgr.last_activity_time == 0.0

    def test_touch_updates_activity(self):
        mgr = TC14SessionManager()
        assert mgr.last_activity_time == 0.0
        mgr.touch()
        assert mgr.last_activity_time > 0.0

    def test_mark_authenticated_and_unauthenticated(self):
        mgr = TC14SessionManager()
        mgr.mark_authenticated("vn_pe03")
        assert mgr.is_authenticated is True
        assert mgr._current_username == "vn_pe03"

        mgr.mark_unauthenticated()
        assert mgr.is_authenticated is False
        assert mgr._current_username is None

    def test_is_session_timed_out(self):
        mgr = TC14SessionManager(config=BrowserConfig(session_timeout_seconds=10))
        # Initial: 0.0 means timed out
        assert mgr.is_session_timed_out() is True

        # Recent touch: active
        mgr.touch()
        assert mgr.is_session_timed_out() is False

        # Simulate past time > 10s
        mgr._last_activity_time = time.time() - 15
        assert mgr.is_session_timed_out() is True

    def test_is_login_page_present_by_url(self):
        mock_driver = MagicMock()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/login"
        mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        mgr.start_session()

        assert mgr.is_login_page_present() is True

    def test_is_login_page_present_by_input(self):
        mock_driver = MagicMock()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/showHome"
        mock_elem = MagicMock()
        mock_elem.is_displayed.return_value = True
        mock_driver.find_elements.return_value = [mock_elem]

        mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        mgr.start_session()
        assert mgr.is_login_page_present() is True

    def test_is_session_alive_conditions(self):
        mock_driver = MagicMock()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/showHome"
        mock_driver.find_elements.return_value = []  # No login inputs

        mgr = TC14SessionManager(
            config=BrowserConfig(session_timeout_seconds=60),
            driver_factory=lambda: mock_driver,
        )
        mgr.start_session()
        mgr.mark_authenticated("vn_pe03")

        assert mgr.is_session_alive() is True

        # When timed out
        mgr._last_activity_time = time.time() - 100
        assert mgr.is_session_alive() is False

        # When login page appears
        mgr.touch()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/login"
        assert mgr.is_session_alive() is False
        assert mgr.is_authenticated is False

    def test_cookie_save_and_load(self, tmp_path):
        cookie_file = tmp_path / "cookies.json"
        mock_driver = MagicMock()
        sample_cookies = [
            {"name": "JSESSIONID", "value": "xyz123", "domain": "tcmp3gwb", "path": "/"},
            {"name": "XSRF-TOKEN", "value": "tok456", "domain": "tcmp3gwb", "path": "/"},
        ]
        mock_driver.get_cookies.return_value = sample_cookies

        mgr = TC14SessionManager(
            config=BrowserConfig(cookie_file=cookie_file),
            driver_factory=lambda: mock_driver,
        )
        mgr.start_session()

        # Save cookies
        saved = mgr.save_cookies()
        assert len(saved) == 2
        assert cookie_file.exists()
        with open(cookie_file, "r") as f:
            data = json.load(f)
        assert data[0]["name"] == "JSESSIONID"

        # Load cookies
        loaded = mgr.load_cookies("http://tcmp3gwb:3000/")
        assert loaded is True
        mock_driver.get.assert_called_with("http://tcmp3gwb:3000/")
        assert mock_driver.add_cookie.call_count == 2

    def test_close_and_context_manager(self):
        mock_driver = MagicMock()
        mgr = TC14SessionManager(driver_factory=lambda: mock_driver)

        with mgr as m:
            assert m._driver is not None

        mock_driver.quit.assert_called_once()
        assert mgr._driver is None


# =============================================================================
# 4. TC14AutomationClient Actions: Login, Search, Navigation, Export
# =============================================================================


class TestTC14AutomationClient:
    """Validate client high-level automation methods with robust mocking."""

    @pytest.fixture
    def mock_session_and_client(self, tmp_path):
        mock_driver = MagicMock()
        mock_driver.find_elements.return_value = []
        mock_driver.find_element.return_value = MagicMock()
        session_mgr = TC14SessionManager(
            config=BrowserConfig(download_dir=tmp_path),
            driver_factory=lambda: mock_driver,
        )
        session_mgr.start_session()
        client = TC14AutomationClient(
            base_url="http://tcmp3gwb:3000/",
            download_dir=tmp_path,
            session_manager=session_mgr,
        )
        return mock_driver, session_mgr, client, tmp_path

    # --- Login Tests ---

    def test_login_success(self, mock_session_and_client):
        mock_driver, session_mgr, client, _ = mock_session_and_client
        mock_driver.current_url = "http://tcmp3gwb:3000/#/showHome"

        # Form elements
        user_input = MagicMock()
        pass_input = MagicMock()
        login_btn = MagicMock()

        def mock_find_element(by, sel):
            if "username" in sel:
                return user_input
            if "password" in sel:
                return pass_input
            if "button" in sel:
                return login_btn
            return MagicMock()

        mock_driver.find_element.side_effect = mock_find_element
        mock_driver.find_elements.return_value = []  # No errors

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            # 1: presence of username, 2: clickable login_btn, 3: _login_succeeded
            wait_instance.until.side_effect = [user_input, login_btn, True]

            result = client.login("vn_pe03", "vn_pe03")
            assert result is True
            assert session_mgr.is_authenticated is True
            user_input.send_keys.assert_called_with("vn_pe03")
            pass_input.send_keys.assert_called_with("vn_pe03")
            login_btn.click.assert_called_once()

    def test_login_already_active_skips(self, mock_session_and_client):
        mock_driver, session_mgr, client, _ = mock_session_and_client
        session_mgr.mark_authenticated("vn_pe03")
        session_mgr.touch()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/showHome"
        mock_driver.find_elements.return_value = []

        result = client.login("vn_pe03", "vn_pe03", force=False)
        assert result is True
        mock_driver.get.assert_not_called()

    def test_login_timeout_raises_error(self, mock_session_and_client):
        mock_driver, _, client, _ = mock_session_and_client

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.side_effect = TimeoutException("Timed out")
            mock_driver.find_elements.return_value = []

            with pytest.raises(TC14AuthenticationError) as exc_info:
                client.login("vn_pe03", "wrong_pass")
            assert "Failed to log in" in str(exc_info.value)

    # --- Search Tests ---

    def test_search_part_direct_hash_success(self, mock_session_and_client):
        mock_driver, session_mgr, client, _ = mock_session_and_client
        session_mgr.mark_authenticated("vn_pe03")
        session_mgr.touch()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/teamcenter.search.search?..."

        # Mock result items
        mock_item = MagicMock()
        mock_title = MagicMock()
        mock_title.text = "110C103NL0 / PBA-00001711 / Revision:A"
        mock_title.get_attribute.return_value = "http://tcmp3gwb:3000/#/showObject?uid=u123"
        mock_item.find_element.return_value = mock_title
        mock_item.is_displayed.return_value = True

        def mock_find_elements(by, sel):
            if "username" in sel:
                return []
            return [mock_item]

        mock_driver.find_elements.side_effect = mock_find_elements

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.return_value = True

            results = client.search_part("110C103NL0")
            assert len(results) == 1
            res = results[0]
            assert res.item_id == "110C103NL0"
            assert res.revision == "A"
            assert res.uid == "u123"
            assert "PBA-00001711" in res.title

    def test_search_part_no_results_raises_error(self, mock_session_and_client):
        mock_driver, session_mgr, client, _ = mock_session_and_client
        session_mgr.mark_authenticated("vn_pe03")
        session_mgr.touch()

        mock_driver.find_elements.return_value = []

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.return_value = True

            with pytest.raises(TC14SearchError) as exc_info:
                client.search_part("NONEXISTENT_PART")
            assert "No search results found" in str(exc_info.value)

    # --- Item Navigation & Tabs ---

    def test_open_item_by_search_result(self, mock_session_and_client):
        mock_driver, _, client, _ = mock_session_and_client

        mock_elem = MagicMock()
        mock_btn = MagicMock()
        mock_elem.find_element.return_value = mock_btn

        result = SearchResult(
            item_id="110C103NL0",
            title="Product Revision",
            revision="A",
            uid="uid_abc_999",
            web_element=mock_elem,
        )

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.return_value = True

            opened = client.open_item(result)
            assert opened is True
            # Should have navigated via direct object URL
            mock_driver.get.assert_called_with(
                "http://tcmp3gwb:3000/#/com.siemens.splm.clientapi.tcui.xrt.showObject?uid=uid_abc_999"
            )

    def test_open_item_by_element_click(self, mock_session_and_client):
        mock_driver, _, client, _ = mock_session_and_client

        mock_elem = MagicMock()
        mock_btn = MagicMock()
        mock_elem.find_element.return_value = mock_btn

        result = SearchResult(
            item_id="110C103NL0",
            title="Product Revision",
            revision="A",
            uid=None,  # No UID forces element click
            web_element=mock_elem,
        )

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.return_value = True

            opened = client.open_item(result)
            assert opened is True
            mock_btn.click.assert_called_once()

    def test_navigate_to_content_tab_success(self, mock_session_and_client):
        mock_driver, _, client, _ = mock_session_and_client

        content_tab = MagicMock()
        tree_table = MagicMock()

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.side_effect = [content_tab, tree_table]

            success = client.navigate_to_content_tab()
            assert success is True
            content_tab.click.assert_called_once()

    def test_navigate_to_content_tab_timeout_raises_error(self, mock_session_and_client):
        mock_driver, _, client, _ = mock_session_and_client

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.side_effect = TimeoutException("Tab missing")
            mock_driver.find_elements.return_value = []

            with pytest.raises(TC14NavigationError) as exc_info:
                client.navigate_to_content_tab()
            assert "Failed to navigate to Content tab" in str(exc_info.value)

    # --- Excel Export Tests ---

    def test_export_bom_excel_success(self, mock_session_and_client):
        mock_driver, _, client, tmp_path = mock_session_and_client

        export_cmd = MagicMock()
        dialog = MagicMock()
        confirm_btn = MagicMock()

        # Simulate downloaded file appearing in download_dir
        dest_file = tmp_path / "TC_Export_123.xlsx"
        wb = openpyxl.Workbook()
        wb.save(str(dest_file))

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.side_effect = [export_cmd, dialog, confirm_btn]

            with patch.object(client, "_poll_for_downloaded_file", return_value=dest_file):
                final_path = client.export_bom_excel(
                    output_dir=tmp_path,
                    target_filename="PLM_110C103NL0.xlsx",
                )
                assert final_path.name == "PLM_110C103NL0.xlsx"
                assert final_path.exists()
                export_cmd.click.assert_called_once()
                confirm_btn.click.assert_called_once()

    def test_export_bom_excel_timeout_raises(self, mock_session_and_client):
        mock_driver, _, client, tmp_path = mock_session_and_client

        with patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_instance = MagicMock()
            mock_wait_cls.return_value = wait_instance
            wait_instance.until.side_effect = [MagicMock(), MagicMock(), MagicMock()]

            with patch.object(client, "_poll_for_downloaded_file", return_value=None):
                with pytest.raises(TC14ExportError) as exc_info:
                    client.export_bom_excel(timeout=2)
                assert "Export timed out" in str(exc_info.value)

    # --- Tree Scraper & OpenXML Synthesis ---

    def test_scrape_bom_tree_interactive(self, mock_session_and_client):
        mock_driver, _, client, _ = mock_session_and_client

        # Row 1: Level 1 parent
        r1 = MagicMock()
        r1.text = "110C103NL0 Main Assembly"
        r1.get_attribute.side_effect = lambda attr: "1" if attr == "aria-level" else "true"
        c1 = MagicMock()
        c1.text = "110C103NL0"
        c2 = MagicMock()
        c2.text = "Main Assembly"
        r1.find_elements.return_value = [c1, c2]

        # Row 2: Level 2 child
        r2 = MagicMock()
        r2.text = "302V493010 Motor Unit"
        r2.get_attribute.side_effect = lambda attr: "2" if attr == "aria-level" else "false"
        c3 = MagicMock()
        c3.text = "302V493010"
        c4 = MagicMock()
        c4.text = "Motor Unit"
        r2.find_elements.return_value = [c3, c4]

        mock_driver.find_elements.return_value = [r1, r2]

        items = client.scrape_bom_tree_interactive(max_depth=3, expand_children=False)
        assert len(items) == 2
        assert items[0].level == 1
        assert items[0].item_id == "110C103NL0"
        assert items[0].has_children is True

        assert items[1].level == 2
        assert items[1].item_id == "302V493010"
        assert items[1].has_children is False

    def test_synthesize_plm_excel_from_tree(self, mock_session_and_client):
        _, _, client, tmp_path = mock_session_and_client
        target_path = tmp_path / "PLM_Synthesized.xlsx"

        items = [
            BOMItem(
                level=1,
                item_id="110C103NL0",
                item_name="Machine Model",
                has_children=True,
                effectivity="UP",
            ),
            BOMItem(
                level=2,
                item_id="302V493010",
                item_name="Sub Unit A",
                has_children=False,
                effectivity="to 2026/12",
            ),
        ]

        result_path = client.synthesize_plm_excel_from_tree(items, target_path)
        assert result_path.exists()

        # Verify columns using openpyxl
        wb = openpyxl.load_workbook(str(result_path))
        ws = wb.active
        assert ws.max_row == 3  # Header + 2 rows
        assert ws.max_column == 14  # 14 columns

        # Check Col B (Level)
        assert ws.cell(row=2, column=2).value == 1
        assert ws.cell(row=3, column=2).value == 2

        # Check Col D (Part Number)
        assert ws.cell(row=2, column=4).value == "110C103NL0"
        assert ws.cell(row=3, column=4).value == "302V493010"

        # Check Col E (hasChildren)
        assert ws.cell(row=2, column=5).value == "True"
        assert ws.cell(row=3, column=5).value == "False"

        # Check Col I (Effectivity)
        assert ws.cell(row=2, column=9).value == "UP"
        assert ws.cell(row=3, column=9).value == "to 2026/12"

        # Check Col K (Item Name)
        assert ws.cell(row=2, column=11).value == "Machine Model"
        assert ws.cell(row=3, column=11).value == "Sub Unit A"

    # --- Project Interface Contract: download_bom_full ---

    def test_download_bom_full_native_export(self, mock_session_and_client):
        mock_driver, _, client, tmp_path = mock_session_and_client
        out_dir = tmp_path / "output_plm"

        fake_exported_file = tmp_path / "downloaded_bom.xlsx"
        fake_exported_file.write_text("dummy")

        with patch.object(client, "ensure_authenticated", return_value=True), \
             patch.object(client, "search_part", return_value=[MagicMock()]), \
             patch.object(client, "open_item", return_value=True), \
             patch.object(client, "navigate_to_content_tab", return_value=True), \
             patch.object(client, "export_bom_excel", return_value=out_dir / "PLM_110C103NL0.xlsx") as mock_export:

            res = client.download_bom_full("110C103NL0", out_dir)
            assert res == out_dir / "PLM_110C103NL0.xlsx"
            mock_export.assert_called_once()

    def test_download_bom_full_fallback_scraper(self, mock_session_and_client):
        """If native export fails with TC14ExportError, fall back to interactive scraper."""
        mock_driver, _, client, tmp_path = mock_session_and_client
        out_dir = tmp_path / "output_plm_fallback"

        scraped_items = [
            BOMItem(level=1, item_id="110C103NL0", item_name="Machine", has_children=True),
        ]

        with patch.object(client, "ensure_authenticated", return_value=True), \
             patch.object(client, "search_part", return_value=[MagicMock()]), \
             patch.object(client, "open_item", return_value=True), \
             patch.object(client, "navigate_to_content_tab", return_value=True), \
             patch.object(client, "export_bom_excel", side_effect=TC14ExportError("Export blocked")), \
             patch.object(client, "scrape_bom_tree_interactive", return_value=scraped_items) as mock_scrape, \
             patch.object(
                 client,
                 "synthesize_plm_excel_from_tree",
                 return_value=out_dir / "PLM_110C103NL0.xlsx",
             ) as mock_synth:

            res = client.download_bom_full("110C103NL0", out_dir)
            assert res == out_dir / "PLM_110C103NL0.xlsx"
            mock_scrape.assert_called_once()
            mock_synth.assert_called_once()

    def test_tc14_client_alias_inheritance(self):
        """Confirm TC14Client is a proper subclass of TC14AutomationClient."""
        assert issubclass(TC14Client, TC14AutomationClient)
        mock_mgr = MagicMock()
        client = TC14Client(session_manager=mock_mgr)
        assert hasattr(client, "download_bom_full")
        assert hasattr(client, "login")
        assert hasattr(client, "search_part")


# =============================================================================
# 5. Exceptions Hierarchy
# =============================================================================


class TestTC14Exceptions:
    """Verify exception inheritance hierarchy."""

    def test_hierarchy(self):
        assert issubclass(TC14AuthenticationError, TC14Error)
        assert issubclass(TC14SearchError, TC14Error)
        assert issubclass(TC14NavigationError, TC14Error)
        assert issubclass(TC14ExportError, TC14Error)
        assert issubclass(TC14TimeoutError, TC14Error)
        assert issubclass(TC14Error, Exception)


# =============================================================================
# 6. Additional Edge Cases & Boundary Coverage
# =============================================================================


class TestTC14AdditionalEdgeCases:
    """Additional edge cases for session management, file polling, and client actions."""

    def test_create_driver_with_extra_args(self, tmp_path):
        with patch("src.automation.tc14.session.webdriver.Edge") as mock_edge:
            mock_inst = MagicMock()
            mock_edge.return_value = mock_inst
            driver = create_driver(
                browser_type="edge",
                download_dir=tmp_path,
                extra_arguments=["--user-agent=TestAgent"],
            )
            assert driver == mock_inst
            options_passed = mock_edge.call_args[1]["options"]
            assert "--user-agent=TestAgent" in options_passed.arguments

    def test_enable_cdp_exception_handled(self, tmp_path):
        mock_driver = MagicMock()
        mock_driver.execute_cdp_cmd.side_effect = Exception("CDP not supported")
        # Should not raise
        _enable_cdp_download_behavior(mock_driver, tmp_path)

    def test_is_login_page_present_exception_handled(self):
        mock_driver = MagicMock()
        mock_driver.current_url = "http://tcmp3gwb:3000/#/showHome"
        mock_driver.find_elements.side_effect = Exception("Driver dead")
        mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        mgr.start_session()
        assert mgr.is_login_page_present() is False

    def test_load_cookies_invalid_paths_and_exceptions(self, tmp_path):
        mock_driver = MagicMock()
        mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        mgr.start_session()

        # Nonexistent file
        assert mgr.load_cookies("http://tcmp3gwb:3000/", source_path=tmp_path / "missing.json") is False

        # Corrupt file
        bad_json = tmp_path / "bad.json"
        bad_json.write_text("NOT_JSON")
        assert mgr.load_cookies("http://tcmp3gwb:3000/", source_path=bad_json) is False

        # Empty list
        assert mgr.load_cookies("http://tcmp3gwb:3000/", cookies=[]) is False

        # Exception during add_cookie handled safely
        mock_driver.add_cookie.side_effect = Exception("Cookie error")
        sample_cookie = [{"name": "test", "value": "val"}]
        assert mgr.load_cookies("http://tcmp3gwb:3000/", cookies=sample_cookie) is True

    def test_client_context_manager(self, tmp_path):
        mock_driver = MagicMock()
        session_mgr = TC14SessionManager(
            config=BrowserConfig(download_dir=tmp_path),
            driver_factory=lambda: mock_driver,
        )
        session_mgr.start_session()
        client = TC14AutomationClient(download_dir=tmp_path, session_manager=session_mgr)

        with client as c:
            assert c == client
        mock_driver.quit.assert_called_once()

    def test_ensure_authenticated_triggers_relogin(self, tmp_path):
        mock_driver = MagicMock()
        mock_driver.find_elements.return_value = []
        session_mgr = TC14SessionManager(
            config=BrowserConfig(download_dir=tmp_path),
            driver_factory=lambda: mock_driver,
        )
        session_mgr.start_session()
        # Not authenticated initially
        client = TC14AutomationClient(download_dir=tmp_path, session_manager=session_mgr)

        with patch.object(client, "login", return_value=True) as mock_login:
            res = client.ensure_authenticated("user", "pass")
            assert res is True
            mock_login.assert_called_once_with(username="user", password="pass")

    def test_open_item_index_out_of_bounds(self, tmp_path):
        mock_driver = MagicMock()
        mock_driver.find_elements.return_value = []
        session_mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        session_mgr.start_session()
        client = TC14AutomationClient(download_dir=tmp_path, session_manager=session_mgr)

        with pytest.raises(TC14NavigationError) as exc_info:
            client.open_item(5)  # Out of range
        assert "out of range" in str(exc_info.value)

    def test_open_item_by_part_number_string(self, tmp_path):
        mock_driver = MagicMock()
        session_mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        session_mgr.start_session()
        client = TC14AutomationClient(download_dir=tmp_path, session_manager=session_mgr)

        mock_search_res = [SearchResult(item_id="110C103NL0", title="Test Item", uid="uid_123")]
        with patch.object(client, "search_part", return_value=mock_search_res) as mock_search, \
             patch("src.automation.tc14.client.WebDriverWait") as mock_wait_cls:
            wait_inst = MagicMock()
            mock_wait_cls.return_value = wait_inst
            wait_inst.until.return_value = True

            res = client.open_item("110C103NL0")
            assert res is True
            mock_search.assert_called_once_with("110C103NL0")
            mock_driver.get.assert_called_with(
                "http://tcmp3gwb:3000/#/com.siemens.splm.clientapi.tcui.xrt.showObject?uid=uid_123"
            )

    def test_poll_for_downloaded_file_detects_complete_file(self, tmp_path):
        mock_driver = MagicMock()
        session_mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        session_mgr.start_session()
        client = TC14AutomationClient(download_dir=tmp_path, session_manager=session_mgr)

        # Create temporary file (should be ignored)
        tmp_file = tmp_path / "download.crdownload"
        tmp_file.write_text("temp")

        # Create real file
        real_file = tmp_path / "valid_download.xlsx"
        real_file.write_text("excel data content")

        found = client._poll_for_downloaded_file(
            directory=tmp_path,
            initial_files=set(),
            timeout=2,
            poll_interval=0.1,
        )
        assert found == real_file

    def test_expand_all_visible_tree_nodes(self, tmp_path):
        mock_driver = MagicMock()
        session_mgr = TC14SessionManager(driver_factory=lambda: mock_driver)
        session_mgr.start_session()
        client = TC14AutomationClient(download_dir=tmp_path, session_manager=session_mgr)

        toggle1 = MagicMock()
        toggle1.is_displayed.return_value = True

        mock_driver.find_elements.side_effect = [
            [toggle1],  # Depth 1 returns 1 toggle
            [],         # Depth 2 returns 0 toggles -> breaks
        ]

        client._expand_all_visible_tree_nodes(max_depth=3)
        toggle1.click.assert_called_once()
        mock_driver.execute_script.assert_called_once()
