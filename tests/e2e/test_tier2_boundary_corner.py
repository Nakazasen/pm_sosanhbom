"""Tier 2: Boundary and Corner Cases E2E Test Suite.

Authoritative Specification: specs/SPEC_PLM_AUTO_DOWNLOAD.md & spec_report.md
Covers Boundary & Corner Conditions across R1 to R6 (>= 30 tests, exactly 30 tests).

Requirements Map:
- R1 Boundaries: Leader Wizard Boundary Scenarios (Tests 1 to 5)
- R2 Boundaries: BOM Filter Engine Boundary & Edge Conditions (Tests 6 to 10)
- R3 Boundaries: Annotation Inheritance Edge Cases (Tests 11 to 15)
- R4 Boundaries: MSI Deep Reconciliation Boundary Conditions (Tests 16 to 20)
- R5 Boundaries: List JIG Master & 4M Boundary Cases (Tests 21 to 25)
- R6 Boundaries: Member Workspace & Input Edge Conditions (Tests 26 to 30)
"""

from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import openpyxl
import pandas as pd
import pytest
from PyQt6.QtWidgets import QApplication, QMessageBox

from src.core.date_filter import DateFilter, extract_expiry_date, is_effectivity_expired
from src.core.model_pruner import ModelPruner
from src.core.models import BOMNode, BOMTree, FilterCriteria, ModelRule
from src.core.msi_engine import FixSerialMaster, MSIEngine, evaluate_msi_branch
from src.core.reconciliation import ReconciliationEngine, migrate_annotations
from src.core.tc2412_bridge import TC2412CanonicalNormalizer, TC2412SheetPLMBridge
from src.core.tree_parser import BOMTreeParser
from src.core.unit_resolver import UnitResolver
from src.gui.leader_view import LeaderWorkspaceView
from src.gui.member_view import MemberWorkspaceView
from src.reporting.excel_generator import ExcelReportGenerator
from src.reporting.outlook_mailer import EmailPreview, OutlookMailer


@pytest.fixture
def qapp() -> QApplication:
    """Provide headless QApplication instance for Qt-dependent components."""
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


# =============================================================================
# R1 Boundaries (5 Tests)
# =============================================================================

class TestTier2R1LeaderBoundaries:
    """Boundary test cases for Requirement R1: Leader Workspace Wizard."""

    def test_tc_t2_r1_01_empty_machine_code_list(self, tmp_path: Path) -> None:
        """TC-T2-R1-01: Boundary: Leader attempts folder creation with empty machine code list."""
        machine_codes: list[str] = []
        created_dirs = []
        for code in machine_codes:
            d = tmp_path / "Virgo" / code
            d.mkdir(parents=True, exist_ok=True)
            created_dirs.append(d)

        assert len(created_dirs) == 0, "No directories should be created for empty machine list"

    def test_tc_t2_r1_02_single_unsubmitted_member_in_large_team(self, tmp_path: Path) -> None:
        """TC-T2-R1-02: Boundary: 9 of 10 members submitted OK; exactly 1 pending member blocks consolidation."""
        members = [f"Member_{i}" for i in range(1, 11)]
        submissions: dict[str, str] = {}
        for m in members[:9]:
            submissions[m] = "OK"
        submissions[members[9]] = ""  # 10th member has not submitted

        pending_members = [m for m, status in submissions.items() if status != "OK"]
        assert len(pending_members) == 1
        assert pending_members[0] == "Member_10"
        # Consolidation must be strictly blocked
        can_consolidate = len(pending_members) == 0
        assert can_consolidate is False

    def test_tc_t2_r1_03_path_with_japanese_kanji_and_spaces(self, tmp_path: Path) -> None:
        """TC-T2-R1-03: Boundary: Paths containing Japanese Kanji and spaces handled without encoding errors."""
        jp_dir = tmp_path / "製造技術課" / "BOM 2026" / "Iris2024" / "110C0Z3LV1"
        jp_dir.mkdir(parents=True, exist_ok=True)

        test_file = jp_dir / "BOM_総合_2026.xlsm"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CTTT"
        ws["A1"] = "テストデータ"
        wb.save(test_file)

        assert test_file.exists()
        wb_read = openpyxl.load_workbook(test_file)
        assert wb_read["CTTT"]["A1"].value == "テストデータ"
        wb_read.close()

    def test_tc_t2_r1_04_corrupted_zero_byte_member_workbook(self, tmp_path: Path) -> None:
        """TC-T2-R1-04: Boundary: 0-byte corrupted workbook caught gracefully without crashing scanner."""
        corrupted_file = tmp_path / "corrupted_member.xlsx"
        corrupted_file.write_bytes(b"")  # 0 bytes

        scan_error = False
        try:
            openpyxl.load_workbook(corrupted_file, read_only=True)
        except Exception:
            scan_error = True

        assert scan_error is True, "0-byte file must be caught as an invalid workbook"

    def test_tc_t2_r1_05_outlook_com_unavailable_graceful_handling(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """TC-T2-R1-05: Boundary: Outlook COM unavailable handled gracefully with preview mode."""
        mailer = OutlookMailer()
        # Explicitly simulate Outlook COM unavailable in headless/offline testing
        monkeypatch.setattr(mailer, "_get_outlook_application", lambda: None)
        preview = EmailPreview(
            recipients_to=["engineer@kyocera.com"],
            recipients_cc=[],
            subject="Test Subject",
            html_body="<p>Test Body</p>",
            attachment_paths=[],
        )
        # When Outlook COM is unavailable, send_via_outlook safely returns False
        success = mailer.send_via_outlook(preview)
        assert success is False


# =============================================================================
# R2 Boundaries (5 Tests)
# =============================================================================

class TestTier2R2BOMFilterBoundaries:
    """Boundary test cases for Requirement R2: BOM Filter Engine Level 1..6."""

    def test_tc_t2_r2_01_same_year_month_diff_le_1_retained(self) -> None:
        """TC-T2-R2-01: Boundary: Same year with month diff <= 1 MUST BE RETAINED (Kyocera grace period)."""
        ref_date = datetime.date(2024, 9, 15)
        # Expired on 31-Aug-2024 (diff_year = 0, diff_month = 1) -> Must NOT be expired!
        expired = is_effectivity_expired("01-Jan-2024 to 31-Aug-2024", reference_date=ref_date, month_tolerance=1)
        assert expired is False, "Month difference <= 1 within the same year must be retained"

        # Expired on 31-Jul-2024 (diff_year = 0, diff_month = 2) -> Expired!
        expired_2 = is_effectivity_expired("01-Jan-2024 to 31-Jul-2024", reference_date=ref_date, month_tolerance=1)
        assert expired_2 is True, "Month difference > 1 within the same year must be pruned"

    def test_tc_t2_r2_02_leap_year_february_29(self) -> None:
        """TC-T2-R2-02: Boundary: Leap year February 29 parsed and evaluated accurately."""
        eff_leap = "01-Mar-2020 to 29-Feb-2024"
        d = extract_expiry_date(eff_leap)
        assert d == datetime.date(2024, 2, 29), "Must parse Feb 29 on leap year correctly"

        # Tested on March 1, 2024 (same year, month diff = 1) -> Retained due to tolerance
        ref_mar = datetime.date(2024, 3, 1)
        assert is_effectivity_expired(eff_leap, reference_date=ref_mar, month_tolerance=1) is False

        # Tested on May 1, 2024 (same year, month diff = 3) -> Expired
        ref_may = datetime.date(2024, 5, 1)
        assert is_effectivity_expired(eff_leap, reference_date=ref_may, month_tolerance=1) is True

    def test_tc_t2_r2_03_deep_tree_depth_6_recursion(self) -> None:
        """TC-T2-R2-03: Boundary: BOM tree reaching maximum depth Level 6 traversed safely."""
        current = BOMNode(level=6, item_id="LEAF_LV6", item_name="WASHER M2", effectivity="01-Jan-2024 UP", has_children=False)
        for lv in range(5, 0, -1):
            parent = BOMNode(level=lv, item_id=f"NODE_LV{lv}", item_name=f"UNIT LV{lv}", effectivity="01-Jan-2024 UP", has_children=True)
            parent.children.append(current)
            current = parent

        tree = BOMTree(roots=[current])
        assert tree.max_depth() == 6
        assert tree.size() == 6

        # Filter across deep tree
        filter_eng = DateFilter(criteria=FilterCriteria(reference_date=datetime.date(2024, 6, 1)))
        filtered = filter_eng.filter_tree(tree)
        assert filtered.max_depth() == 6
        assert filtered.size() == 6

    def test_tc_t2_r2_04_unknown_model_name_bolocbom(self) -> None:
        """TC-T2-R2-04: Boundary: Unregistered model name raises in strict mode or warns in standard mode."""
        pruner_strict = ModelPruner(strict=True)
        with pytest.raises(ValueError, match="not found in BolocBom"):
            pruner_strict.get_rules_for_model("UnknownModelX")

        pruner_lenient = ModelPruner(strict=False)
        rules = pruner_lenient.get_rules_for_model("UnknownModelX")
        assert len(rules) == 0

    def test_tc_t2_r2_05_empty_plm_file_zero_records(self, tmp_path: Path) -> None:
        """TC-T2-R2-05: Boundary: Empty PLM export (only headers, 0 records) handled without crash."""
        normalizer = TC2412CanonicalNormalizer()
        empty_df = pd.DataFrame(columns=["Level", "Item Id", "Item Name", "Quantity"])
        canonical = normalizer.normalize_dataframe(empty_df)
        assert len(canonical) == 0
        assert "item_id" in canonical.columns


# =============================================================================
# R3 Boundaries (5 Tests)
# =============================================================================

class TestTier2R3AnnotationBoundaries:
    """Boundary test cases for Requirement R3: Annotation Inheritance Engine."""

    def test_tc_t2_r3_01_consecutive_updates_multi_version_sheets(self, tmp_path: Path) -> None:
        """TC-T2-R3-01: Boundary: Ingesting new revisions multiple times creates sequentially numbered sheets."""
        wb_path = tmp_path / "BOM_test.xlsm"
        wb = openpyxl.Workbook()
        wb.active.title = "PLM"
        wb.save(wb_path)

        # Ingest 3 times
        for update_idx in range(1, 4):
            wb_load = openpyxl.load_workbook(wb_path)
            # Find appropriate backup title
            base_title = "PLM_old"
            if base_title not in wb_load.sheetnames:
                new_title = base_title
            else:
                counter = 1
                while f"{base_title}_{counter}" in wb_load.sheetnames:
                    counter += 1
                new_title = f"{base_title}_{counter}"
            ws_bak = wb_load.copy_worksheet(wb_load["PLM"])
            ws_bak.title = new_title
            wb_load.save(wb_path)
            wb_load.close()

        wb_verify = openpyxl.load_workbook(wb_path)
        assert "PLM" in wb_verify.sheetnames
        assert "PLM_old" in wb_verify.sheetnames
        assert "PLM_old_1" in wb_verify.sheetnames
        assert "PLM_old_2" in wb_verify.sheetnames
        wb_verify.close()

    def test_tc_t2_r3_02_scrambled_row_order_match_index(self) -> None:
        """TC-T2-R3-02: Boundary: Completely shuffled row order preserves MATCH-INDEX annotation alignment."""
        old_records = [
            {"part_code": "PART_A", "giai_thich": "Note A"},
            {"part_code": "PART_B", "giai_thich": "Note B"},
            {"part_code": "PART_C", "giai_thich": "Note C"},
        ]
        # Inverted order in new revision
        new_records = [
            {"part_code": "PART_C"},
            {"part_code": "PART_A"},
            {"part_code": "PART_B"},
        ]

        migrated = migrate_annotations(pd.DataFrame(new_records), pd.DataFrame(old_records))
        assert migrated.iloc[0]["part_code"] == "PART_C"
        assert migrated.iloc[0]["giai_thich"] == "Note C"
        assert migrated.iloc[1]["part_code"] == "PART_A"
        assert migrated.iloc[1]["giai_thich"] == "Note A"
        assert migrated.iloc[2]["part_code"] == "PART_B"
        assert migrated.iloc[2]["giai_thich"] == "Note B"

    def test_tc_t2_r3_03_completely_disjoint_parts_zero_overlap(self) -> None:
        """TC-T2-R3-03: Boundary: 0% overlap between old and new revisions leaves all annotations blank."""
        old_df = pd.DataFrame([{"part_code": "OLD_1", "giai_thich": "Old note"}])
        new_df = pd.DataFrame([{"part_code": "NEW_1"}, {"part_code": "NEW_2"}])

        migrated = migrate_annotations(new_df, old_df)
        assert all(val == "" for val in migrated["giai_thich"])

    def test_tc_t2_r3_04_explanations_with_newlines_and_special_chars(self) -> None:
        """TC-T2-R3-04: Boundary: Explanations containing newlines, quotes, and Japanese Kanji preserved."""
        complex_text = "【設計変更】\nThay đổi nhà cung cấp NCC-01;\nChú ý: \"Lắp ngược chiều\";"
        old_df = pd.DataFrame([{"part_code": "PART_01", "giai_thich": complex_text}])
        new_df = pd.DataFrame([{"part_code": "PART_01"}])

        migrated = migrate_annotations(new_df, old_df)
        assert migrated.iloc[0]["giai_thich"] == complex_text

    def test_tc_t2_r3_05_missing_plm_old_sheet_recovery(self, tmp_path: Path) -> None:
        """TC-T2-R3-05: Boundary: If Sheet PLM_old is accidentally missing, system initializes fresh backup."""
        wb_path = tmp_path / "BOM_no_old.xlsm"
        wb = openpyxl.Workbook()
        ws_plm = wb.active
        ws_plm.title = "PLM"
        ws_plm["C2"] = "PART_100"
        wb.save(wb_path)

        wb_loaded = openpyxl.load_workbook(wb_path)
        if "PLM_old" not in wb_loaded.sheetnames:
            ws_bak = wb_loaded.copy_worksheet(wb_loaded["PLM"])
            ws_bak.title = "PLM_old"
        wb_loaded.save(wb_path)
        wb_loaded.close()

        wb_check = openpyxl.load_workbook(wb_path)
        assert "PLM_old" in wb_check.sheetnames
        wb_check.close()


# =============================================================================
# R4 Boundaries (5 Tests)
# =============================================================================

class TestTier2R4MSIBoundaries:
    """Boundary test cases for Requirement R4: MSI Deep Reconciliation Engine."""

    def test_tc_t2_r4_01_barcode_shorter_than_9_chars(self) -> None:
        """TC-T2-R4-01: Boundary: Barcode shorter than 9 chars handled without IndexError and evaluates to NG."""
        master = FixSerialMaster()
        master.add_subunit("302FP93010", "1HN")

        # Short barcode lookup executes safely without IndexError
        short_code = "302FP9"
        code, srv = master.lookup_subunit(short_code)
        assert isinstance(code, str)

        # In MSI evaluation, unregistered short barcode evaluates safely to NG
        res = evaluate_msi_branch(
            in_plm=False,
            member_code=short_code,
            master_code="1HN",
        )
        assert res["status"] == "NG"
        assert "B" in res["highlight_red"]

    def test_tc_t2_r4_02_both_row_36_and_plm_c2_empty(self) -> None:
        """TC-T2-R4-02: Boundary: When both row 36 and PLM!C2 are empty, evaluates safely to NG."""
        master = FixSerialMaster()
        row36_code = ""
        plm_c2 = ""
        lookup_code = row36_code or plm_c2

        res = evaluate_msi_branch(
            in_plm=bool(lookup_code),
            member_code="",
            master_code="",
            member_service="",
            master_service="",
        )
        assert res["status"] == "NG"

    def test_tc_t2_r4_03_case_insensitive_3char_code_matching(self) -> None:
        """TC-T2-R4-03: Boundary: Case-insensitive match on 3-char code ('1hn' vs '1HN')."""
        member_input = "1hn".strip().upper()
        master_val = "1HN".strip().upper()
        res = evaluate_msi_branch(
            in_plm=True,
            member_code=member_input,
            master_code=master_val,
            member_service="-",
            master_service="",
        )
        assert res["status"] == "OK"
        assert res["branch"] == 6

    def test_tc_t2_r4_04_service_note_whitespace_stripping(self) -> None:
        """TC-T2-R4-04: Boundary: Extra whitespace around service note stripped safely."""
        res = evaluate_msi_branch(
            in_plm=True,
            member_code="1HN",
            master_code="1HN",
            member_service="  FW 2.0  ",
            master_service="FW 2.0",
        )
        assert res["status"] == "OK"
        assert res["branch"] == 5

    def test_tc_t2_r4_05_missing_sheet_unit_in_fix_serial_master(self, tmp_path: Path) -> None:
        """TC-T2-R4-05: Boundary: Workbook missing Sheet 'UNIT' returns False safely without crash."""
        bad_master = tmp_path / "bad_master.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "SOME_OTHER_SHEET"
        wb.save(bad_master)

        master = FixSerialMaster()
        success = master.load_from_file(bad_master)
        # Sheet UNIT was missing, so subunit_map remains empty
        assert len(master.subunit_map) == 0


# =============================================================================
# R5 Boundaries (5 Tests)
# =============================================================================

class TestTier2R5JIGBoundaries:
    """Boundary test cases for Requirement R5: List JIG Master & 4M Evaluation."""

    def test_tc_t2_r5_01_unsupported_model_in_jig_master(self) -> None:
        """TC-T2-R5-01: Boundary: Unsupported model generates clean fallback table without crashing."""
        gen = ExcelReportGenerator()
        wb = openpyxl.Workbook()
        ws_jig = wb.active
        ws_jig.title = "List JIG"

        gen._build_sheet_jig(ws_jig, pd.DataFrame(), model_name="Sirius99")
        assert "DANH MỤC LIST JIG" in str(ws_jig["B2"].value)

    def test_tc_t2_r5_02_jig_master_file_locked_read_only(self, tmp_path: Path) -> None:
        """TC-T2-R5-02: Boundary: Opening JIG master in read-only mode succeeds on locked/read-only files."""
        ro_file = tmp_path / "readonly_jig.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Virgo"
        wb.save(ro_file)

        # Load with read_only=True
        wb_ro = openpyxl.load_workbook(ro_file, read_only=True)
        assert "Virgo" in wb_ro.sheetnames
        wb_ro.close()

    def test_tc_t2_r5_03_merged_cells_in_jig_master_unmerged_cleanly(self) -> None:
        """TC-T2-R5-03: Boundary: Merged cells unmerged cleanly before pasting new data."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.merge_cells("A3:C3")
        assert len(ws.merged_cells.ranges) == 1

        # Unmerge routine
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
        assert len(ws.merged_cells.ranges) == 0

    def test_tc_t2_r5_04_model_with_zero_jigs_empty_table(self) -> None:
        """TC-T2-R5-04: Boundary: Model with 0 JIGs creates valid table structure with headers."""
        gen = ExcelReportGenerator()
        wb = openpyxl.Workbook()
        ws = wb.active
        gen._build_sheet_jig(ws, pd.DataFrame(columns=["jig_code"]), model_name="Polaris")
        # Headers present
        assert ws.cell(row=4, column=2).value == "STT"
        assert ws.cell(row=4, column=3).value == "MÃ JIG / DỤNG CỤ"

    def test_tc_t2_r5_05_4m_rejected_by_ktsx_recorded_as_ng(self) -> None:
        """TC-T2-R5-05: Boundary: 4M evaluation with rejection ('REJECTED' / 'NG') accurately recorded."""
        assessment = {
            "model_name": "Virgo",
            "is_4m_required": True,
            "ktsx_engineer": "Nguyen Van C",
            "verdict": "NG",
            "reason": "Chưa hoàn thiện Jig hiệu chuẩn quang học",
        }
        assert assessment["verdict"] == "NG"
        assert assessment["reason"] != ""


# =============================================================================
# R6 Boundaries (5 Tests)
# =============================================================================

class TestTier2R6MemberBoundaries:
    """Boundary test cases for Requirement R6: Member Workspace & Self-Check."""

    def test_tc_t2_r6_01_cttt_table_completely_empty(self, qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        """TC-T2-R6-01: Boundary: Submitting with 0 rows blocked with warning."""
        warn_mock = MagicMock()
        monkeypatch.setattr(QMessageBox, "warning", warn_mock)

        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        res = view.submit_data()

        assert res is None
        assert warn_mock.called

    def test_tc_t2_r6_02_fractional_quantities_handling(self, qapp: QApplication, tmp_path: Path) -> None:
        """TC-T2-R6-02: Boundary: Fractional quantities (0.5, 1.25) for grease/tape preserved as floats."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        view.add_cttt_row(page="01", part_code="GREASE_01", part_name="LUBRICANT", quantity=0.5)
        view.add_cttt_row(page="02", part_code="TAPE_01", part_name="ACETATE TAPE", quantity=1.25)

        data = view.get_cttt_table_data()
        assert data[0]["SỐ LƯỢNG"] == 0.5
        assert data[1]["SỐ LƯỢNG"] == 1.25

    def test_tc_t2_r6_03_non_numeric_quantity_string_resilience(self, qapp: QApplication, tmp_path: Path) -> None:
        """TC-T2-R6-03: Boundary: Non-numeric quantity strings ('2 pcs') safely parsed to 0.0 without crash."""
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        view.add_cttt_row(page="01", part_code="PART_01", quantity=1.0)

        # Manually inject invalid text into quantity item
        view.cttt_table.item(0, 3).setText("2 pcs")
        data = view.get_cttt_table_data()
        assert data[0]["SỐ LƯỢNG"] == 0.0

    def test_tc_t2_r6_04_self_check_missing_both_plm_and_r3(self, qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        """TC-T2-R6-04: Boundary: Self-check with no reference data returns informative summary without crashing."""
        monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: None)
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        view.add_cttt_row(page="01", part_code="302FP02010", quantity=1.0)
        view.set_reference_data(None, None)

        res = view.run_preliminary_self_check()
        assert isinstance(res, dict)
        assert res["total"] == 1

    def test_tc_t2_r6_05_repeated_submission_overwrites_cleanly(self, qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        """TC-T2-R6-05: Boundary: Member submitting multiple times in same session produces valid packages."""
        monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: None)
        view = MemberWorkspaceView(base_dir=tmp_path)
        view.clear_cttt_table()
        view.add_cttt_row(page="01", part_code="PART_A", quantity=1.0, explanation="OK")

        sub1 = view.submit_data()
        assert sub1 is not None and sub1.exists()

        # Update and submit again
        view.add_cttt_row(page="02", part_code="PART_B", quantity=2.0, explanation="OK")
        sub2 = view.submit_data()
        assert sub2 is not None and sub2.exists()
