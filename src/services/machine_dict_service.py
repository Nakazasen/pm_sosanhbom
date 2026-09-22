"""Machine dictionary and 2-way sync service with file_loaimay_nhommail.xlsx."""

from __future__ import annotations

import logging
import os
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

logger = logging.getLogger(__name__)

DEFAULT_DICT_PATH = (
    r"\\fstvn01\Data\00_KDTVN Common(KDTVN共通)\⑤Production Engineering(製造技術)"
    r"\Hang muc can luu\Vinh\Pm_sosanhBOM\file_loaimay_nhommail.xlsx"
)


@dataclass
class MachineInfo:
    """Represents machine mapping information from Excel dictionary."""

    machine_name: str           # Cột A: Tên máy (ví dụ: 6th A3, Virgo, Iris 2024...)
    variant: str = ""           # Cột B: Kiểu (ví dụ: Thường, 6th A3S, PRT, MFP...)
    brand_segment: str = ""     # Cột C: Thương hiệu/Phân khúc (ví dụ: KDC, OEM, H, L...)
    machine_codes: list[str] = field(default_factory=list)  # Cột D: Danh sách mã máy 4 ký tự
    specs: str = ""             # Cột E: Thông số kỹ thuật (ví dụ: 250, 500, 40/50/60/70...)
    row_index: int = -1         # 1-indexed row number in Master sheet


@dataclass
class EmailGroupInfo:
    """Represents mail distribution groups from Excel dictionary."""

    group_label: str    # ví dụ: To, CC1(nhóm cơ 1), CC2(nhóm cơ 2)...
    email_string: str   # ví dụ: KDTVN-Production Engineering_Mecha 1(Local) <...>
    clean_email: str    # ví dụ: KDTVN-ProductionEngineering_Mecha1_Local_@kdcf.onmicrosoft.com


class MachineDictService:
    """Service to load, cache, lookup, and write-back machine types and email groups."""

    def __init__(self, excel_path: str | Path | None = None) -> None:
        self.excel_path = Path(excel_path or DEFAULT_DICT_PATH)
        self._machines: list[MachineInfo] = []
        self._code_to_machine: dict[str, MachineInfo] = {}
        self._email_groups: dict[str, EmailGroupInfo] = {}
        self._is_loaded = False

    def load(self, force_reload: bool = False) -> bool:
        """Load and parse the Master sheet from file_loaimay_nhommail.xlsx."""
        if self._is_loaded and not force_reload:
            return True

        if not self.excel_path.exists():
            logger.warning("Machine dictionary file does not exist: %s", self.excel_path)
            return False

        try:
            wb = openpyxl.load_workbook(str(self.excel_path), data_only=True)
            sheet_name = "Master" if "Master" in wb.sheetnames else wb.sheetnames[0]
            ws = wb[sheet_name]

            machines: list[MachineInfo] = []
            code_map: dict[str, MachineInfo] = {}
            email_groups: dict[str, EmailGroupInfo] = {}

            # Read rows starting from row 2 (skipping header)
            for r_idx in range(2, ws.max_row + 1):
                col_a = ws.cell(row=r_idx, column=1).value
                col_b = ws.cell(row=r_idx, column=2).value
                col_c = ws.cell(row=r_idx, column=3).value
                col_d = ws.cell(row=r_idx, column=4).value
                col_e = ws.cell(row=r_idx, column=5).value
                col_f = ws.cell(row=r_idx, column=6).value
                col_g = ws.cell(row=r_idx, column=7).value

                # Parse Email groups (Cols F & G)
                if col_f and col_g:
                    label = str(col_f).strip()
                    raw_email = str(col_g).strip()
                    match = re.search(r"<([^>]+)>", raw_email)
                    clean_email = match.group(1).strip() if match else raw_email
                    email_groups[label] = EmailGroupInfo(
                        group_label=label,
                        email_string=raw_email,
                        clean_email=clean_email,
                    )

                # Parse Machine Info (Cols A - E)
                if col_a or col_d:
                    raw_name = str(col_a).strip() if col_a else ""
                    # Strip all whitespace from machine name: "Iris 2024" -> "Iris2024", "Polaris Next" -> "PolarisNext"
                    machine_name = re.sub(r"\s+", "", raw_name)
                    variant = str(col_b).strip() if col_b else ""
                    brand_segment = str(col_c).strip() if col_c else ""
                    raw_codes = str(col_d).strip() if col_d else ""
                    specs = str(col_e).strip() if col_e else ""

                    codes = [c.strip().upper() for c in raw_codes.split(";") if c.strip()]

                    info = MachineInfo(
                        machine_name=machine_name,
                        variant=variant,
                        brand_segment=brand_segment,
                        machine_codes=codes,
                        specs=specs,
                        row_index=r_idx,
                    )
                    machines.append(info)

                    for code in codes:
                        code_map[code] = info

            wb.close()

            self._machines = machines
            self._code_to_machine = code_map
            self._email_groups = email_groups
            self._is_loaded = True
            logger.info("Loaded %d machines and %d email groups from %s", len(machines), len(email_groups), self.excel_path)
            return True
        except Exception as exc:
            logger.error("Failed to load machine dictionary from %s: %s", self.excel_path, exc)
            return False

    def lookup_machine_code(self, code_4char: str) -> Optional[MachineInfo]:
        """Lookup by 4-character machine code (e.g. '0C0T', '02Y4')."""
        if not self._is_loaded:
            self.load()
        cleaned = code_4char.strip().upper()
        return self._code_to_machine.get(cleaned)

    def extract_and_lookup_material(self, material_code: str) -> Optional[MachineInfo]:
        """Extract machine code from material (e.g. 'T10C0TZUS0' -> '0C0T' or '1102Y43AX0' -> '02Y4').

        T1 + [4-char code] + ... or 11 + [4-char code] + ...
        """
        if not self._is_loaded:
            self.load()

        clean_mat = material_code.strip().upper()

        # Strategy 1: Check prefix T1 or 11
        if len(clean_mat) >= 6:
            candidate_4char = clean_mat[2:6]
            if candidate_4char in self._code_to_machine:
                return self._code_to_machine[candidate_4char]

        # Strategy 2: Search across all known machine codes
        for code, info in self._code_to_machine.items():
            if code in clean_mat:
                return info

        return None

    def get_all_machines(self) -> list[MachineInfo]:
        """Return all loaded machines."""
        if not self._is_loaded:
            self.load()
        return list(self._machines)

    def get_all_email_groups(self) -> dict[str, EmailGroupInfo]:
        """Return all email groups."""
        if not self._is_loaded:
            self.load()
        return dict(self._email_groups)

    def get_model_names(self) -> list[str]:
        """Return sorted list of all unique machine/model names from Excel dictionary.

        Includes all machine models from Column A in file_loaimay_nhommail.xlsx with all
        whitespace removed (e.g. 'Iris 2024' -> 'Iris2024', 'Polaris Next' -> 'PolarisNext'),
        plus baseline models for complete backward compatibility.
        """
        if not self._is_loaded:
            self.load()
        names: set[str] = set()
        for m in self._machines:
            if m.machine_name:
                cleaned = re.sub(r"\s+", "", m.machine_name.strip())
                if cleaned:
                    names.add(cleaned)
        legacy_defaults = ["Virgo", "Libra2", "Iris2024", "Sirius2", "Mebius", "Polaris"]
        for leg in legacy_defaults:
            names.add(re.sub(r"\s+", "", leg))
        return sorted(names, key=lambda s: s.lower())

    def add_or_update_machine_code(
        self,
        machine_name: str,
        new_code: str,
        variant: str = "",
        brand_segment: str = "",
        specs: str = "",
    ) -> bool:
        """2-way write-back: Add a new machine code to file_loaimay_nhommail.xlsx."""
        if not self.excel_path.exists():
            logger.error("Dictionary file does not exist for write-back: %s", self.excel_path)
            return False

        cleaned_code = new_code.strip().upper()
        if not cleaned_code:
            return False

        try:
            # Backup before writing
            backup_path = self.excel_path.with_suffix(".xlsx.bak")
            try:
                shutil.copy2(str(self.excel_path), str(backup_path))
            except Exception as e:
                logger.warning("Could not create backup of dictionary: %s", e)

            wb = openpyxl.load_workbook(str(self.excel_path))
            sheet_name = "Master" if "Master" in wb.sheetnames else wb.sheetnames[0]
            ws = wb[sheet_name]

            target_row = -1
            found_existing = False

            # Search if machine_name already exists in Column A
            clean_search_name = re.sub(r"\s+", "", machine_name).lower()
            for r_idx in range(2, ws.max_row + 1):
                val_a = ws.cell(row=r_idx, column=1).value
                val_b = ws.cell(row=r_idx, column=2).value

                if val_a and re.sub(r"\s+", "", str(val_a)).lower() == clean_search_name:
                    # If variant is specified, try to match variant as well
                    if variant and val_b and str(val_b).strip().lower() != variant.strip().lower():
                        continue

                    target_row = r_idx
                    found_existing = True
                    break

            if found_existing:
                # Append to existing row Column D
                current_codes_val = ws.cell(row=target_row, column=4).value or ""
                current_codes = [c.strip().upper() for c in str(current_codes_val).split(";") if c.strip()]
                if cleaned_code not in current_codes:
                    current_codes.append(cleaned_code)
                    ws.cell(row=target_row, column=4, value="; ".join(current_codes))
            else:
                # Append a new row at the end with whitespace-free machine name
                target_row = ws.max_row + 1
                ws.cell(row=target_row, column=1, value=re.sub(r"\s+", "", machine_name.strip()))
                ws.cell(row=target_row, column=2, value=variant.strip())
                ws.cell(row=target_row, column=3, value=brand_segment.strip())
                ws.cell(row=target_row, column=4, value=cleaned_code)
                ws.cell(row=target_row, column=5, value=specs.strip())

            wb.save(str(self.excel_path))
            wb.close()
            logger.info("Successfully updated machine code %s for machine %s in %s", cleaned_code, machine_name, self.excel_path)

            # Reload internal cache
            self.load(force_reload=True)
            return True
        except Exception as exc:
            logger.error("Failed to write machine code to %s: %s", self.excel_path, exc)
            return False
