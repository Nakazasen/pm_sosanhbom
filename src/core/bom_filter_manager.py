"""BOM Filter Rules Database & Manager Module.

Provides centralized SQLite persistence, offline cache synchronization,
and CRUD operations for PLM BOM decomposition / pruning rules (BolocBom).
"""

from __future__ import annotations

from contextlib import contextmanager
from dataclasses import asdict, dataclass
import datetime
import json
import logging
import os
from pathlib import Path
import random
import re
import shutil
import sqlite3
import time
from typing import Any

from src.core.default_rules import DEFAULT_MODEL_RULES
from src.core.models import ModelRule
from src.core.path_utils import (
    get_app_base_dir,
    is_unc_path_accessible,
    safe_exists,
    safe_is_file,
)

logger = logging.getLogger(__name__)

# Default UNC Network path requested by User
DEFAULT_SHARED_DB_PATH = (
    r"\\fstvn01\Data\00_KDTVN Common(KDTVN共通)\⑤Production Engineering(製造技術)"
    r"\Hang muc can luu\Vinh\Pm_sosanhBOM\ssbom_master.db"
)


@dataclass
class BOMRuleRecord:
    """Represents a single PLM BOM filtering rule for a machine model."""

    id: int | None
    model_name: str
    item_name: str | None
    match_mode: str = "Full_name"  # 'Full_name' | 'Part_name'
    part_code: str | None = None
    notes: str = ""
    is_active: bool = True
    created_at: str | None = None
    updated_at: str | None = None

    def to_model_rule(self) -> ModelRule:
        """Convert to ModelRule domain entity used by ModelPruner."""
        action = None
        parent_part_code = None
        if self.notes:
            n_low = self.notes.lower()
            if "loại bỏ cả cụm" in n_low or "prune_node" in n_low:
                action = "prune_node"
            elif "cắt con" in n_low or "prune_children" in n_low:
                action = "prune_children"
            m = re.search(r"\[branch:\s*([^\]]+)\]", self.notes)
            if m:
                parent_part_code = m.group(1).strip()
        return ModelRule(
            item_name=self.item_name,
            match_mode=self.match_mode or "Full_name",
            part_code=self.part_code,
            action=action,
            parent_part_code=parent_part_code,
        )

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class BOMFilterManager:
    """Manages BOM pruning rules in SQLite with LAN UNC sharing and offline fallback."""

    def __init__(
        self,
        base_dir: Path | str | None = None,
        remote_db_path: Path | str | None = None,
    ) -> None:
        self.base_dir = Path(base_dir) if base_dir else get_app_base_dir()
        self._remote_check_cache: bool | None = None
        self._remote_check_time: float = 0.0

        if remote_db_path:
            self.remote_db_path = Path(remote_db_path)
        else:
            local_master = self.base_dir / "ssbom_master.db"
            cfg_path = self.base_dir / "config" / "settings.json"
            cfg_shared = ""
            if safe_exists(cfg_path):
                try:
                    with open(cfg_path, "r", encoding="utf-8") as fp:
                        data = json.load(fp)
                        cfg_shared = data.get("paths", {}).get("shared_db_path", "")
                except Exception:
                    pass
            is_test_dir = any(t in str(self.base_dir).lower() for t in ("pytest", "temp", "tmp"))
            if cfg_shared:
                self.remote_db_path = Path(cfg_shared)
            else:
                local_master = self.base_dir / "ssbom_master.db"
                if is_test_dir or safe_exists(local_master):
                    self.remote_db_path = local_master
                else:
                    self.remote_db_path = Path(DEFAULT_SHARED_DB_PATH)

        self.is_test_dir = any(t in str(self.base_dir).lower() for t in ("pytest", "temp", "tmp"))
        self._last_cache_sync_time: float = 0.0

        self.local_cache_dir = self.base_dir / "data"
        self.local_cache_dir.mkdir(parents=True, exist_ok=True)
        self.local_cache_path = self.local_cache_dir / "ssbom_master_cache.db"

        # Initialize schema and seed default data
        self.initialize_database()

    def is_remote_available(self) -> bool:
        """Check if remote UNC share or directory is accessible with caching."""
        now = time.time()
        if self._remote_check_cache is not None and (now - self._remote_check_time < 15.0):
            return self._remote_check_cache

        is_avail = False
        try:
            parent_dir = self.remote_db_path.parent
            if is_unc_path_accessible(parent_dir, timeout=0.5):
                is_avail = safe_exists(parent_dir)
        except Exception:
            is_avail = False

        self._remote_check_cache = is_avail
        self._remote_check_time = now
        return is_avail

    def _ensure_cache_fresh(self, force: bool = False) -> None:
        """Ensure local cache is up-to-date with remote DB using throttled mtime check."""
        now = time.time()
        throttle = 0.0 if self.is_test_dir else 5.0
        if not force and (now - self._last_cache_sync_time < throttle):
            return
        self._last_cache_sync_time = now

        if not self.is_remote_available():
            return

        remote_path = self.remote_db_path
        if remote_path == self.local_cache_path or not safe_exists(remote_path):
            return

        try:
            if not safe_exists(self.local_cache_path):
                self._safe_backup_db(remote_path, self.local_cache_path)
                return

            r_mtime = remote_path.stat().st_mtime
            l_mtime = self.local_cache_path.stat().st_mtime
            if r_mtime > l_mtime:
                logger.debug("Remote BOM DB is newer (%.1f > %.1f), syncing cache...", r_mtime, l_mtime)
                self._safe_backup_db(remote_path, self.local_cache_path)
        except Exception as exc:
            logger.debug("Error checking BOM cache freshness: %s", exc)

    def get_active_db_path(self) -> tuple[Path, bool]:
        """Determine whether to use remote network database or local cache.

        Returns:
            (active_path, is_remote_connected)
        """
        if self.is_remote_available():
            return self.remote_db_path, True
        return self.local_cache_path, False

    def get_read_db_path(self, force_refresh: bool = False) -> Path:
        """Return local cache path for local-first zero-latency read access."""
        self._ensure_cache_fresh(force=force_refresh)
        if safe_exists(self.local_cache_path):
            return self.local_cache_path
        active_path, _ = self.get_active_db_path()
        return active_path

    @staticmethod
    def _safe_backup_db(src_path: Path, dst_path: Path) -> None:
        """Safely copy live SQLite database using sqlite3 online backup API."""
        if not safe_exists(src_path):
            return
        src_conn = None
        dst_conn = None
        try:
            dst_path.parent.mkdir(parents=True, exist_ok=True)
            src_conn = sqlite3.connect(str(src_path), timeout=10.0)
            dst_conn = sqlite3.connect(str(dst_path), timeout=10.0)
            src_conn.backup(dst_conn)
        except Exception as exc:
            logger.debug("Safe SQLite backup skipped: %s", exc)
        finally:
            if src_conn is not None:
                try:
                    src_conn.close()
                except Exception:
                    pass
            if dst_conn is not None:
                try:
                    dst_conn.close()
                except Exception:
                    pass

    @contextmanager
    def _get_connection(self, db_path: Path):
        """Open an SQLite connection with 30s busy timeout and clean closure."""
        db_path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(str(db_path), timeout=30.0)
        conn.row_factory = sqlite3.Row
        try:
            conn.execute("PRAGMA busy_timeout = 30000;")
            yield conn
        finally:
            try:
                conn.close()
            except Exception:
                pass

    @contextmanager
    def _write_transaction(self, db_path: Path):
        """Open an SQLite write transaction using BEGIN IMMEDIATE to avoid upgrade deadlocks."""
        db_path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(str(db_path), timeout=30.0, isolation_level=None)
        conn.row_factory = sqlite3.Row
        try:
            conn.execute("PRAGMA busy_timeout = 30000;")
            conn.execute("BEGIN IMMEDIATE;")
            yield conn
            conn.execute("COMMIT;")
        except Exception:
            try:
                conn.execute("ROLLBACK;")
            except Exception:
                pass
            raise
        finally:
            try:
                conn.close()
            except Exception:
                pass

    def _execute_dual_write(self, write_action, operation_name: str = "ghi quy tắc BOM") -> tuple[bool, Any, bool]:
        """Execute write_action(conn) on remote (with jitter retry) and immediately on local cache.

        write_action: callable(sqlite3.Connection) -> result (e.g. lastrowid, rowcount, or custom).
        Returns:
            (success: bool, result_or_error_msg: Any, used_fallback: bool)
        """
        remote_path, is_remote = self.get_active_db_path()
        used_fallback = False
        result = None

        # 1. Write to remote if available
        if is_remote and remote_path != self.local_cache_path:
            max_retries = 5
            remote_success = False
            for attempt in range(max_retries):
                try:
                    with self._write_transaction(remote_path) as conn:
                        result = write_action(conn)
                    remote_success = True
                    break
                except sqlite3.IntegrityError as int_err:
                    return False, str(int_err), False
                except LookupError as lk_err:
                    return False, str(lk_err), False
                except sqlite3.OperationalError as op_err:
                    err_str = str(op_err).lower()
                    if ("locked" in err_str or "busy" in err_str) and attempt < max_retries - 1:
                        delay = 0.04 * (1.5 ** attempt) + random.uniform(0.01, 0.06)
                        logger.warning(
                            "Remote DB locked during %s, retrying in %.3fs (attempt %d/%d)...",
                            operation_name,
                            delay,
                            attempt + 1,
                            max_retries,
                        )
                        time.sleep(delay)
                        continue
                    logger.warning("Remote DB write failed after retries: %s; falling back to local cache.", op_err)
                    used_fallback = True
                    break
                except Exception as ex:
                    logger.warning("Remote write error: %s; falling back to local cache.", ex)
                    used_fallback = True
                    break

        # 2. Dual-write directly to local cache
        try:
            with self._write_transaction(self.local_cache_path) as conn:
                local_res = write_action(conn)
                if result is None:
                    result = local_res
        except sqlite3.IntegrityError as int_err:
            return False, str(int_err), used_fallback
        except LookupError as lk_err:
            return False, str(lk_err), used_fallback
        except Exception as local_err:
            logger.error("Local cache dual-write error: %s", local_err)
            if used_fallback and result is None:
                return False, f"Lỗi ghi dữ liệu cục bộ: {local_err}", used_fallback

        self._last_cache_sync_time = time.time()
        return True, result, used_fallback

    def _sync_to_local_cache(self, source_path: Path) -> None:
        """Sync remote database file to local cache for offline resiliency."""
        try:
            if source_path != self.local_cache_path and safe_exists(source_path):
                should_sync = True
                if safe_exists(self.local_cache_path):
                    try:
                        should_sync = source_path.stat().st_mtime > self.local_cache_path.stat().st_mtime
                    except OSError:
                        should_sync = False
                if should_sync:
                    self._safe_backup_db(source_path, self.local_cache_path)
        except Exception as ex:
            logger.warning("Could not sync remote DB to local cache: %s", ex)

    def initialize_database(self) -> None:
        """Create tables if not existing, seed default 357 rules, and sync cache."""
        target_path, is_remote = self.get_active_db_path()

        try:
            with self._get_connection(target_path) as conn:
                is_network = str(target_path).startswith(r"\\") or str(target_path).startswith("//")
                try:
                    if is_network:
                        conn.execute("PRAGMA journal_mode = TRUNCATE;")
                    else:
                        conn.execute("PRAGMA journal_mode = WAL;")
                except sqlite3.OperationalError:
                    pass

                cursor = conn.cursor()
                cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='bolocbom_rules';")
                table_exists = cursor.fetchone() is not None

                if not table_exists:
                    conn.execute(
                        """
                        CREATE TABLE IF NOT EXISTS bolocbom_rules (
                            id INTEGER PRIMARY KEY AUTOINCREMENT,
                            model_name TEXT NOT NULL,
                            item_name TEXT,
                            match_mode TEXT DEFAULT 'Full_name',
                            part_code TEXT,
                            notes TEXT DEFAULT '',
                            is_active INTEGER DEFAULT 1,
                            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                        );
                        """
                    )
                    conn.execute(
                        "CREATE INDEX IF NOT EXISTS idx_bolocbom_model ON bolocbom_rules(model_name);"
                    )
                    conn.execute(
                        "CREATE INDEX IF NOT EXISTS idx_bolocbom_active ON bolocbom_rules(is_active);"
                    )

                    # Check if bolocbom_rules has existing rows
                    cursor.execute("SELECT COUNT(*) FROM bolocbom_rules;")
                    count = cursor.fetchone()[0]

                    if count == 0:
                        self._seed_default_rules(conn)
                        conn.commit()

            if is_remote:
                self._sync_to_local_cache(target_path)

        except Exception as ex:
            logger.error("Failed to initialize BOM Filter database: %s", ex)
            if is_remote:
                logger.info("Falling back to local cache initialization.")
                try:
                    with self._get_connection(self.local_cache_path) as conn:
                        cursor = conn.cursor()
                        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='bolocbom_rules';")
                        if not cursor.fetchone():
                            conn.execute(
                                """
                                CREATE TABLE IF NOT EXISTS bolocbom_rules (
                                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                                    model_name TEXT NOT NULL,
                                    item_name TEXT,
                                    match_mode TEXT DEFAULT 'Full_name',
                                    part_code TEXT,
                                    notes TEXT DEFAULT '',
                                    is_active INTEGER DEFAULT 1,
                                    created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                                    updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                                );
                                """
                            )
                            conn.execute("CREATE INDEX IF NOT EXISTS idx_bolocbom_model ON bolocbom_rules(model_name);")
                            cursor = conn.execute("SELECT COUNT(*) FROM bolocbom_rules;")
                            if cursor.fetchone()[0] == 0:
                                self._seed_default_rules(conn)
                                conn.commit()
                except Exception as local_ex:
                    logger.error("Failed to initialize local cache BOM Filter: %s", local_ex)

    def _seed_default_rules(self, conn: sqlite3.Connection) -> None:
        """Seed all 357 rules from DEFAULT_MODEL_RULES into database."""
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        records = []
        for model_name, rule_dicts in DEFAULT_MODEL_RULES.items():
            for r in rule_dicts:
                records.append(
                    (
                        model_name,
                        r.get("item_name"),
                        r.get("match_mode", "Full_name"),
                        r.get("part_code"),
                        "Quy tắc mặc định ban đầu",
                        1,
                        now,
                        now,
                    )
                )

        conn.executemany(
            """
            INSERT INTO bolocbom_rules (model_name, item_name, match_mode, part_code, notes, is_active, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?);
            """,
            records,
        )
        logger.info("Seeded %d default BOM filter rules into database.", len(records))

    def get_all_models(self) -> list[str]:
        """Return list of distinct machine model names registered in the rules table (Local-First fast read)."""
        target_path = self.get_read_db_path()
        try:
            with self._get_connection(target_path) as conn:
                cursor = conn.execute(
                    "SELECT DISTINCT model_name FROM bolocbom_rules ORDER BY model_name COLLATE NOCASE;"
                )
                models = sorted(list({re.sub(r"\s+", "", row[0]) for row in cursor.fetchall() if row[0]}), key=lambda s: s.lower())
                return models
        except Exception as ex:
            logger.error("Error fetching models from DB: %s", ex)
            return [re.sub(r"\s+", "", k) for k in DEFAULT_MODEL_RULES.keys()]

    def get_rules_for_model(self, model_name: str, active_only: bool = True) -> list[BOMRuleRecord]:
        """Fetch all BOM filter rules for a given machine model (Local-First fast read)."""
        target_path = self.get_read_db_path()
        if not model_name:
            return []

        clean_name = re.sub(r"\s+", "", model_name.strip())
        try:
            with self._get_connection(target_path) as conn:
                query = "SELECT * FROM bolocbom_rules WHERE REPLACE(LOWER(model_name), ' ', '') = LOWER(?)"
                params: list[Any] = [clean_name]
                if active_only:
                    query += " AND is_active = 1"
                query += " ORDER BY id ASC;"

                cursor = conn.execute(query, params)
                rows = cursor.fetchall()

                # If no direct match, check if prefix matches
                if not rows:
                    query_prefix = "SELECT * FROM bolocbom_rules WHERE LOWER(?) LIKE REPLACE(LOWER(model_name), ' ', '') || '%'"
                    if active_only:
                        query_prefix += " AND is_active = 1"
                    query_prefix += " ORDER BY id ASC;"
                    cursor = conn.execute(query_prefix, [clean_name])
                    rows = cursor.fetchall()

                results = []
                for r in rows:
                    results.append(
                        BOMRuleRecord(
                            id=r["id"],
                            model_name=r["model_name"],
                            item_name=r["item_name"],
                            match_mode=r["match_mode"],
                            part_code=r["part_code"],
                            notes=r["notes"],
                            is_active=bool(r["is_active"]),
                            created_at=r["created_at"],
                            updated_at=r["updated_at"],
                        )
                    )
                return results
        except Exception as ex:
            logger.error("Error fetching rules for model %s: %s", model_name, ex)
            return []

    def get_model_rules_for_pruner(self, model_name: str) -> list[ModelRule]:
        """Fetch rules converted directly to ModelRule entities for ModelPruner."""
        records = self.get_rules_for_model(model_name, active_only=True)
        if records:
            return [r.to_model_rule() for r in records]
        return []

    def add_rule(
        self,
        model_name: str,
        item_name: str | None,
        match_mode: str = "Full_name",
        part_code: str | None = None,
        notes: str = "",
    ) -> int:
        """Add a new BOM filtering rule via fast dual-write micro-transaction."""
        clean_model = model_name.strip()
        clean_item = item_name.strip() if item_name else None
        clean_code = part_code.strip() if part_code else None
        clean_mode = "Part_name" if match_mode == "Part_name" else "Full_name"
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        def _do_add(conn: sqlite3.Connection):
            cursor = conn.execute(
                """
                INSERT INTO bolocbom_rules (model_name, item_name, match_mode, part_code, notes, is_active, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, 1, ?, ?);
                """,
                (clean_model, clean_item, clean_mode, clean_code, notes.strip(), now, now),
            )
            return cursor.lastrowid or 0

        ok, result, _ = self._execute_dual_write(_do_add, operation_name=f"thêm quy tắc '{clean_model}'")
        return result if ok and isinstance(result, int) else 0

    def update_rule(
        self,
        rule_id: int,
        item_name: str | None,
        match_mode: str = "Full_name",
        part_code: str | None = None,
        notes: str = "",
    ) -> bool:
        """Update an existing BOM filtering rule via fast dual-write micro-transaction."""
        clean_item = item_name.strip() if item_name else None
        clean_code = part_code.strip() if part_code else None
        clean_mode = "Part_name" if match_mode == "Part_name" else "Full_name"
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        def _do_update(conn: sqlite3.Connection):
            cursor = conn.execute(
                """
                UPDATE bolocbom_rules
                SET item_name = ?, match_mode = ?, part_code = ?, notes = ?, updated_at = ?
                WHERE id = ?;
                """,
                (clean_item, clean_mode, clean_code, notes.strip(), now, rule_id),
            )
            if cursor.rowcount == 0:
                raise LookupError(f"Không tìm thấy quy tắc {rule_id}")
            return cursor.rowcount > 0

        ok, result, _ = self._execute_dual_write(_do_update, operation_name=f"cập nhật quy tắc ID {rule_id}")
        return bool(ok and result)

    def delete_rule(self, rule_id: int) -> bool:
        """Delete a rule by ID via fast dual-write micro-transaction."""
        def _do_del(conn: sqlite3.Connection):
            cursor = conn.execute("DELETE FROM bolocbom_rules WHERE id = ?;", (rule_id,))
            if cursor.rowcount == 0:
                raise LookupError(f"Không tìm thấy quy tắc {rule_id}")
            return cursor.rowcount > 0

        ok, result, _ = self._execute_dual_write(_do_del, operation_name=f"xóa quy tắc ID {rule_id}")
        return bool(ok and result)

    def delete_model(self, model_name: str) -> bool:
        """Delete all rules associated with a machine model via fast dual-write micro-transaction."""
        def _do_del_model(conn: sqlite3.Connection):
            cursor = conn.execute("DELETE FROM bolocbom_rules WHERE LOWER(model_name) = LOWER(?);", (model_name.strip(),))
            return cursor.rowcount > 0

        ok, result, _ = self._execute_dual_write(_do_del_model, operation_name=f"xóa model '{model_name}'")
        return bool(ok and result)

    def reset_model_to_default(self, model_name: str) -> int:
        """Reset a model's rules to the factory default rules via fast dual-write micro-transaction."""
        clean = model_name.strip().lower()
        matched_model = None
        for def_model in DEFAULT_MODEL_RULES:
            if def_model.lower() == clean:
                matched_model = def_model
                break

        if not matched_model:
            return 0

        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        records = [
            (
                matched_model,
                r.get("item_name"),
                r.get("match_mode", "Full_name"),
                r.get("part_code"),
                "Khôi phục mặc định ban đầu",
                1,
                now,
                now,
            )
            for r in DEFAULT_MODEL_RULES[matched_model]
        ]

        def _do_reset_model(conn: sqlite3.Connection):
            conn.execute("DELETE FROM bolocbom_rules WHERE LOWER(model_name) = LOWER(?);", (clean,))
            conn.executemany(
                """
                INSERT INTO bolocbom_rules (model_name, item_name, match_mode, part_code, notes, is_active, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?);
                """,
                records,
            )
            return len(records)

        ok, result, _ = self._execute_dual_write(_do_reset_model, operation_name=f"khôi phục model '{model_name}'")
        return result if ok and isinstance(result, int) else 0

    def reset_all_to_default(self) -> int:
        """Clear table and re-seed all 357 default rules via fast dual-write micro-transaction."""
        def _do_reset_all(conn: sqlite3.Connection):
            conn.execute("DELETE FROM bolocbom_rules;")
            self._seed_default_rules(conn)
            return sum(len(rules) for rules in DEFAULT_MODEL_RULES.values())

        ok, result, _ = self._execute_dual_write(_do_reset_all, operation_name="khôi phục tất cả quy tắc")
        return result if ok and isinstance(result, int) else 0

    def import_from_excel(self, file_path: Path | str, sheet_name: str = "BolocBom") -> int:
        """Import BOM filtering rules from an Excel workbook (e.g. form_ssbom.xlsm).

        Expected columns:
        - Col A: Model Name
        - Col B: Item Name
        - Col C: Match Mode (Full_name / Part_name)
        - Col D: Part Code
        - Col E (optional): Notes
        """
        import openpyxl

        p = Path(file_path)
        if not p.exists():
            raise FileNotFoundError(f"File không tồn tại: {p}")

        wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
        try:
            ws = None
            for s in wb.sheetnames:
                if s.strip().lower() == sheet_name.strip().lower():
                    ws = wb[s]
                    break
            if ws is None:
                ws = wb.active

            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            records_to_insert = []

            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_idx == 1:
                    # Check if row 1 is a header (e.g. contains 'model', 'loại máy', etc.)
                    val0 = str(row[0] or "").strip().lower()
                    if any(k in val0 for k in ["model", "loại", "máy", "stt", "tên"]):
                        continue

                if not row or len(row) < 2:
                    continue

                model = str(row[0] or "").strip()
                item_name = str(row[1] or "").strip() if row[1] is not None else None
                match_mode = str(row[2] or "Full_name").strip() if len(row) > 2 and row[2] else "Full_name"
                part_code = str(row[3] or "").strip() if len(row) > 3 and row[3] else None
                notes = str(row[4] or "").strip() if len(row) > 4 and row[4] else "Nhập từ Excel"

                if not model or (not item_name and not part_code):
                    continue

                clean_mode = "Part_name" if "part" in match_mode.lower() else "Full_name"
                records_to_insert.append(
                    (model, item_name, clean_mode, part_code, notes, 1, now, now)
                )

            if not records_to_insert:
                return 0

            target_path, is_remote = self.get_active_db_path()
            with self._write_transaction(target_path) as conn:
                conn.executemany(
                    """
                    INSERT INTO bolocbom_rules (model_name, item_name, match_mode, part_code, notes, is_active, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?);
                    """,
                    records_to_insert,
                )

            if is_remote:
                self._sync_to_local_cache(target_path)

            return len(records_to_insert)
        finally:
            wb.close()

    def export_to_excel(self, output_path: Path | str, model_name: str | None = None) -> Path:
        """Export rules to an Excel file with standard 'BolocBom' sheet format."""
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BolocBom"

        # Headers
        headers = ["Model", "Item Name", "Match Mode", "Part Code", "Ghi Chú"]
        ws.append(headers)

        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3D59", end_color="1E3D59", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Fetch records
        target_path, _ = self.get_active_db_path()
        with self._get_connection(target_path) as conn:
            if model_name:
                cursor = conn.execute(
                    "SELECT model_name, item_name, match_mode, part_code, notes FROM bolocbom_rules WHERE LOWER(model_name) = LOWER(?) ORDER BY id ASC;",
                    (model_name.strip(),),
                )
            else:
                cursor = conn.execute(
                    "SELECT model_name, item_name, match_mode, part_code, notes FROM bolocbom_rules ORDER BY model_name, id ASC;"
                )
            rows = cursor.fetchall()

        data_font = Font(name="Segoe UI", size=10)
        for r_idx, r in enumerate(rows, start=2):
            ws.append([r[0], r[1] or "", r[2] or "Full_name", r[3] or "", r[4] or ""])
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.font = data_font
                c.border = thin_border
                if c_idx in (1, 3, 4):
                    c.alignment = Alignment(horizontal="center", vertical="center")

        # Column widths
        widths = [16, 42, 16, 18, 30]
        for idx, w in enumerate(widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = w

        wb.save(str(out_p))
        wb.close()
        return out_p


_bom_filter_instances: dict[str, BOMFilterManager] = {}


def get_bom_filter_manager(
    base_dir: Path | str | None = None,
    remote_db_path: Path | str | None = None,
) -> BOMFilterManager:
    """Return a cached or new BOMFilterManager instance."""
    key = f"{base_dir}_{remote_db_path}"
    if key not in _bom_filter_instances:
        _bom_filter_instances[key] = BOMFilterManager(base_dir, remote_db_path)
    return _bom_filter_instances[key]
