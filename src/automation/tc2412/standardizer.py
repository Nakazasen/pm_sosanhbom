"""TC2412 BOM Canonical 14-Column Standardizer.

Standardizes raw 24-column Teamcenter 2412 Excel BOM exports into the canonical
14-column layout expected by downstream BOM reconciliation tools and legacy Excel macros.
Columns A-N:
  Col A (1): Home (1-based row index)
  Col B (2): Level (integer)
  Col C (3): Item Type ('Parts')
  Col D (4): Item Id (Part Number)
  Col E (5): Has Children (None)
  Col F (6): Quantity (formatted string, e.g. '1.000')
  Col G (7): 1st Parts (None)
  Col H (8): 2nd BOM Flag ('False' / 'True')
  Col I (9): Occurrence Effectivities (None)
  Col J (10): Item Revision Project List (None)
  Col K (11): Item Name (Parts Text)
  Col L (12): Notice No (ECN Notice No)
  Col M (13): Revision (e.g. '01')
  Col N (14): Item Rev Status ('Released')
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CANONICAL_14_COLUMNS: List[str] = [
    "Home",
    "Level",
    "Item Type",
    "Item Id",
    "Has Children",
    "Quantity",
    "1st Parts",
    "2nd BOM Flag",
    "Occurrence Effectivities",
    "Item Revision Project List",
    "Item Name",
    "Notice No",
    "Revision",
    "Item Rev Status",
]

CANONICAL_14_WIDTHS: dict[str, float] = {
    "A": 5.7,   # Home
    "B": 19.7,  # Level
    "C": 13.0,  # Item Type
    "D": 15.7,  # Item Id
    "E": 13.0,  # Has Children
    "F": 11.7,  # Quantity
    "G": 18.7,  # 1st Parts
    "H": 20.7,  # 2nd BOM Flag
    "I": 14.8,  # Occurrence Effectivities
    "J": 19.7,  # Item Revision Project List
    "K": 40.7,  # Item Name
    "L": 29.7,  # Notice No
    "M": 23.7,  # Revision
    "N": 26.7,  # Item Rev Status
}

THIN_SIDE = Side(style="thin", color="000000")
ALL_BORDER = Border(top=THIN_SIDE, bottom=THIN_SIDE, left=THIN_SIDE, right=THIN_SIDE)

HEADER_FONT = Font(name="MS Gothic", size=11, bold=True)
HEADER_FILL = PatternFill(fill_type="solid", start_color="C0C0C0", end_color="C0C0C0")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
DATA_FONT = Font(name="MS Gothic", size=11, bold=False)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)


def prune_mechanical_scope(raw_rows: List[List[Any]]) -> List[List[Any]]:
    """Prune electrical units and unneeded accessory sub-levels for PE Dept 2 (Mechanical).

    - Retains Root (Level 0) and all mechanical assemblies under SET ASSY ELEMENTS (DLP, DRIVE, FRAME, etc.)
    - Excludes ELEC UNIT (Level 2 unit handled exclusively by PE Dept 4 / Electrical) and all its sub-children
    - For ACCESSORIES (Level 1): retains Level 1 ACCESSORIES and Level 2 units (FILL UP CONTAINER ASSY,
      WASTE BOTTLE ASSY, AC CORD ASSY) while collapsing their lower-level sub-components (Level >= 3).
    """
    if len(raw_rows) <= 2:
        return raw_rows

    res: List[List[Any]] = [raw_rows[0], raw_rows[1]]

    in_elec = False
    in_accessories = False

    for idx in range(2, len(raw_rows)):
        r = raw_rows[idx]
        try:
            lvl = int(r[1] if len(r) > 1 and r[1] is not None else 0)
        except Exception:
            lvl = 0

        pid = str(r[3] if len(r) > 3 and r[3] is not None else "").strip()
        pname = str(r[7] if len(r) > 7 and r[7] is not None else "").strip()

        # Reset state at Level 1 or Level 0
        if lvl <= 1:
            in_elec = False
            in_accessories = ("accessories" in pname.lower())
            res.append(r)
            continue

        # If inside accessories, only keep Level 2 units
        if in_accessories:
            if lvl == 2:
                res.append(r)
            continue

        # At Level 2 (under SET ASSY ELEMENTS):
        if lvl == 2:
            if "elec unit" in pname.lower() or ("3vc" in pid.lower() and "58500" in pid.lower()):
                in_elec = True
                continue
            else:
                in_elec = False

        if in_elec:
            continue

        res.append(r)

    return res


def transform_24_to_14_columns(
    raw_rows: List[List[Any]],
    prune_electrical: bool = True,
) -> List[List[Any]]:
    """Transform raw 24-column rows into canonical 14-column format.

    Maps:
      Home -> 1-based sequential row index
      Level -> Level.1 (Col 2)
      Item Type -> Item Type (Col 3, e.g. 'Parts')
      Item Id -> Name (Col 4, Part Number)
      Has Children -> None
      Quantity -> Quantity (Col 6)
      1st Parts -> None
      2nd BOM Flag -> 2nd BOM Flag (Col 7)
      Occurrence Effectivities -> None
      Item Revision Project List -> None
      Item Name -> Parts Text (Col 8)
      Notice No -> Notice No (Col 9)
      Revision -> Revision (Col 10)
      Item Rev Status -> Release Status (Col 11)
    """
    if not raw_rows:
        return []

    source_rows = prune_mechanical_scope(raw_rows) if prune_electrical else raw_rows

    res: List[List[Any]] = [CANONICAL_14_COLUMNS]

    for idx in range(1, len(source_rows)):
        r = source_rows[idx]
        home_idx = idx

        # Level
        try:
            lvl_val = int(r[1] if len(r) > 1 and r[1] is not None else 0)
        except Exception:
            lvl_val = 0

        item_type = str(r[2]).strip() if len(r) > 2 and r[2] is not None else "Parts"
        item_id = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""

        # Quantity
        qty_raw = r[5] if len(r) > 5 else None
        if qty_raw is not None and str(qty_raw).strip() != "":
            try:
                # Format as string with 3 decimals if float, or keep string
                qty_str = f"{float(qty_raw):.3f}"
            except Exception:
                qty_str = str(qty_raw).strip()
        else:
            qty_str = None

        # 2nd BOM Flag
        flag_raw = r[6] if len(r) > 6 else None
        flag_str = str(flag_raw).strip() if flag_raw is not None and str(flag_raw).strip() != "" else None

        # Item Name (Parts Text)
        name_raw = r[7] if len(r) > 7 else None
        name_str = str(name_raw).strip() if name_raw is not None else ""

        # Notice No
        notice_raw = r[8] if len(r) > 8 else None
        notice_str = str(notice_raw).strip() if notice_raw is not None and str(notice_raw).strip() != "" else None

        # Revision
        rev_raw = r[9] if len(r) > 9 else None
        if rev_raw is not None and str(rev_raw).strip() != "":
            rev_str = str(rev_raw).strip()
            if len(rev_str) == 1 and rev_str.isdigit():
                rev_str = f"0{rev_str}"
        else:
            rev_str = ""

        # Release Status
        status_raw = r[10] if len(r) > 10 else None
        status_str = str(status_raw).strip() if status_raw is not None and str(status_raw).strip() != "" else None

        row_14 = [
            home_idx,        # Home (1, 2, 3...)
            lvl_val,         # Level (0, 1, 2...)
            item_type,       # Item Type
            item_id,         # Item Id
            None,            # Has Children
            qty_str,         # Quantity
            None,            # 1st Parts
            flag_str,        # 2nd BOM Flag
            None,            # Occurrence Effectivities
            None,            # Item Revision Project List
            name_str,        # Item Name
            notice_str,      # Notice No
            rev_str,         # Revision
            status_str,      # Item Rev Status
        ]
        res.append(row_14)

    return res


def write_14_column_workbook(rows_14: List[List[Any]], output_path: Path) -> Path:
    """Save 14-column data with authentic formatting, fonts, and column widths."""
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.views.sheetView[0].showGridLines = True

    for r in rows_14:
        ws.append(r)

    # Style header (Row 1)
    for col_idx in range(1, 15):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = ALL_BORDER

    # Set column widths
    for col_letter, width in CANONICAL_14_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # Style data cells (Row 2 to max_row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, 15):
            cell = ws.cell(row=r, column=c)
            cell.font = DATA_FONT
            cell.border = ALL_BORDER

            # Alignments: Item Name (11) and Notice No (12) are left-aligned; others centered
            if c in (11, 12):
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_CENTER

            # Number formats: Col 1 & 2 General, Col 3 to 14 Text (@)
            if c in (1, 2):
                cell.number_format = "General"
            else:
                cell.number_format = "@"

    wb.save(output_path)
    wb.close()
    return output_path


def standardize_plm_file(
    src_file: Path,
    dest_file: Path,
    prune_electrical: bool = True,
) -> Tuple[int, int]:
    """Standardize a single raw 24-column PLM BOM Excel file into 14-column format.

    Returns:
        (original_row_count, final_row_count)
    """
    src_file = Path(src_file).resolve()
    dest_file = Path(dest_file).resolve()

    wb_in = openpyxl.load_workbook(src_file, data_only=True)
    ws_in = wb_in.active
    raw_rows = [list(row) for row in ws_in.iter_rows(values_only=True)]
    wb_in.close()

    rows_14 = transform_24_to_14_columns(raw_rows, prune_electrical=prune_electrical)
    write_14_column_workbook(rows_14, dest_file)

    orig_count = len(raw_rows)
    final_count = len(rows_14)
    return orig_count, final_count
