"""Excel report generation service matching legacy form_ssbom.xlsm layout."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)


class ExcelReportExporter:
    """Exports BOM reconciliation results to standard Excel format."""

    def __init__(self) -> None:
        self.green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        self.green_font = Font(name="Calibri", size=10, bold=True, color="006100")

        self.red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.red_font = Font(name="Calibri", size=10, bold=True, color="9C0006")

        self.yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        self.yellow_font = Font(name="Calibri", size=10, bold=True, color="9C6500")

        self.header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        thin_side = Side(style="thin", color="D9D9D9")
        self.thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    def export(
        self,
        output_path: Path,
        summary_data: Dict[str, Any],
        cttt_df: pd.DataFrame,
        msi_df: Optional[pd.DataFrame] = None,
        missing_df: Optional[pd.DataFrame] = None,
    ) -> Path:
        """Create structured Excel workbook with TongHop, CTTT, MSI, and Thieu_PLM sheets."""
        output_path = Path(output_path).resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()

        # Sheet 1: TongHop
        ws_tonghop = wb.active
        ws_tonghop.title = "TongHop"
        self._build_tonghop_sheet(ws_tonghop, summary_data)

        # Sheet 2: CTTT
        ws_cttt = wb.create_sheet(title="CTTT")
        self._build_table_sheet(ws_cttt, cttt_df, is_cttt=True)

        # Sheet 3: MSI
        if msi_df is not None and not msi_df.empty:
            ws_msi = wb.create_sheet(title="MSI")
            self._build_table_sheet(ws_msi, msi_df, is_cttt=False)

        # Sheet 4: Thieu_PLM
        if missing_df is not None and not missing_df.empty:
            ws_thieu = wb.create_sheet(title="Thieu_PLM")
            self._build_table_sheet(ws_thieu, missing_df, is_cttt=False)

        wb.save(output_path)
        logger.info("Saved reconciliation report to: %s", output_path)
        return output_path

    def _build_tonghop_sheet(self, ws: Any, summary: Dict[str, Any]) -> None:
        """Create summary KPI cards and model metadata."""
        ws.views.sheetView[0].showGridLines = True

        ws["B2"] = "BÁO CÁO TỔNG HỢP SO SÁNH BOM (CTTT vs PLM vs R3)"
        ws["B2"].font = Font(name="Calibri", size=14, bold=True, color="1F497D")

        items = [
            ("Mã Model sản phẩm:", summary.get("model", "N/A")),
            ("Ngày hiệu lực đối soát:", summary.get("target_date", "N/A")),
            ("Thời gian chạy:", summary.get("timestamp", "N/A")),
            ("Tổng số linh kiện đối soát:", summary.get("total_parts", 0)),
            ("Số linh kiện hoàn toàn Khớp (OK):", summary.get("ok_count", 0)),
            ("Số linh kiện Sai Khác (NG):", summary.get("ng_count", 0)),
            ("Số cảnh báo Service / MSI:", summary.get("warning_count", 0)),
            ("Trạng thái chung:", "HOÀN TẤT - OK" if summary.get("ng_count", 0) == 0 else "CẦN GIẢI TRÌNH (NG)"),
        ]

        row = 4
        for label, val in items:
            ws.cell(row=row, column=2, value=label).font = Font(name="Calibri", size=11, bold=True)
            val_cell = ws.cell(row=row, column=3, value=val)
            val_cell.font = Font(name="Calibri", size=11)
            val_cell.border = self.thin_border

            if label == "Trạng thái chung:":
                val_cell.font = Font(name="Calibri", size=11, bold=True)
                if val == "HOÀN TẤT - OK":
                    val_cell.fill = self.green_fill
                    val_cell.font = self.green_font
                else:
                    val_cell.fill = self.red_fill
                    val_cell.font = self.red_font

            row += 1

        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 30

    def _build_table_sheet(self, ws: Any, df: pd.DataFrame, is_cttt: bool = False) -> None:
        """Write DataFrame into sheet with styling, auto-width, and conditional formatting."""
        ws.views.sheetView[0].showGridLines = True

        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = self.thin_border

                if r_idx == 1:
                    cell.fill = self.header_fill
                    cell.font = self.header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    val_str = str(value).strip().upper()
                    if val_str == "OK":
                        cell.fill = self.green_fill
                        cell.font = self.green_font
                        cell.alignment = Alignment(horizontal="center")
                    elif val_str == "NG":
                        cell.fill = self.red_fill
                        cell.font = self.red_font
                        cell.alignment = Alignment(horizontal="center")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

