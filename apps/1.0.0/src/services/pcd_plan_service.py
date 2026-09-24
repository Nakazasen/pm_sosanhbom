"""PCD Monthly Production Plan scanner service.

Scans PCD SharePoint / OneDrive folder structure:
Theo tháng (月別) / <Year> / <YearMonth> /
Picks end-of-period file (*定期計画後明細計画.xlsx) over beginning-of-period file (*月初明細計画.xlsx).
Filters sheet '詳細日程' where Column H (履歴) == 'NEW' and Column D (Material) starts with 'T1' or '11'.
"""

from __future__ import annotations

import datetime
import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, List, Optional

import openpyxl

from src.services.machine_dict_service import MachineDictService, MachineInfo

logger = logging.getLogger(__name__)

DEFAULT_ONEDRIVE_BASE = Path(r"C:\Users\tvn183660\OneDrive - KYOCERA Document Solutions")


class PCDPlanReadError(Exception):
    """Raised when PCD plan file cannot be read, hydrated, or parsed."""
    pass


def is_cloud_placeholder(path: Path) -> bool:
    """Check if file is a dehydrated OneDrive cloud placeholder (FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS)."""
    try:
        attrs = path.stat().st_file_attributes
        # 0x400000 = FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS, 0x1000 = FILE_ATTRIBUTE_OFFLINE
        return bool(attrs & (0x400000 | 0x1000))
    except Exception:
        return False


def is_onedrive_running() -> bool:
    """Check if OneDrive.exe process is currently active."""
    try:
        import psutil
        for proc in psutil.process_iter(["name"]):
            name = proc.info.get("name") or ""
            if name.lower() in ("onedrive.exe", "onedrive"):
                return True
    except Exception:
        pass
    return False


def start_onedrive() -> bool:
    """Attempt to launch OneDrive application."""
    import subprocess
    candidates = [
        Path(r"C:\Program Files\Microsoft OneDrive\OneDrive.exe"),
        Path(r"C:\Program Files (x86)\Microsoft OneDrive\OneDrive.exe"),
        Path(os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\OneDrive\OneDrive.exe")),
    ]
    for c in candidates:
        if c.exists():
            try:
                subprocess.Popen([str(c)], shell=False)
                return True
            except Exception as e:
                logger.warning("Could not launch %s: %s", c, e)
    return False


def check_file_accessible(path: Path) -> tuple[bool, str]:
    """Check whether a plan file is physically accessible and readable."""
    if not path.exists():
        return False, f"Tệp không tồn tại: {path}"

    if path.suffix.lower() == ".url":
        return False, "Đây là lối tắt Internet (.url) tới SharePoint trên trình duyệt, không phải tệp bảng tính Excel (.xlsx)."

    try:
        size = path.stat().st_size
        if size == 0:
            return False, "Tệp có kích thước 0 byte (tệp rỗng hoặc chưa hoàn tất tải về)."
    except Exception as e:
        return False, f"Không thể lấy thông tin tệp: {e}"

    if is_cloud_placeholder(path):
        if not is_onedrive_running():
            return False, (
                "Tệp này đang lưu trên đám mây OneDrive ở chế độ trực tuyến (Cloud-only) "
                "và ứng dụng OneDrive hiện chưa được mở trên máy tính.\n"
                "Vui lòng mở ứng dụng OneDrive để tệp được tải về máy trước khi quét."
            )

    return True, ""


OFFICIAL_PCD_SHAREPOINT_URL = (
    "https://kdcf.sharepoint.com/sites/kdtvn_PCD/ProductionPlan/Forms/AllItems.aspx"
    "?id=%2Fsites%2Fkdtvn%5FPCD%2FProductionPlan%2FTheo%20th%C3%A1ng%20%28%E6%9C%88%E5%88%A5%29"
    "&viewid=bbba5c4e%2D26b9%2D47fd%2Daf3a%2Dd6f0aed08994"
)


def open_pcd_sharepoint_url() -> bool:
    """Open official PCD Production Plan SharePoint library in the user's default web browser."""
    import webbrowser
    try:
        return webbrowser.open(OFFICIAL_PCD_SHAREPOINT_URL)
    except Exception as exc:
        logger.warning("Could not open browser for URL %s: %s", OFFICIAL_PCD_SHAREPOINT_URL, exc)
        return False


def _get_configured_pcd_base_dir() -> Optional[Path]:
    """Retrieve custom pcd_base_dir from config/settings.json if defined."""
    import json
    cfg_paths = [Path("config/settings.json"), Path("../config/settings.json")]
    for cp in cfg_paths:
        if cp.exists():
            try:
                with open(cp, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                val = cfg.get("paths", {}).get("pcd_base_dir")
                if val:
                    p = Path(val)
                    if p.exists():
                        return p
            except Exception:
                pass
    return None


@dataclass
class PCDPlanItem:
    """Represents a filtered production plan item."""

    material_code: str          # Cột D (Material): e.g. T10C0TZUS0, 1102Y43AX0
    machine_code_4char: str     # 4-character machine code e.g. 0C0T, 02Y4
    history_status: str         # Cột H (履歴): e.g. NEW
    is_trial: bool              # True if starts with T1 (DMT/PMT), False if starts with 11 (MP)
    suggested_phase: str        # 'DMT/PMT' for T1, 'MP' for 11
    machine_info: Optional[MachineInfo] = None  # Matched from dictionary
    machine_name: str = ""      # e.g. 6th Next, Virgo, Libra 2...
    raw_data: dict[str, Any] = field(default_factory=dict)


@dataclass
class PCDPlanScanResult:
    """Result of scanning a monthly plan folder/file."""

    year: int
    month: int
    plan_file: Path
    file_type: str              # 'end_of_period' (*定期計画後明細計画) or 'beginning_of_period' (*月初明細計画)
    total_rows_scanned: int = 0
    items: list[PCDPlanItem] = field(default_factory=list)


class PCDPlanService:
    """Service to discover and parse PCD monthly production plans."""

    def __init__(
        self,
        base_dir: str | Path | None = None,
        dict_service: MachineDictService | None = None,
    ) -> None:
        self._is_custom_base = base_dir is not None
        self.base_dir = Path(base_dir or DEFAULT_ONEDRIVE_BASE)
        self.dict_service = dict_service or MachineDictService()

    def find_monthly_plan_file(
        self,
        year: int | None = None,
        month: int | None = None,
        custom_folder: str | Path | None = None,
    ) -> Optional[tuple[Path, str]]:
        """Find the plan file for given year and month.

        Priority:
        1. End-of-period (*定期計画後明細計画.xlsx)
        2. Beginning-of-period (*月初明細計画.xlsx)
        Strictly excludes quality status reports or non-plan documents.
        """
        now = datetime.datetime.now()
        target_year = year or now.year
        target_month = month or now.month
        year_str = str(target_year)
        month_str = f"{target_year}{target_month:02d}"

        search_dirs: list[Path] = []

        if custom_folder:
            search_dirs.append(Path(custom_folder))

        candidate_paths: list[Path] = []

        # 1. If base_dir was not explicitly provided by caller, check user configured custom pcd_base_dir
        if not self._is_custom_base:
            conf_dir = _get_configured_pcd_base_dir()
            if conf_dir:
                candidate_paths.extend([
                    conf_dir / "Theo tháng (月別)" / year_str / month_str,
                    conf_dir / year_str / month_str,
                    conf_dir,
                ])

        # 2. Check in OneDrive sync locations (Personal & SharePoint shortcut roots)
        user_home = Path(os.path.expanduser("~"))
        candidate_paths.extend([
            self.base_dir / "Theo tháng (月別)" / year_str / month_str,
            self.base_dir / "ProductionPlan" / "Theo tháng (月別)" / year_str / month_str,
            self.base_dir / "ProductionPlan" / year_str / month_str,
            self.base_dir / year_str / month_str,
            user_home / "KYOCERA Document Solutions" / "kdtvn_PCD - ProductionPlan" / "Theo tháng (月別)" / year_str / month_str,
            user_home / "KYOCERA Document Solutions" / "kdtvn_PCD - Theo tháng (月別)" / year_str / month_str,
            user_home / "KYOCERA Document Solutions" / "Theo tháng (月別)" / year_str / month_str,
        ])

        # 3. Check LAN directories
        lan_roots = [
            Path(r"\\fstvn01\Data\10_Production Engineering Department(製造技術部)\02.製造技術課\PE Dept\4A. QUAN LY BOM-TDTK-BOM管理-設計変更\SO SANH PLM-CTTT-R3"),
            Path(r"\\fstvn01\Data\10_Production Engineering Department(製造技術部)\02.製造技術課\PE Dept\4A. QUAN LY BOM-TDTK-BOM管理-設計変更"),
        ]
        for lr in lan_roots:
            if lr.exists():
                candidate_paths.append(lr / "PCD" / year_str / month_str)
                candidate_paths.append(lr / year_str / month_str)

        # 4. Base dir as fallback
        candidate_paths.append(self.base_dir)

        for p in candidate_paths:
            if p.exists() and p.is_dir() and p not in search_dirs:
                search_dirs.append(p)

        logger.info("Searching for monthly plan in %d candidate directories for %s/%s", len(search_dirs), year_str, month_str)

        PLAN_KEYWORDS = ["定期計画後明細計画", "月初明細計画", "明細計画", "ke hoach", "kehoach", "pcd", "plan"]
        EXCLUDE_KEYWORDS = [
            "品質状況", "chat luong", "chất lượng", "kaizo", "thay the", "thay thế",
            "dieutra", "điều tra", "leaflet", "rating label", "tham dinh", "thẩm định",
        ]

        for sdir in search_dirs:
            try:
                files = [
                    f for f in sdir.glob("*.xls*")
                    if not f.name.endswith(".url") and not f.name.startswith("~$")
                ]

                def matches_year_and_not_excluded(f: Path) -> bool:
                    fname_lower = f.name.lower()
                    if any(ex in fname_lower for ex in EXCLUDE_KEYWORDS):
                        return False
                    # If file has a 4-digit year pattern (e.g. 2024, 2025, 2026), it must match target_year
                    years_in_name = re.findall(r"20\d\d", f.name)
                    if years_in_name and year_str not in years_in_name:
                        return False
                    return True

                # Priority 1: End of period
                for f in files:
                    if "定期計画後明細計画" in f.name and matches_year_and_not_excluded(f):
                        if month_str in f.name or f"{target_month}月" in f.name or f"{target_month:02d}" in f.name:
                            return f, "end_of_period"

                # Priority 2: Beginning of period
                for f in files:
                    if "月初明細計画" in f.name and matches_year_and_not_excluded(f):
                        if month_str in f.name or f"{target_month}月" in f.name or f"{target_month:02d}" in f.name:
                            return f, "beginning_of_period"

                # Fallback: strictly plan file matching year AND month pattern
                for f in files:
                    if matches_year_and_not_excluded(f):
                        has_plan_kw = any(k in f.name.lower() for k in PLAN_KEYWORDS) or ("計画" in f.name and "品質" not in f.name)
                        if has_plan_kw and (month_str in f.name or f"{target_month}月" in f.name):
                            return f, "custom_plan"
            except Exception as e:
                logger.debug("Error checking directory %s: %s", sdir, e)

        return None

    def _read_via_excel_com(self, path: Path) -> tuple[str, list[list[Any]]]:
        """Read workbook via background Excel COM instance.

        Supports files protected by Microsoft AIP/RMS, .xls, and .xlsb.
        """
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        excel = None
        wb = None
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
            wb = excel.Workbooks.Open(str(path), ReadOnly=True, UpdateLinks=False)

            target_ws = None
            for sheet in wb.Sheets:
                if "詳細日程" in sheet.Name:
                    target_ws = sheet
                    break
            if not target_ws:
                target_ws = wb.Sheets(1)

            sheet_name = target_ws.Name
            used_range = target_ws.UsedRange
            raw_vals = used_range.Value

            rows: list[list[Any]] = []
            if raw_vals is not None:
                if isinstance(raw_vals, (list, tuple)):
                    for r in raw_vals:
                        if isinstance(r, (list, tuple)):
                            rows.append(list(r))
                        else:
                            rows.append([r])
                else:
                    rows.append([raw_vals])
            return sheet_name, rows
        finally:
            if wb:
                try:
                    wb.Close(SaveChanges=False)
                except Exception:
                    pass
            if excel:
                try:
                    excel.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()

    def parse_plan_file(self, file_path: str | Path) -> PCDPlanScanResult:
        """Parse Excel plan file, read sheet '詳細日程', filter Column H='NEW' & Column D starts with 'T1' or '11'."""
        path = Path(file_path).resolve()

        # 1. Pre-flight accessibility and cloud status check
        ok, reason = check_file_accessible(path)
        if not ok:
            raise PCDPlanReadError(reason)

        # Ensure dictionary is loaded
        self.dict_service.load()

        file_type = "end_of_period" if "定期計画後明細計画" in path.name else "beginning_of_period"
        now = datetime.datetime.now()

        # 2. Try loading with openpyxl first
        rows_data: list[list[Any]] = []
        target_sheet_name: str = ""
        load_success = False
        exc_openpyxl_err = ""

        try:
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            for sname in wb.sheetnames:
                if "詳細日程" in sname:
                    target_sheet_name = sname
                    break
            if not target_sheet_name and wb.sheetnames:
                target_sheet_name = wb.sheetnames[0]
                logger.warning("Sheet '詳細日程' not found in %s, falling back to '%s'", path.name, target_sheet_name)

            ws = wb[target_sheet_name]
            for row in ws.iter_rows(values_only=True):
                rows_data.append(list(row))
            wb.close()
            load_success = True
        except Exception as exc_openpyxl:
            exc_openpyxl_err = str(exc_openpyxl)
            logger.warning("openpyxl failed to load %s (%s). Attempting Excel COM fallback...", path.name, exc_openpyxl)

        # 3. Fallback to win32com (Excel COM) if openpyxl failed (supports AIP/RMS protected files, .xls, .xlsb)
        if not load_success:
            try:
                sheet_found, com_rows = self._read_via_excel_com(path)
                target_sheet_name = sheet_found
                rows_data = com_rows
                load_success = True
                logger.info("Successfully loaded %s via Excel COM fallback (%d rows)", path.name, len(rows_data))
            except Exception as exc_com:
                logger.error("Excel COM fallback also failed for %s: %s", path.name, exc_com)
                raise PCDPlanReadError(
                    f"Không thể đọc tệp Excel '{path.name}'.\n"
                    f"Nguyên nhân: Tệp có thể đang bị khóa, chưa được tải về từ đám mây, "
                    f"hoặc không đúng định dạng bảng tính Excel hợp lệ.\n"
                    f"Chi tiết: {exc_openpyxl_err or exc_com}"
                )

        if not rows_data:
            raise PCDPlanReadError(f"Tệp Excel '{path.name}' không có dữ liệu trong sheet '{target_sheet_name}'.")

        # 4. Scan header row to identify Material (Col D) and History (Col H)
        header_row_idx = 0
        mat_col_idx = 4  # Column D (1-indexed)
        history_col_idx = 8  # Column H (1-indexed)

        for r_idx in range(min(5, len(rows_data))):
            row = rows_data[r_idx]
            for c_idx, val in enumerate(row, 1):
                sval = str(val or "").strip()
                if sval in ("Material", "品目", "Mã hàng", "Mã vật tư"):
                    mat_col_idx = c_idx
                    header_row_idx = r_idx
                elif sval in ("履歴", "History", "Trạng thái", "Status"):
                    history_col_idx = c_idx

        items: list[PCDPlanItem] = []
        total_rows = 0

        for r_idx in range(header_row_idx + 1, len(rows_data)):
            row = rows_data[r_idx]
            total_rows += 1
            if not row or len(row) < max(mat_col_idx, history_col_idx):
                continue

            val_mat = str(row[mat_col_idx - 1] or "").strip().upper()
            val_history = str(row[history_col_idx - 1] or "").strip().upper()

            # Filter Condition: History == 'NEW' and Material starts with 'T1' or '11'
            if val_history == "NEW" and (val_mat.startswith("T1") or val_mat.startswith("11")):
                is_trial = val_mat.startswith("T1")
                suggested_phase = "DMT/PMT" if is_trial else "MP"

                matched_machine = self.dict_service.extract_and_lookup_material(val_mat)
                machine_name = matched_machine.machine_name if matched_machine else ""

                # Extract 4-char machine code:
                # 1. Match against known codes of the machine
                machine_code_4char = ""
                if matched_machine:
                    for c in matched_machine.machine_codes:
                        if c and c.upper() in val_mat:
                            machine_code_4char = c.upper()
                            break

                # 2. Extract 4-char machine code by Kyocera standard:
                # Material code starts with T1 (Trial) or 11 (MP).
                # 4-char machine code starts from character 3 (index 2: val_mat[2:6]),
                # e.g., T10C452US0 -> 0C45, 110C3M3AK0 -> 0C3M, T10C0P3NL0 -> 0C0P.
                # For codes with two zeros like 11002YJNL0, index 3:7 matches 02YJ.
                if not machine_code_4char:
                    c2 = val_mat[2:6] if len(val_mat) >= 6 else ""
                    c3 = val_mat[3:7] if len(val_mat) >= 7 else ""
                    if c2.startswith(("0C", "02", "03")):
                        machine_code_4char = c2
                    elif c3.startswith(("0C", "02", "03")):
                        machine_code_4char = c3
                    elif c2:
                        machine_code_4char = c2

                item = PCDPlanItem(
                    material_code=val_mat,
                    machine_code_4char=machine_code_4char,
                    history_status=val_history,
                    is_trial=is_trial,
                    suggested_phase=suggested_phase,
                    machine_info=matched_machine,
                    machine_name=machine_name,
                    raw_data={"row": total_rows, "material": val_mat, "history": val_history},
                )
                items.append(item)

        logger.info("Parsed %d rows from %s, found %d NEW materials (T1/11)", total_rows, path.name, len(items))

        return PCDPlanScanResult(
            year=now.year,
            month=now.month,
            plan_file=path,
            file_type=file_type,
            total_rows_scanned=total_rows,
            items=items,
        )

