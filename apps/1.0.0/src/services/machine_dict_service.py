"""Machine dictionary and 2-way sync service with file_loaimay_nhommail.xlsx."""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

logger = logging.getLogger(__name__)

DEFAULT_DICT_PATH = (
    r"\\fstvn01\Data\00_KDTVN Common(KDTVN共通)\⑤Production Engineering(製造技術)"
    r"\Hang muc can luu\Vinh\Pm_sosanhBOM\file_loaimay_nhommail.xlsx"
)


@dataclass
class MachineInfo:
    """Represents machine mapping information from Excel dictionary."""

    machine_name: str           # Cột A: Tên máy (ví dụ: 6th A3, Virgo, Iris 2024...)
    variant: str = ""           # Cột B: Kiểu (ví dụ: Thường, 6th A3S, PRT, MFP...)
    brand_segment: str = ""     # Cột C: Thương hiệu/Phân khúc (ví dụ: KDC, OEM, H, L...)
    machine_codes: list[str] = field(default_factory=list)  # Cột D: Danh sách mã máy 4 ký tự
    specs: str = ""             # Cột E: Thông số kỹ thuật (ví dụ: 250, 500, 40/50/60/70...)
    row_index: int = -1         # 1-indexed row number in Master sheet


@dataclass
class EmailGroupInfo:
    """Represents mail distribution groups from Excel dictionary."""

    group_label: str    # ví dụ: To, CC1(nhóm cơ 1), CC2(nhóm cơ 2)...
    email_string: str   # ví dụ: KDTVN-Production Engineering_Mecha 1(Local) <...>
    clean_email: str    # ví dụ: KDTVN-ProductionEngineering_Mecha1_Local_@kdcf.onmicrosoft.com


class MachineDictService:
    """Service to load, cache, lookup, and write-back machine types and email groups."""

    def __init__(self, excel_path: str | Path | None = None) -> None:
        self.excel_path = Path(excel_path or DEFAULT_DICT_PATH)
        self._machines: list[MachineInfo] = []
        self._code_to_machine: dict[str, MachineInfo] = {}
        self._email_groups: dict[str, EmailGroupInfo] = {}
        self._is_loaded = False

    def load(self, force_reload: bool = False) -> bool:
        """Load and parse the Master sheet from file_loaimay_nhommail.xlsx or local cache."""
        if self._is_loaded and not force_reload:
            return True

        target_path: Optional[Path] = None
        if self.excel_path.exists():
            target_path = self.excel_path
        else:
            # Check candidate fallback paths
            candidates = [
                Path("file_loaimay_nhommail.xlsx"),
                Path("data/file_loaimay_nhommail.xlsx"),
                Path("config/file_loaimay_nhommail.xlsx"),
                Path("../config/file_loaimay_nhommail.xlsx"),
            ]
            for cand in candidates:
                if cand.exists():
                    target_path = cand
                    break

        if target_path and target_path.exists():
            try:
                wb = openpyxl.load_workbook(str(target_path), data_only=True)
                sheet_name = "Master" if "Master" in wb.sheetnames else wb.sheetnames[0]
                ws = wb[sheet_name]

                machines: list[MachineInfo] = []
                code_map: dict[str, MachineInfo] = {}
                email_groups: dict[str, EmailGroupInfo] = {}

                # Read rows starting from row 2 (skipping header)
                for r_idx in range(2, ws.max_row + 1):
                    col_a = ws.cell(row=r_idx, column=1).value
                    col_b = ws.cell(row=r_idx, column=2).value
                    col_c = ws.cell(row=r_idx, column=3).value
                    col_d = ws.cell(row=r_idx, column=4).value
                    col_e = ws.cell(row=r_idx, column=5).value
                    col_f = ws.cell(row=r_idx, column=6).value
                    col_g = ws.cell(row=r_idx, column=7).value

                    # Parse Email groups (Cols F & G)
                    if col_f and col_g:
                        label = str(col_f).strip()
                        raw_email = str(col_g).strip()
                        match = re.search(r"<([^>]+)>", raw_email)
                        clean_email = match.group(1).strip() if match else raw_email
                        email_groups[label] = EmailGroupInfo(
                            group_label=label,
                            email_string=raw_email,
                            clean_email=clean_email,
                        )

                    # Parse Machine Info (Cols A - E)
                    if col_a or col_d:
                        raw_name = str(col_a).strip() if col_a else ""
                        machine_name = re.sub(r"\s+", "", raw_name)
                        variant = str(col_b).strip() if col_b else ""
                        brand_segment = str(col_c).strip() if col_c else ""
                        raw_codes = str(col_d).strip() if col_d else ""
                        specs = str(col_e).strip() if col_e else ""

                        codes = [c.strip().upper() for c in raw_codes.split(";") if c.strip()]

                        info = MachineInfo(
                            machine_name=machine_name,
                            variant=variant,
                            brand_segment=brand_segment,
                            machine_codes=codes,
                            specs=specs,
                            row_index=r_idx,
                        )
                        machines.append(info)

                        for code in codes:
                            code_map[code] = info

                wb.close()

                self._machines = machines
                self._code_to_machine = code_map
                self._email_groups = email_groups
                self._is_loaded = True
                logger.info("Loaded %d machines and %d email groups from %s", len(machines), len(email_groups), target_path)
                return True
            except Exception as exc:
                logger.error("Failed to load machine dictionary from %s: %s", target_path, exc)

        # Fallback to JSON cache if available
        base_dir = Path(__file__).resolve().parent
        json_candidates = [
            Path("config/loaimay_cache.json"),
            Path("../config/loaimay_cache.json"),
            base_dir.parent.parent / "config" / "loaimay_cache.json",
            base_dir.parent / "config" / "loaimay_cache.json",
            Path("scratch/loaimay_data.json"),
        ]
        for jcand in json_candidates:
            if jcand.exists():
                try:
                    with open(jcand, "r", encoding="utf-8") as f:
                        jdata = json.load(f)
                    master_rows = jdata.get("Master", [])
                    machines = []
                    code_map = {}
                    email_groups = {}
                    for r_idx, row in enumerate(master_rows[1:], start=2):
                        col_a = row[0] if len(row) > 0 else None
                        col_b = row[1] if len(row) > 1 else None
                        col_c = row[2] if len(row) > 2 else None
                        col_d = row[3] if len(row) > 3 else None
                        col_e = row[4] if len(row) > 4 else None
                        col_f = row[5] if len(row) > 5 else None
                        col_g = row[6] if len(row) > 6 else None

                        if col_f and col_g:
                            label = str(col_f).strip()
                            raw_email = str(col_g).strip()
                            match = re.search(r"<([^>]+)>", raw_email)
                            clean_email = match.group(1).strip() if match else raw_email
                            email_groups[label] = EmailGroupInfo(
                                group_label=label,
                                email_string=raw_email,
                                clean_email=clean_email,
                            )

                        if col_a or col_d:
                            raw_name = str(col_a).strip() if col_a else ""
                            machine_name = re.sub(r"\s+", "", raw_name)
                            variant = str(col_b).strip() if col_b else ""
                            brand_segment = str(col_c).strip() if col_c else ""
                            raw_codes = str(col_d).strip() if col_d else ""
                            specs = str(col_e).strip() if col_e else ""
                            codes = [c.strip().upper() for c in raw_codes.split(";") if c.strip()]
                            info = MachineInfo(
                                machine_name=machine_name,
                                variant=variant,
                                brand_segment=brand_segment,
                                machine_codes=codes,
                                specs=specs,
                                row_index=r_idx,
                            )
                            machines.append(info)
                            for code in codes:
                                code_map[code] = info

                    self._machines = machines
                    self._code_to_machine = code_map
                    self._email_groups = email_groups
                    self._is_loaded = True

                    # Check settings.json override for email_groups
                    cfg_path = self._get_local_settings_path()
                    if cfg_path.exists():
                        try:
                            with open(cfg_path, "r", encoding="utf-8") as f:
                                cfg = json.load(f)
                            custom_eg = cfg.get("email_groups")
                            if isinstance(custom_eg, dict):
                                for label, raw_email in custom_eg.items():
                                    match = re.search(r"<([^>]+)>", str(raw_email))
                                    clean_email = match.group(1).strip() if match else str(raw_email).strip()
                                    self._email_groups[str(label).strip()] = EmailGroupInfo(
                                        group_label=str(label).strip(),
                                        email_string=str(raw_email).strip(),
                                        clean_email=clean_email,
                                    )
                        except Exception:
                            pass

                    logger.info("Loaded %d machines from JSON cache %s", len(machines), jcand)
                    return True
                except Exception as exc:
                    logger.error("Failed loading JSON cache %s: %s", jcand, exc)

        logger.warning("Machine dictionary file does not exist: %s and no local cache found", self.excel_path)
        return False

    def lookup_machine_code(self, code_4char: str) -> Optional[MachineInfo]:
        """Lookup by 4-character machine code (e.g. '0C0T', '02Y4')."""
        if not self._is_loaded:
            self.load()
        cleaned = code_4char.strip().upper()
        return self._code_to_machine.get(cleaned)

    def get_model_for_machine_code(self, machine_code: str, fallback_model: str = "") -> str:
        """Resolve machine code (e.g. '110C103NL0', 'T10C0TZUS0', '0C10') to canonical model name."""
        if not self._is_loaded:
            self.load()
        cleaned = machine_code.strip().upper()
        # 1. Try direct extract
        info = self.extract_and_lookup_material(cleaned)
        if info and info.machine_name:
            return re.sub(r"\s+", "", info.machine_name.strip())
        # 2. Try 4-char lookup
        if len(cleaned) == 4:
            info = self.lookup_machine_code(cleaned)
            if info and info.machine_name:
                return re.sub(r"\s+", "", info.machine_name.strip())
        # 3. Try slicing chars 2..6
        if len(cleaned) >= 6:
            code_4 = cleaned[2:6]
            info = self.lookup_machine_code(code_4)
            if info and info.machine_name:
                return re.sub(r"\s+", "", info.machine_name.strip())
        return fallback_model

    def extract_and_lookup_material(self, material_code: str) -> Optional[MachineInfo]:
        """Extract machine code from material (e.g. 'T10C0TZUS0' -> '0C0T' or '1102Y43AX0' -> '02Y4').

        T1 + [4-char code] + ... or 11 + [4-char code] + ...
        """
        if not self._is_loaded:
            self.load()

        clean_mat = material_code.strip().upper()

        # Strategy 1: Check prefix T1 or 11
        if len(clean_mat) >= 6:
            candidate_4char = clean_mat[2:6]
            if candidate_4char in self._code_to_machine:
                return self._code_to_machine[candidate_4char]

        # Strategy 2: Search across all known machine codes
        for code, info in self._code_to_machine.items():
            if code in clean_mat:
                return info

        return None

    def get_all_machines(self) -> list[MachineInfo]:
        """Return all loaded machines."""
        if not self._is_loaded:
            self.load()
        return list(self._machines)

    def get_all_email_groups(self) -> dict[str, EmailGroupInfo]:
        """Return all email groups."""
        if not self._is_loaded:
            self.load()
        return dict(self._email_groups)

    def update_email_groups(self, email_groups: dict[str, str]) -> bool:
        """Update email distribution groups in settings.json and 2-way write-back to Excel."""
        if not email_groups:
            return False

        # 1. Update in-memory
        new_groups: dict[str, EmailGroupInfo] = {}
        for label, raw_email in email_groups.items():
            clean_label = str(label).strip()
            clean_raw = str(raw_email).strip()
            if not clean_label or not clean_raw:
                continue
            match = re.search(r"<([^>]+)>", clean_raw)
            clean_email = match.group(1).strip() if match else clean_raw
            new_groups[clean_label] = EmailGroupInfo(
                group_label=clean_label,
                email_string=clean_raw,
                clean_email=clean_email,
            )
        if not new_groups:
            return False

        self._email_groups = new_groups

        # 2. Persist to settings.json
        cfg_path = self._get_local_settings_path()
        try:
            cfg: dict[str, Any] = {}
            if cfg_path.exists():
                with open(cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
            cfg["email_groups"] = {lbl: info.email_string for lbl, info in new_groups.items()}
            cfg_path.parent.mkdir(parents=True, exist_ok=True)
            tmp = cfg_path.with_suffix(".tmp")
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(cfg, f, indent=2, ensure_ascii=False)
            tmp.replace(cfg_path)
            logger.info("Saved email groups to %s", cfg_path)
        except Exception as e:
            logger.warning("Could not persist email groups to %s: %s", cfg_path, e)

        # 3. Write-back to file_loaimay_nhommail.xlsx if exists
        target_path: Optional[Path] = None
        if self.excel_path.exists():
            target_path = self.excel_path
        else:
            candidates = [
                Path("file_loaimay_nhommail.xlsx"),
                Path("data/file_loaimay_nhommail.xlsx"),
                Path("config/file_loaimay_nhommail.xlsx"),
                Path("../config/file_loaimay_nhommail.xlsx"),
            ]
            for c in candidates:
                if c.exists():
                    target_path = c
                    break

        if target_path and target_path.exists():
            try:
                backup_path = target_path.with_suffix(".xlsx.bak")
                try:
                    shutil.copy2(str(target_path), str(backup_path))
                except Exception:
                    pass

                wb = openpyxl.load_workbook(str(target_path))
                sheet_name = "Master" if "Master" in wb.sheetnames else wb.sheetnames[0]
                ws = wb[sheet_name]

                # Clear previous email columns F and G from row 2 onwards
                for r in range(2, max(ws.max_row + 1, len(new_groups) + 2)):
                    ws.cell(row=r, column=6).value = None
                    ws.cell(row=r, column=7).value = None

                # Write new email groups
                for r_idx, (lbl, info) in enumerate(new_groups.items(), start=2):
                    ws.cell(row=r_idx, column=6).value = lbl
                    ws.cell(row=r_idx, column=7).value = info.email_string

                wb.save(str(target_path))
                wb.close()
                logger.info("Wrote %d email groups to %s", len(new_groups), target_path)
            except Exception as e:
                logger.error("Failed writing email groups to %s: %s", target_path, e)

        # 4. Also update loaimay_cache.json if present
        cache_paths = [Path("config/loaimay_cache.json"), Path("../config/loaimay_cache.json")]
        for cp in cache_paths:
            if cp.exists():
                try:
                    with open(cp, "r", encoding="utf-8") as f:
                        cdata = json.load(f)
                    master = cdata.get("Master", [])
                    email_list = list(new_groups.items())
                    for idx, row in enumerate(master[1:], start=0):
                        while len(row) < 7:
                            row.append(None)
                        if idx < len(email_list):
                            lbl, info = email_list[idx]
                            row[5] = lbl
                            row[6] = info.email_string
                        else:
                            row[5] = None
                            row[6] = None
                    if len(email_list) > len(master) - 1:
                        for idx in range(len(master) - 1, len(email_list)):
                            lbl, info = email_list[idx]
                            new_row = [None, None, None, None, None, lbl, info.email_string]
                            master.append(new_row)
                    cdata["Master"] = master
                    with open(cp, "w", encoding="utf-8") as f:
                        json.dump(cdata, f, indent=2, ensure_ascii=False)
                    logger.info("Updated cache %s with new email groups", cp)
                except Exception as e:
                    logger.debug("Could not update cache %s: %s", cp, e)

        return True

    def _get_local_settings_path(self) -> Path:
        cand = Path("config/settings.json")
        if cand.exists():
            return cand
        cand_up = Path("../config/settings.json")
        if cand_up.exists():
            return cand_up
        return Path("config/settings.json")

    def get_model_names(self) -> list[str]:
        """Return sorted list of all unique machine/model names from Excel dictionary and local config.

        Includes all machine models from Column A in file_loaimay_nhommail.xlsx with all
        whitespace removed (e.g. 'Iris 2024' -> 'Iris2024', 'Polaris Next' -> 'PolarisNext'),
        plus baseline models and any user-configured custom models from settings.json.
        """
        if not self._is_loaded:
            self.load()
        names: set[str] = set()
        for m in self._machines:
            if m.machine_name:
                cleaned = re.sub(r"\s+", "", m.machine_name.strip())
                if cleaned:
                    names.add(cleaned)
        legacy_defaults = ["Virgo", "Libra2", "Iris2024", "Sirius2", "Mebius", "Polaris"]
        for leg in legacy_defaults:
            names.add(re.sub(r"\s+", "", leg))

        # Include custom models and project configs from settings.json
        cfg_path = self._get_local_settings_path()
        if cfg_path.exists():
            try:
                with open(cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                for c in cfg.get("custom_models", []):
                    cleaned = re.sub(r"\s+", "", str(c).strip())
                    if cleaned:
                        names.add(cleaned)
                for p in cfg.get("project_configs", {}).keys():
                    cleaned = re.sub(r"\s+", "", str(p).strip())
                    if cleaned:
                        names.add(cleaned)
            except Exception as e:
                logger.debug("Could not read custom models from %s: %s", cfg_path, e)

        return sorted(names, key=lambda s: s.lower())

    def add_model_name(self, model_name: str) -> bool:
        """Register a new machine model both locally in settings.json and in master Excel if accessible."""
        cleaned = re.sub(r"\s+", "", model_name.strip())
        if not cleaned:
            return False

        # 1. Save to settings.json
        cfg_path = self._get_local_settings_path()
        try:
            cfg: dict[str, Any] = {}
            if cfg_path.exists():
                with open(cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
            custom_models = cfg.get("custom_models", [])
            if not isinstance(custom_models, list):
                custom_models = []
            if cleaned not in custom_models:
                custom_models.append(cleaned)
                cfg["custom_models"] = custom_models

                cfg_path.parent.mkdir(parents=True, exist_ok=True)
                tmp_path = cfg_path.with_suffix(".tmp")
                with open(tmp_path, "w", encoding="utf-8") as f:
                    json.dump(cfg, f, indent=2, ensure_ascii=False)
                tmp_path.replace(cfg_path)
                logger.info("Saved custom model '%s' to %s", cleaned, cfg_path)
        except Exception as exc:
            logger.warning("Could not persist custom model to %s: %s", cfg_path, exc)

        # 2. Try to sync to Master Excel if accessible
        if self.excel_path.exists():
            try:
                if not self._is_loaded:
                    self.load()
                exists_in_excel = any(m.machine_name.lower() == cleaned.lower() for m in self._machines)
                if not exists_in_excel:
                    wb = openpyxl.load_workbook(str(self.excel_path))
                    sheet_name = "Master" if "Master" in wb.sheetnames else wb.sheetnames[0]
                    ws = wb[sheet_name]
                    target_row = ws.max_row + 1
                    ws.cell(row=target_row, column=1, value=cleaned)
                    wb.save(str(self.excel_path))
                    wb.close()
                    logger.info("Synced new model '%s' to master Excel %s", cleaned, self.excel_path)
                    self.load(force_reload=True)
            except Exception as e:
                logger.warning("Could not sync model to master Excel %s: %s", self.excel_path, e)

        return True

    def add_or_update_machine_code(
        self,
        machine_name: str,
        new_code: str,
        variant: str = "",
        brand_segment: str = "",
        specs: str = "",
    ) -> bool:
        """Add or update a machine code in local configuration and sync to file_loaimay_nhommail.xlsx if accessible."""
        cleaned_code = new_code.strip().upper()
        if not cleaned_code:
            return False

        clean_m = re.sub(r"\s+", "", machine_name.strip())
        if not clean_m:
            return False

        # 1. Persist to local settings.json custom_model_codes (works offline, in tests, or read-only LAN)
        cfg_path = self._get_local_settings_path()
        try:
            cfg: dict[str, Any] = {}
            if cfg_path.exists():
                with open(cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
            custom_model_codes = cfg.get("custom_model_codes", {})
            if not isinstance(custom_model_codes, dict):
                custom_model_codes = {}
            curr_list = custom_model_codes.get(clean_m, [])
            if not isinstance(curr_list, list):
                curr_list = []
            if cleaned_code not in curr_list:
                curr_list.append(cleaned_code)
                custom_model_codes[clean_m] = curr_list
                cfg["custom_model_codes"] = custom_model_codes

                cfg_path.parent.mkdir(parents=True, exist_ok=True)
                tmp_path = cfg_path.with_suffix(".tmp")
                with open(tmp_path, "w", encoding="utf-8") as f:
                    json.dump(cfg, f, indent=2, ensure_ascii=False)
                tmp_path.replace(cfg_path)
                logger.info("Saved code '%s' for model '%s' to %s", cleaned_code, clean_m, cfg_path)
        except Exception as exc:
            logger.warning("Could not persist code to %s: %s", cfg_path, exc)

        # 2. Try 2-way write-back to file_loaimay_nhommail.xlsx if accessible
        if self.excel_path.exists():
            try:
                # Backup before writing
                backup_path = self.excel_path.with_suffix(".xlsx.bak")
                try:
                    shutil.copy2(str(self.excel_path), str(backup_path))
                except Exception as e:
                    logger.warning("Could not create backup of dictionary: %s", e)

                wb = openpyxl.load_workbook(str(self.excel_path))
                sheet_name = "Master" if "Master" in wb.sheetnames else wb.sheetnames[0]
                ws = wb[sheet_name]

                target_row = -1
                found_existing = False

                # Search if machine_name already exists in Column A
                clean_search_name = clean_m.lower()
                for r_idx in range(2, ws.max_row + 1):
                    val_a = ws.cell(row=r_idx, column=1).value
                    val_b = ws.cell(row=r_idx, column=2).value

                    if val_a and re.sub(r"\s+", "", str(val_a)).lower() == clean_search_name:
                        # If variant is specified, try to match variant as well
                        if variant and val_b and str(val_b).strip().lower() != variant.strip().lower():
                            continue

                        target_row = r_idx
                        found_existing = True
                        break

                if found_existing:
                    # Append to existing row Column D
                    current_codes_val = ws.cell(row=target_row, column=4).value or ""
                    current_codes = [c.strip().upper() for c in str(current_codes_val).split(";") if c.strip()]
                    if cleaned_code not in current_codes:
                        current_codes.append(cleaned_code)
                        ws.cell(row=target_row, column=4, value="; ".join(current_codes))
                else:
                    # Append a new row at the end with whitespace-free machine name
                    target_row = ws.max_row + 1
                    ws.cell(row=target_row, column=1, value=clean_m)
                    ws.cell(row=target_row, column=2, value=variant.strip())
                    ws.cell(row=target_row, column=3, value=brand_segment.strip())
                    ws.cell(row=target_row, column=4, value=cleaned_code)
                    ws.cell(row=target_row, column=5, value=specs.strip())

                wb.save(str(self.excel_path))
                wb.close()
                logger.info("Successfully updated machine code %s for machine %s in %s", cleaned_code, clean_m, self.excel_path)

                # Reload internal cache
                self.load(force_reload=True)
            except Exception as exc:
                logger.error("Failed to write machine code to %s: %s", self.excel_path, exc)
        else:
            logger.debug("Dictionary file does not exist, skipping Excel sync: %s", self.excel_path)

        return True

    def get_machine_codes_for_model(self, model_name: str) -> list[str]:
        """Return all genuine 4-character machine codes associated with this model or model family."""
        if not self._is_loaded:
            self.load()
        cleaned_search = re.sub(r"\s+", "", model_name.strip()).lower()
        if not cleaned_search:
            return []

        codes: list[str] = []

        # 1. Exact match from loaded Excel dictionary
        for m in self._machines:
            if m.machine_name and re.sub(r"\s+", "", m.machine_name.strip()).lower() == cleaned_search:
                for c in m.machine_codes:
                    clean_c = c.strip().upper()
                    if clean_c and clean_c not in codes:
                        codes.append(clean_c)

        # 2. Exact match from settings.json custom_model_codes
        cfg_path = self._get_local_settings_path()
        if cfg_path.exists():
            try:
                with open(cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                custom_codes_map = cfg.get("custom_model_codes", {})
                if isinstance(custom_codes_map, dict):
                    for k, v in custom_codes_map.items():
                        if re.sub(r"\s+", "", k.strip()).lower() == cleaned_search:
                            if isinstance(v, list):
                                for c in v:
                                    clean_c = str(c).strip().upper()
                                    if clean_c and clean_c not in codes:
                                        codes.append(clean_c)
            except Exception as e:
                logger.debug("Could not read custom_model_codes from %s: %s", cfg_path, e)

        # 3. If no exact match found, match by model family prefix / substring
        # (e.g. 'Sirius2' matches 'Sirius 2 (21ppm)', 'Sirius 2 (26ppm)';
        #       'Polaris' matches 'Polaris Next', 'Polaris E Model', 'Polaris E Plus';
        #       'Mebius' matches 'Mebius E', 'Mebius E Plus';
        #       '6thA4' matches '6th A4 Next', '6th A4 Plus', '6th A4 E')
        if not codes:
            for m in self._machines:
                m_clean = re.sub(r"\s+", "", m.machine_name.strip()).lower() if m.machine_name else ""
                if m_clean and (m_clean.startswith(cleaned_search) or cleaned_search.startswith(m_clean)):
                    for c in m.machine_codes:
                        clean_c = c.strip().upper()
                        if clean_c and clean_c not in codes:
                            codes.append(clean_c)

            if cfg_path.exists():
                try:
                    with open(cfg_path, "r", encoding="utf-8") as f:
                        cfg = json.load(f)
                    custom_codes_map = cfg.get("custom_model_codes", {})
                    if isinstance(custom_codes_map, dict):
                        for k, v in custom_codes_map.items():
                            k_clean = re.sub(r"\s+", "", k.strip()).lower()
                            if k_clean and (k_clean.startswith(cleaned_search) or cleaned_search.startswith(k_clean)):
                                if isinstance(v, list):
                                    for c in v:
                                        clean_c = str(c).strip().upper()
                                        if clean_c and clean_c not in codes:
                                            codes.append(clean_c)
                except Exception:
                    pass

        return codes

    def get_model_info_by_name(self, model_name: str) -> list[MachineInfo]:
        """Return all MachineInfo entries for this model."""
        if not self._is_loaded:
            self.load()
        cleaned = re.sub(r"\s+", "", model_name.strip()).lower()
        return [m for m in self._machines if re.sub(r"\s+", "", m.machine_name.strip()).lower() == cleaned]
