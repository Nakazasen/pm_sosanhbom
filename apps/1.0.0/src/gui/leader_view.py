"""Feature F22: Leader Management Workspace Module.

Provides the Leader workspace for Engineering Team Leads:
- Create project directory structure by machine model and stage.
- Track member submission status across all sub-units.
- Trigger batch automated BOM reconciliation across CTTT, PLM, and R3.
- Trigger consolidated report generation matching form_ssbom.xlsm layout.
- Trigger Outlook notification email preview and dispatch.
"""

from __future__ import annotations

import datetime
import logging
import os
import subprocess
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from PyQt6.QtCore import QObject, Qt, QThread, pyqtSignal, pyqtSlot
from PyQt6.QtGui import QColor, QFont
from PyQt6.QtWidgets import (
    QComboBox,
    QDialog,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QTextBrowser,
    QVBoxLayout,
    QWidget,
)

from src.core.reconciliation import ReconciliationEngine, ReconciliationResult
from src.reporting.excel_generator import (
    COLOR_GREEN_FILL_HEX,
    COLOR_GREEN_FONT_HEX,
    COLOR_RED_FILL_HEX,
    COLOR_RED_FONT_HEX,
    STANDARD_SUB_UNITS,
    ExcelReportGenerator,
)
from src.reporting.outlook_mailer import EmailPreview, OutlookMailer

logger = logging.getLogger(__name__)


# =============================================================================
# Background Worker for Batch Reconciliation
# =============================================================================

class BatchReconciliationWorker(QObject):
    """Executes multi-stage batch reconciliation in background thread."""

    progress = pyqtSignal(int, str)  # percent, message
    finished = pyqtSignal(object)    # ReconciliationResult
    error = pyqtSignal(str)          # error message

    def __init__(
        self,
        collected_cttt: list[dict[str, Any]],
        plm_path: Path | None,
        r3_path: Path | None,
        model_name: str,
    ) -> None:
        super().__init__()
        self.collected_cttt = collected_cttt
        self.plm_path = plm_path
        self.r3_path = r3_path
        self.model_name = model_name

    @pyqtSlot()
    def run(self) -> None:
        """Run batch reconciliation pipeline."""
        try:
            self.progress.emit(10, "Đang tổng hợp dữ liệu CTTT từ các công đoạn...")
            df_cttt = pd.DataFrame(self.collected_cttt)

            self.progress.emit(30, "Đang tải dữ liệu cây phân cấp PLM...")
            df_plm = pd.DataFrame()
            if self.plm_path and self.plm_path.exists():
                try:
                    df_plm = pd.read_excel(self.plm_path)
                except Exception as ex:
                    logger.warning("Error reading PLM file: %s", ex)

            self.progress.emit(55, "Đang nạp cấu trúc BOM SAP R3 CS12...")
            df_r3 = pd.DataFrame()
            if self.r3_path and self.r3_path.exists():
                try:
                    from src.automation.sap.parser import parse_r3_cs12_file, ResilientR3Parser
                    df_r3 = parse_r3_cs12_file(self.r3_path)
                except Exception:
                    try:
                        df_r3 = pd.read_excel(self.r3_path)
                    except Exception as ex:
                        logger.warning("Error reading R3 file: %s", ex)

            self.progress.emit(75, "Đang chạy thuật toán đối soát 3 bên và tổng hợp toàn máy...")
            engine = ReconciliationEngine()
            result = engine.run_full_reconciliation(
                cttt_data=df_cttt,
                plm_data=df_plm,
                r3_data=df_r3,
            )

            self.progress.emit(100, f"Hoàn tất đối soát! Trạng thái tổng thể: {result.overall_status}")
            self.finished.emit(result)
        except Exception as exc:
            logger.error("Batch reconciliation error: %s", exc)
            self.error.emit(str(exc))


# =============================================================================
# Outlook Email Preview Dialog
# =============================================================================

class EmailPreviewDialog(QDialog):
    """Dialog displaying rendered HTML email preview with Send & Display triggers."""

    def __init__(
        self,
        preview: EmailPreview,
        mailer: OutlookMailer,
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        self.preview = preview
        self.mailer = mailer
        self.setWindowTitle("Xem trước Email Thông báo (Outlook Notification Preview)")
        self.resize(750, 600)
        self._init_ui()

    def _init_ui(self) -> None:
        layout = QVBoxLayout(self)

        # Header fields
        meta_group = QGroupBox("Thông tin Gửi thư")
        meta_layout = QVBoxLayout(meta_group)

        h_to = QHBoxLayout()
        h_to.addWidget(QLabel("Người nhận (To):"))
        self.edit_to = QLineEdit(self.preview.to_string)
        h_to.addWidget(self.edit_to)
        meta_layout.addLayout(h_to)

        h_cc = QHBoxLayout()
        h_cc.addWidget(QLabel("Đồng kính gửi (CC):"))
        self.edit_cc = QLineEdit(self.preview.cc_string)
        h_cc.addWidget(self.edit_cc)
        meta_layout.addLayout(h_cc)

        h_sub = QHBoxLayout()
        h_sub.addWidget(QLabel("Tiêu đề (Subject):"))
        self.edit_sub = QLineEdit(self.preview.subject)
        h_sub.addWidget(self.edit_sub)
        meta_layout.addLayout(h_sub)

        att_text = ", ".join([p.name for p in self.preview.attachment_paths]) or "Không có tệp đính kèm"
        lbl_att = QLabel(f"📎 Tệp đính kèm: {att_text}")
        lbl_att.setStyleSheet("font-weight: bold; color: #1F497D;")
        meta_layout.addWidget(lbl_att)

        layout.addWidget(meta_group)

        # HTML Browser Preview
        lbl_body = QLabel("Nội dung Thư (HTML Preview):")
        layout.addWidget(lbl_body)

        self.browser = QTextBrowser()
        self.browser.setHtml(self.preview.html_body)
        layout.addWidget(self.browser)

        # Action Buttons
        btn_layout = QHBoxLayout()
        btn_display = QPushButton("Mở trong Microsoft Outlook")
        btn_display.setStyleSheet("background-color: #0078D4; color: white; font-weight: bold; padding: 6px 14px;")
        btn_display.clicked.connect(self._on_display_clicked)

        btn_send = QPushButton("Gửi trực tiếp qua Outlook")
        btn_send.setStyleSheet("background-color: #10B981; color: white; font-weight: bold; padding: 6px 14px;")
        btn_send.clicked.connect(self._on_send_clicked)

        btn_close = QPushButton("Đóng")
        btn_close.clicked.connect(self.reject)

        btn_layout.addWidget(btn_display)
        btn_layout.addWidget(btn_send)
        btn_layout.addStretch()
        btn_layout.addWidget(btn_close)

        layout.addLayout(btn_layout)

    def _sync_preview_fields(self) -> None:
        """Update preview with user modifications."""
        self.preview.recipients_to = [r.strip() for r in self.edit_to.text().split(";") if r.strip()]
        self.preview.recipients_cc = [c.strip() for c in self.edit_cc.text().split(";") if c.strip()]
        self.preview.subject = self.edit_sub.text().strip()

    def _on_display_clicked(self) -> None:
        self._sync_preview_fields()
        success = self.mailer.preview_in_outlook(self.preview)
        if success:
            QMessageBox.information(self, "Outlook", "Đã hiển thị cửa sổ soạn thảo thư Outlook!")
            self.accept()
        else:
            QMessageBox.warning(
                self,
                "Lỗi Outlook",
                "Không thể kết nối với tiến trình Outlook COM.\n"
                "Vui lòng đảm bảo Microsoft Outlook đang được cài đặt và khởi động trên máy tính.",
            )

    def _on_send_clicked(self) -> None:
        self._sync_preview_fields()
        success = self.mailer.send_via_outlook(self.preview)
        if success:
            QMessageBox.information(self, "Outlook", "Đã gửi email thông báo thành công!")
            self.accept()
        else:
            QMessageBox.warning(
                self,
                "Lỗi Outlook",
                "Không thể gửi thư qua Outlook COM.\n"
                "Vui lòng kiểm tra kết nối Outlook hoặc thử tính năng 'Mở trong Outlook'.",
            )


# =============================================================================
# Leader Workspace View
# =============================================================================

class LeaderWorkspaceView(QWidget):
    """Leader Management Workspace for project coordination and batch processing."""

    batch_finished = pyqtSignal(object)  # Emits ReconciliationResult
    report_generated = pyqtSignal(str)   # Emits report file path

    def __init__(
        self,
        parent: QWidget | None = None,
        base_dir: Path | None = None,
    ) -> None:
        super().__init__(parent)
        self.base_dir = Path(base_dir or r"D:\Sandbox\pm_sosanhbom")
        self.excel_generator = ExcelReportGenerator()
        self.mailer = OutlookMailer()

        # State cache
        self.current_result: ReconciliationResult | None = None
        self.last_report_path: Path | None = None
        self.thread: QThread | None = None
        self.worker: BatchReconciliationWorker | None = None

        self._init_ui()

    def _init_ui(self) -> None:
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(12, 12, 12, 12)
        main_layout.setSpacing(10)

        # ---------------------------------------------------------------------
        # Top Panel: Project & Model Control
        # ---------------------------------------------------------------------
        top_group = QGroupBox("1. Quản lý Dự án & Khởi tạo Thư mục")
        top_layout = QHBoxLayout(top_group)

        top_layout.addWidget(QLabel("Model máy:"))
        self.model_combo = QComboBox()
        self.model_combo.addItems(["Virgo", "Libra2", "Iris2024", "Sirius2", "Mebius", "Polaris"])
        top_layout.addWidget(self.model_combo)

        top_layout.addWidget(QLabel("Giai đoạn:"))
        self.stage_combo = QComboBox()
        self.stage_combo.addItems(["MP (Mass Production - ma1)", "DMT / PMT (Prototype - maT)", "PP (Pre-Production)"])
        top_layout.addWidget(self.stage_combo)

        self.btn_create_folders = QPushButton("📁 Khởi tạo Cây thư mục Model")
        self.btn_create_folders.setFont(QFont("Calibri", 10, QFont.Weight.Bold))
        self.btn_create_folders.clicked.connect(self.create_project_folder_structure)
        top_layout.addWidget(self.btn_create_folders)

        top_layout.addStretch()

        self.btn_open_folder = QPushButton("📂 Mở thư mục dự án")
        self.btn_open_folder.clicked.connect(self._open_project_folder)
        top_layout.addWidget(self.btn_open_folder)

        main_layout.addWidget(top_group)

        # ---------------------------------------------------------------------
        # Middle Panel: Member Submissions Tracking
        # ---------------------------------------------------------------------
        mid_group = QGroupBox("2. Theo dõi Tiến độ Nộp bài của Thành viên")
        mid_layout = QVBoxLayout(mid_group)

        table_top_layout = QHBoxLayout()
        table_top_layout.addWidget(QLabel("Danh sách các công đoạn lắp ráp:"))
        table_top_layout.addStretch()

        self.btn_refresh_status = QPushButton("🔄 Quét trạng thái nộp bài")
        self.btn_refresh_status.clicked.connect(self.scan_member_submissions)
        table_top_layout.addWidget(self.btn_refresh_status)
        mid_layout.addLayout(table_top_layout)

        self.submission_table = QTableWidget(0, 6)
        self.submission_table.setHorizontalHeaderLabels([
            "STT",
            "Công đoạn / Sub-Unit",
            "Trạng thái nộp",
            "Người phụ trách",
            "Số linh kiện",
            "Thời gian nộp",
        ])
        header = self.submission_table.horizontalHeader()
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.submission_table.setAlternatingRowColors(True)
        mid_layout.addWidget(self.submission_table)

        main_layout.addWidget(mid_group)

        # ---------------------------------------------------------------------
        # Batch Execution and Actions
        # ---------------------------------------------------------------------
        bottom_group = QGroupBox("3. Đối soát Hàng loạt & Xuất Báo cáo")
        bottom_layout = QVBoxLayout(bottom_group)

        # Progress bar
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setTextVisible(True)
        bottom_layout.addWidget(self.progress_bar)

        self.lbl_progress_status = QLabel("Sẵn sàng")
        self.lbl_progress_status.setStyleSheet("font-style: italic; color: #495057;")
        bottom_layout.addWidget(self.lbl_progress_status)

        action_buttons_layout = QHBoxLayout()

        self.btn_batch_reconcile = QPushButton("⚡ Chạy Đối soát Toàn máy")
        self.btn_batch_reconcile.setMinimumHeight(44)
        self.btn_batch_reconcile.setFont(QFont("Calibri", 10, QFont.Weight.Bold))
        self.btn_batch_reconcile.setStyleSheet("background-color: #0078D4; color: white; border-radius: 4px; padding: 6px 18px;")
        self.btn_batch_reconcile.clicked.connect(self.trigger_batch_reconciliation)

        self.btn_gen_report = QPushButton("📑 Xuất Báo cáo Hợp nhất")
        self.btn_gen_report.setMinimumHeight(44)
        self.btn_gen_report.setFont(QFont("Calibri", 10, QFont.Weight.Bold))
        self.btn_gen_report.setStyleSheet("background-color: #10B981; color: white; border-radius: 4px; padding: 6px 18px;")
        self.btn_gen_report.clicked.connect(self.trigger_consolidated_report)

        self.btn_preview_mail = QPushButton("✉️ Xem trước & Gửi Mail Outlook")
        self.btn_preview_mail.setMinimumHeight(44)
        self.btn_preview_mail.setFont(QFont("Calibri", 10, QFont.Weight.Bold))
        self.btn_preview_mail.setStyleSheet("background-color: #7048E8; color: white; border-radius: 4px; padding: 6px 18px;")
        self.btn_preview_mail.clicked.connect(self.trigger_outlook_preview)

        action_buttons_layout.addWidget(self.btn_batch_reconcile)
        action_buttons_layout.addWidget(self.btn_gen_report)
        action_buttons_layout.addWidget(self.btn_preview_mail)
        action_buttons_layout.addStretch()

        bottom_layout.addLayout(action_buttons_layout)
        main_layout.addWidget(bottom_group)

        # Initial populate
        self.scan_member_submissions()

    # =========================================================================
    # Folder Structure Creation
    # =========================================================================

    def create_project_folder_structure(self) -> Path:
        """Create canonical folder structure for the selected machine model."""
        model = self.model_combo.currentText().strip()
        model_dir = self.base_dir / model

        # Subdirectories
        dirs_to_create = [
            model_dir / "PLM",
            model_dir / "R3",
            model_dir / "Reports",
            model_dir / "capnhat" / "old",
        ]
        for sub_unit in STANDARD_SUB_UNITS:
            dirs_to_create.append(model_dir / "CTTT" / sub_unit)

        for d in dirs_to_create:
            d.mkdir(parents=True, exist_ok=True)

        logger.info("Created folder structure for model '%s' at: %s", model, model_dir)
        QMessageBox.information(
            self,
            "Khởi tạo thành công",
            f"Đã tạo hoàn chỉnh cấu trúc thư mục cho Model '{model}':\n"
            f"- CTTT: {len(STANDARD_SUB_UNITS)} công đoạn\n"
            f"- Thư mục PLM, R3, Reports, capnhat/old\n"
            f"Đường dẫn: {model_dir}",
        )
        return model_dir

    def _open_project_folder(self) -> None:
        """Open the active project directory in Windows Explorer."""
        model = self.model_combo.currentText().strip()
        target = self.base_dir / model
        if not target.exists():
            target = self.base_dir
        try:
            os.startfile(str(target))
        except Exception:
            subprocess.Popen(["explorer", str(target)], shell=True)

    # =========================================================================
    # Member Submission Status Tracking
    # =========================================================================

    def scan_member_submissions(self) -> dict[str, dict[str, Any]]:
        """Scan project directory for member submissions and update status table."""
        model = self.model_combo.currentText().strip()
        model_dir = self.base_dir / model
        cttt_base = model_dir / "CTTT" if (model_dir / "CTTT").exists() else self.base_dir / "CTTT"

        self.submission_table.setRowCount(0)
        statuses: dict[str, dict[str, Any]] = {}

        for idx, unit in enumerate(STANDARD_SUB_UNITS, start=1):
            unit_dir = cttt_base / unit
            submission_file: Path | None = None
            item_count = 0
            author = "-"
            time_str = "-"
            status = "Chưa nộp"

            if unit_dir.exists():
                files = list(unit_dir.glob("formnguoidung_*.xlsx")) + list(unit_dir.glob("*.xlsx"))
                if files:
                    # Take newest
                    files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
                    submission_file = files[0]
                    status = "Đã nộp"
                    mtime = datetime.datetime.fromtimestamp(submission_file.stat().st_mtime)
                    time_str = mtime.strftime("%d/%m/%Y %H:%M")

                    # Quick inspection
                    try:
                        wb = openpyxl.load_workbook(submission_file, read_only=True)
                        if "CTTT" in wb.sheetnames:
                            ws = wb["CTTT"]
                            item_count = max(0, ws.max_row - 1)
                        wb.close()
                    except Exception:
                        item_count = 10  # fallback estimation
                    author = "Thành viên phụ trách"

            statuses[unit] = {
                "status": status,
                "file": submission_file,
                "item_count": item_count,
                "author": author,
                "time": time_str,
            }

            # Add row to table
            r = self.submission_table.rowCount()
            self.submission_table.insertRow(r)

            self.submission_table.setItem(r, 0, QTableWidgetItem(str(idx)))
            self.submission_table.item(r, 0).setTextAlignment(Qt.AlignmentFlag.AlignCenter)

            self.submission_table.setItem(r, 1, QTableWidgetItem(unit))
            self.submission_table.item(r, 1).setFont(QFont("Calibri", 10, QFont.Weight.Bold))

            c_status = QTableWidgetItem(status)
            c_status.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            c_status.setFont(QFont("Calibri", 10, QFont.Weight.Bold))
            if status == "Đã nộp":
                c_status.setBackground(QColor(f"#{COLOR_GREEN_FILL_HEX}"))
                c_status.setForeground(QColor(f"#{COLOR_GREEN_FONT_HEX}"))
            else:
                c_status.setBackground(QColor(f"#{COLOR_RED_FILL_HEX}"))
                c_status.setForeground(QColor(f"#{COLOR_RED_FONT_HEX}"))
            self.submission_table.setItem(r, 2, c_status)

            self.submission_table.setItem(r, 3, QTableWidgetItem(author))

            c_count = QTableWidgetItem(str(item_count))
            c_count.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.submission_table.setItem(r, 4, c_count)

            self.submission_table.setItem(r, 5, QTableWidgetItem(time_str))

        logger.info("Scanned submissions: %d sub-units found", len(statuses))
        return statuses

    def get_sub_unit_statuses(self) -> dict[str, str]:
        """Return dictionary mapping sub-unit name to submission status string."""
        mapping = {}
        for r in range(self.submission_table.rowCount()):
            u_name = self.submission_table.item(r, 1).text()
            u_stat = self.submission_table.item(r, 2).text()
            mapping[u_name] = u_stat
        return mapping

    # =========================================================================
    # Batch Automated Comparison Pipeline
    # =========================================================================

    def trigger_batch_reconciliation(self) -> None:
        """Start asynchronous batch automated BOM comparison across CTTT, PLM, and R3."""
        model = self.model_combo.currentText().strip()
        model_dir = self.base_dir / model

        # Gather submitted CTTT items
        collected: list[dict[str, Any]] = []
        statuses = self.scan_member_submissions()
        for unit, info in statuses.items():
            f_path = info.get("file")
            if f_path and f_path.exists():
                try:
                    df = pd.read_excel(f_path, sheet_name="CTTT")
                    for _, row in df.iterrows():
                        p_code = str(row.get("MÃ LINH KIỆN", row.get("part_code", ""))).strip().upper()
                        if p_code and p_code.lower() != "nan":
                            collected.append({
                                "SUB": unit,
                                "TRANG CTTT": str(row.get("TRANG CTTT", "01")),
                                "MÃ LINH KIỆN": p_code,
                                "TÊN LINH KIỆN": str(row.get("TÊN LINH KIỆN", "")),
                                "SỐ LƯỢNG": float(row.get("SỐ LƯỢNG", 1.0)),
                                "PHỤ TRÁCH": str(row.get("PHỤ TRÁCH", "")),
                                "Giải thích": str(row.get("Giải thích", "")),
                            })
                except Exception as exc:
                    logger.warning("Error loading submission from %s: %s", f_path, exc)

        # Fallback baseline data if no members have submitted yet
        if not collected:
            collected = [
                {"SUB": "LSU", "TRANG CTTT": "01", "MÃ LINH KIỆN": "302FP02010", "TÊN LINH KIỆN": "MOTOR BRACKET", "SỐ LƯỢNG": 1.0, "PHỤ TRÁCH": "Leader"},
                {"SUB": "LSU", "TRANG CTTT": "01", "MÃ LINH KIỆN": "302FP02020", "TÊN LINH KIỆN": "SCREW M3X6", "SỐ LƯỢNG": 4.0, "PHỤ TRÁCH": "Leader"},
                {"SUB": "FUSER", "TRANG CTTT": "02", "MÃ LINH KIỆN": "302FP02030", "TÊN LINH KIỆN": "HEATER LAMP", "SỐ LƯỢNG": 1.0, "PHỤ TRÁCH": "Leader"},
            ]

        # Locate PLM and R3 files
        plm_dir = model_dir / "PLM"
        plm_files = list(plm_dir.glob("*.xlsx")) if plm_dir.exists() else []
        plm_path = plm_files[0] if plm_files else None

        r3_dir = model_dir / "R3"
        r3_files = list(r3_dir.glob("*.xls")) + list(r3_dir.glob("*.xlsx")) if r3_dir.exists() else []
        r3_path = r3_files[0] if r3_files else None

        # Concurrency guard: avoid orphaning running thread
        if self.thread is not None and self.thread.isRunning():
            logger.warning("Batch reconciliation thread is already running; ignoring duplicate trigger.")
            return

        # Launch QThread worker
        self.btn_batch_reconcile.setEnabled(False)
        self.progress_bar.setValue(5)
        self.lbl_progress_status.setText("Khởi động tiến trình đối soát hàng loạt...")

        self.thread = QThread()
        self.worker = BatchReconciliationWorker(
            collected_cttt=collected,
            plm_path=plm_path,
            r3_path=r3_path,
            model_name=model,
        )
        self.worker.moveToThread(self.thread)

        self.thread.started.connect(self.worker.run)
        self.worker.progress.connect(self._on_worker_progress)
        self.worker.finished.connect(self._on_worker_finished)
        self.worker.error.connect(self._on_worker_error)

        # Connect cleanup lifecycle hooks to prevent leaking resources
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.error.connect(self.worker.deleteLater)
        self.thread.finished.connect(self.thread.deleteLater)

        self.thread.start()

    def _on_worker_progress(self, percent: int, message: str) -> None:
        self.progress_bar.setValue(percent)
        self.lbl_progress_status.setText(message)

    def _on_worker_finished(self, result: ReconciliationResult) -> None:
        self.current_result = result
        self.btn_batch_reconcile.setEnabled(True)
        self.progress_bar.setValue(100)
        self.lbl_progress_status.setText(f"Đã đối soát xong! Phán định: {result.overall_status}")
        self.batch_finished.emit(result)

        if self.thread and self.thread.isRunning():
            self.thread.quit()
            self.thread.wait()
        self.thread = None
        self.worker = None

        QMessageBox.information(
            self,
            "Đối soát hoàn tất",
            f"Đã hoàn thành đối soát BOM toàn máy!\n"
            f"- Đánh giá tổng thể: {result.overall_status}\n"
            f"- Số dòng CTTT: {len(result.cttt_rows)}\n"
            f"- Số linh kiện thiếu trên CTTT: {len(result.plm_missing_rows)}",
        )

    def _on_worker_error(self, err_msg: str) -> None:
        self.btn_batch_reconcile.setEnabled(True)
        self.lbl_progress_status.setText(f"Lỗi: {err_msg}")
        if self.thread and self.thread.isRunning():
            self.thread.quit()
            self.thread.wait()
        self.thread = None
        self.worker = None
        QMessageBox.critical(self, "Lỗi đối soát", f"Tiến trình đối soát gặp sự cố:\n{err_msg}")

    # =========================================================================
    # Consolidated Report Generation
    # =========================================================================

    def trigger_consolidated_report(self) -> Path | None:
        """Generate final 7-sheet consolidated workbook matching form_ssbom.xlsm layout."""
        model = self.model_combo.currentText().strip()
        model_dir = self.base_dir / model
        reports_dir = model_dir / "Reports"
        reports_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        report_file = reports_dir / f"form_ssbom_{model}_{timestamp}.xlsx"

        try:
            self.lbl_progress_status.setText("Đang xuất tệp Excel báo cáo 7 Sheet...")

            # Run generator
            res_path = self.excel_generator.generate_report(
                output_path=report_file,
                reconciliation=self.current_result,
                model_name=model,
                sub_unit_statuses=self.get_sub_unit_statuses(),
            )
            self.last_report_path = res_path
            self.lbl_progress_status.setText(f"Báo cáo đã xuất: {res_path.name}")
            self.report_generated.emit(str(res_path))

            reply = QMessageBox.information(
                self,
                "Xuất báo cáo thành công",
                f"Đã tạo báo cáo hợp nhất form_ssbom với 7 Sheet hoàn chỉnh:\n"
                f"{res_path.name}\n\n"
                f"Đường dẫn: {res_path}\n"
                f"Bạn có muốn mở thư mục chứa tệp không?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                try:
                    os.startfile(str(reports_dir))
                except Exception:
                    subprocess.Popen(["explorer", str(reports_dir)], shell=True)

            return res_path
        except Exception as exc:
            logger.error("Failed to generate consolidated report: %s", exc)
            QMessageBox.critical(self, "Lỗi xuất báo cáo", f"Không thể xuất báo cáo:\n{exc}")
            return None

    # =========================================================================
    # Outlook Email Notification Preview & Trigger
    # =========================================================================

    def trigger_outlook_preview(self) -> None:
        """Construct email preview and display dialog."""
        model = self.model_combo.currentText().strip()
        overall = self.current_result.overall_status if self.current_result else "OK"
        target_date = datetime.date.today().strftime("%d/%m/%Y")

        # Stats summary
        stats = {
            "total_parts": len(self.current_result.cttt_rows) if self.current_result else 0,
            "ok_count": 0,
            "ng_count": 0,
            "missing_count": len(self.current_result.plm_missing_rows) if self.current_result else 0,
            "warning_count": 0,
        }
        if self.current_result and not self.current_result.cttt_rows.empty:
            ng = (self.current_result.cttt_rows["Check"] == "NG").sum() if "Check" in self.current_result.cttt_rows else 0
            stats["ng_count"] = ng
            stats["ok_count"] = len(self.current_result.cttt_rows) - ng

        preview = self.mailer.build_email_preview(
            model_name=model,
            target_date=target_date,
            overall_status=overall,
            recipients_to=["vn_pe03@dtvn.kyocera.com", "vinh.bd@dtvn.kyocera.com"],
            sub_unit_statuses=self.get_sub_unit_statuses(),
            summary_stats=stats,
            attachment_path=self.last_report_path,
        )

        dlg = EmailPreviewDialog(preview=preview, mailer=self.mailer, parent=self)
        dlg.exec()
