"""TC2412 Canonical Normalizer and Sheet PLM Excel Bridge.

Conforms strictly to Section 13 of specs/SPEC_PLM_AUTO_DOWNLOAD.md:
- Supports both 14-column canonical PLM files and 24-column TC2412 new format files.
- Protects downstream Excel formulas in master workbooks (e.g. BOM_110C0Z3LV1.xlsm):
  1. Tongket!C5 & CTTT!A1: `=PLM!C2` -> Col C MUST be Root Machine Part Code (item_id).
  2. CTTT!I3: `=VLOOKUP(C3, PLM!C:L, 10, 0)` -> Col L MUST be Revision (10th col from C).
  3. PLM!R2: `=IF(C2="","",C2)` (PART CODE).
  4. PLM!S2: `=IF(E2="","",E2)` (Q.TY).
  5. CTTT!G3: `=VLOOKUP(C3, PLM!T:U, 2, 0)` -> Cols T:U reserved for Pivot Table summary.
"""

from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Union

import openpyxl
import pandas as pd

logger = logging.getLogger(__name__)


class TC2412CanonicalNormalizer:
    """Normalizes raw PLM BOM records from either 14-column canonical or 24-column TC2412 exports."""

    @classmethod
    def is_tc2412_24col(cls, headers: Sequence[Any]) -> bool:
        """Detect if headers belong to the new 24-column Teamcenter 2412 export format."""
        normalized = [str(h or "").strip().lower() for h in headers]
        has_name = "name" in normalized
        has_parts_text = any("parts text" in h or "partstext" in h for h in normalized)
        return has_name and has_parts_text

    @classmethod
    def normalize_row(cls, raw_row: dict[str, Any]) -> dict[str, Any]:
        """Normalize a single raw record into the canonical schema dictionary."""
        # Detect if raw_row uses 24-column TC2412 headers
        keys_lower = {str(k).strip().lower(): k for k in raw_row.keys()}

        is_24col = "name" in keys_lower and ("parts text" in keys_lower or "partstext" in keys_lower)

        if is_24col:
            # TC2412 24-column mapping
            name_key = keys_lower.get("name")
            parts_text_key = keys_lower.get("parts text") or keys_lower.get("partstext")
            rel_status_key = keys_lower.get("release status") or keys_lower.get("releasestatus")
            rev_key = keys_lower.get("revision") or keys_lower.get("rev")
            level_key = keys_lower.get("level") or keys_lower.get("lv")
            type_key = keys_lower.get("item type") or keys_lower.get("itemtype")
            has_child_key = keys_lower.get("has children") or keys_lower.get("haschildren")
            qty_key = keys_lower.get("quantity") or keys_lower.get("qty")
            occ_eff_key = keys_lower.get("occurrence effectivities") or keys_lower.get("occurrenceeffectivities")
            proj_key = keys_lower.get("item revision project list") or keys_lower.get("itemrevisionprojectlist")

            return {
                "home": raw_row.get(keys_lower.get("home", ""), ""),
                "level": raw_row.get(level_key, 0) if level_key else 0,
                "item_type": raw_row.get(type_key, "") if type_key else "",
                "item_id": str(raw_row.get(name_key, "") or "").strip() if name_key else "",
                "has_children": raw_row.get(has_child_key, False) if has_child_key else False,
                "quantity": raw_row.get(qty_key, 1) if qty_key else 1,
                "item_name": str(raw_row.get(parts_text_key, "") or "").strip() if parts_text_key else "",
                "revision": str(raw_row.get(rev_key, "") or "").strip() if rev_key else "",
                "item_rev_status": str(raw_row.get(rel_status_key, "") or "").strip() if rel_status_key else "",
                "occurrence_effectivities": str(raw_row.get(occ_eff_key, "") or "").strip() if occ_eff_key else "",
                "project_list": str(raw_row.get(proj_key, "") or "").strip() if proj_key else "",
                "first_parts": str(raw_row.get(keys_lower.get("1st parts", ""), "") or "").strip(),
                "second_bom_flag": str(raw_row.get(keys_lower.get("2nd bom flag", ""), "") or "").strip(),
                "notice_no": str(raw_row.get(keys_lower.get("notice no", ""), "") or "").strip(),
                "_raw": raw_row,
            }

        # Canonical 14-column mapping
        item_id_key = keys_lower.get("item id") or keys_lower.get("itemid") or keys_lower.get("partcode") or keys_lower.get("item_id")
        item_name_key = keys_lower.get("item name") or keys_lower.get("itemname") or keys_lower.get("name") or keys_lower.get("partname")
        rev_key = keys_lower.get("revision") or keys_lower.get("rev")
        status_key = keys_lower.get("item rev status") or keys_lower.get("itemrevstatus") or keys_lower.get("status")
        level_key = keys_lower.get("level") or keys_lower.get("lv")
        type_key = keys_lower.get("item type") or keys_lower.get("itemtype")
        has_child_key = keys_lower.get("has children") or keys_lower.get("haschildren")
        qty_key = keys_lower.get("quantity") or keys_lower.get("qty")
        first_p_key = keys_lower.get("1st parts") or keys_lower.get("1stparts")
        second_bf_key = keys_lower.get("2nd bom flag") or keys_lower.get("2ndbomflag")
        occ_eff_key = keys_lower.get("occurrence effectivities") or keys_lower.get("occurrenceeffectivities")
        proj_key = keys_lower.get("item revision project list") or keys_lower.get("itemrevisionprojectlist")
        notice_key = keys_lower.get("notice no") or keys_lower.get("noticeno")
        home_key = keys_lower.get("home")

        return {
            "home": raw_row.get(home_key, "") if home_key else "",
            "level": raw_row.get(level_key, 0) if level_key else 0,
            "item_type": raw_row.get(type_key, "") if type_key else "",
            "item_id": str(raw_row.get(item_id_key, "") or "").strip() if item_id_key else "",
            "has_children": raw_row.get(has_child_key, False) if has_child_key else False,
            "quantity": raw_row.get(qty_key, 1) if qty_key else 1,
            "first_parts": str(raw_row.get(first_p_key, "") or "").strip() if first_p_key else "",
            "second_bom_flag": str(raw_row.get(second_bf_key, "") or "").strip() if second_bf_key else "",
            "occurrence_effectivities": str(raw_row.get(occ_eff_key, "") or "").strip() if occ_eff_key else "",
            "project_list": str(raw_row.get(proj_key, "") or "").strip() if proj_key else "",
            "item_name": str(raw_row.get(item_name_key, "") or "").strip() if item_name_key else "",
            "notice_no": str(raw_row.get(notice_key, "") or "").strip() if notice_key else "",
            "revision": str(raw_row.get(rev_key, "") or "").strip() if rev_key else "",
            "item_rev_status": str(raw_row.get(status_key, "") or "").strip() if status_key else "",
            "_raw": raw_row,
        }

    @classmethod
    def normalize_dataframe(cls, df: pd.DataFrame) -> pd.DataFrame:
        """Convert any PLM DataFrame (14-column or 24-column) to a Canonical DataFrame."""
        if df.empty:
            return pd.DataFrame(columns=[
                "level", "item_type", "item_id", "has_children", "quantity",
                "first_parts", "second_bom_flag", "occurrence_effectivities",
                "project_list", "item_name", "notice_no", "revision", "item_rev_status", "home"
            ])

        records = df.to_dict(orient="records")
        normalized = [cls.normalize_row(r) for r in records]
        return pd.DataFrame(normalized)


class TC2412SheetPLMBridge:
    """Bridges normalized TC2412 data into Sheet 'PLM' while strictly preserving Excel formulas."""

    SHEET_PLM_COLUMNS = [
        "Level",                     # Col A
        "Item Type",                  # Col B
        "Item Id",                    # Col C  (CRITICAL: Root in C2 -> =PLM!C2, start of VLOOKUP C:L)
        "Has Children",               # Col D
        "Quantity",                   # Col E  (CRITICAL: =IF(E2="","",E2) in S2)
        "1st Parts",                  # Col F
        "2nd BOM Flag",               # Col G
        "Occurrence Effectivities",   # Col H
        "Item Revision Project List", # Col I
        "Item Name",                  # Col J  (Japanese Kanji font MS Gothic)
        "Notice No",                  # Col K
        "Revision",                   # Col L  (CRITICAL: Col 10 in VLOOKUP(C3, PLM!C:L, 10, 0))
        "Item Rev Status",            # Col M
        "Home",                       # Col N
    ]

    @classmethod
    def populate_sheet_plm(
        cls,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        records: list[dict[str, Any]],
        preserve_formulas: bool = True,
    ) -> int:
        """Write canonical records to worksheet starting at Row 2 while preserving formula columns.

        Cols R (PART CODE) and S (Q.TY) will be populated with standard dynamic formulas:
        - Col R: `=IF(C{r}="","",C{r})`
        - Col S: `=IF(E{r}="","",E{r})`
        Cols T:U are left untouched for Pivot Table summaries.

        Args:
            ws: openpyxl Worksheet for Sheet 'PLM'.
            records: List of normalized dictionary records.
            preserve_formulas: If True, writes standard formulas into Cols R and S.

        Returns:
            int: Number of rows populated.
        """
        # Set Headers on Row 1 for Cols A-N
        for col_idx, col_name in enumerate(cls.SHEET_PLM_COLUMNS, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)

        if preserve_formulas:
            ws.cell(row=1, column=18, value="PART CODE")  # Col R
            ws.cell(row=1, column=19, value="Q.TY")       # Col S

        for row_idx, r in enumerate(records, start=2):
            ws.cell(row=row_idx, column=1, value=r.get("level", 0))
            ws.cell(row=row_idx, column=2, value=r.get("item_type", ""))
            ws.cell(row=row_idx, column=3, value=r.get("item_id", ""))
            ws.cell(row=row_idx, column=4, value=r.get("has_children", False))
            ws.cell(row=row_idx, column=5, value=r.get("quantity", 1))
            ws.cell(row=row_idx, column=6, value=r.get("first_parts", ""))
            ws.cell(row=row_idx, column=7, value=r.get("second_bom_flag", ""))
            ws.cell(row=row_idx, column=8, value=r.get("occurrence_effectivities", ""))
            ws.cell(row=row_idx, column=9, value=r.get("project_list", ""))
            ws.cell(row=row_idx, column=10, value=r.get("item_name", ""))
            ws.cell(row=row_idx, column=11, value=r.get("notice_no", ""))
            ws.cell(row=row_idx, column=12, value=r.get("revision", ""))
            ws.cell(row=row_idx, column=13, value=r.get("item_rev_status", ""))
            ws.cell(row=row_idx, column=14, value=r.get("home", ""))

            if preserve_formulas:
                ws.cell(row=row_idx, column=18, value=f'=IF(C{row_idx}="","",C{row_idx})')
                ws.cell(row=row_idx, column=19, value=f'=IF(E{row_idx}="","",E{row_idx})')

        return len(records)

    @classmethod
    def update_workbook_sheet_plm(
        cls,
        target_workbook_path: Union[str, Path],
        source_plm_path: Union[str, Path],
        backup: bool = True,
    ) -> Path:
        """Read source PLM file (14-col or 24-col) and write to target workbook Sheet 'PLM'.

        Creates a backup copy before modifying the target workbook.

        Args:
            target_workbook_path: Path to master workbook (e.g. BOM_110C0Z3LV1.xlsm).
            source_plm_path: Path to exported TC2412 file.
            backup: Whether to create a .bak copy before modifying.

        Returns:
            Path to the updated workbook.
        """
        target_path = Path(target_workbook_path).resolve()
        source_path = Path(source_plm_path).resolve()

        if not target_path.exists():
            raise FileNotFoundError(f"Target workbook does not exist: {target_path}")
        if not source_path.exists():
            raise FileNotFoundError(f"Source PLM file does not exist: {source_path}")

        # Create safe backup
        if backup:
            bak_path = target_path.with_suffix(target_path.suffix + ".bak")
            shutil.copy2(target_path, bak_path)
            logger.info("Created backup of master workbook at: %s", bak_path)

        # Read source data
        src_wb = openpyxl.load_workbook(filename=str(source_path), data_only=True)
        src_ws = src_wb.active
        if src_ws is None:
            raise ValueError(f"Source file has no active worksheet: {source_path}")

        raw_headers = [str(src_ws.cell(row=1, column=c).value or "").strip() for c in range(1, src_ws.max_column + 1)]
        raw_rows = []
        for r in range(2, src_ws.max_row + 1):
            row_dict = {}
            for c_idx, h in enumerate(raw_headers, start=1):
                row_dict[h] = src_ws.cell(row=r, column=c_idx).value
            raw_rows.append(row_dict)
        src_wb.close()

        # Normalize rows
        normalized_records = [TC2412CanonicalNormalizer.normalize_row(row) for row in raw_rows]

        # Update target workbook
        tgt_wb = openpyxl.load_workbook(filename=str(target_path), keep_vba=target_path.suffix.lower() == ".xlsm")

        if "PLM" in tgt_wb.sheetnames:
            plm_ws = tgt_wb["PLM"]
        else:
            plm_ws = tgt_wb.create_sheet("PLM")

        cls.populate_sheet_plm(plm_ws, normalized_records, preserve_formulas=True)
        tgt_wb.save(str(target_path))
        tgt_wb.close()

        logger.info(
            "Successfully bridged %d normalized TC2412 records into Sheet 'PLM' of %s.",
            len(normalized_records),
            target_path.name,
        )
        return target_path

