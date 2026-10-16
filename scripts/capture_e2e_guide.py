"""End-to-end capture driver: compares BOM for machine 110C3M3AK0 (C3MAKO / Virgo)
by driving the real SSBOM GUI (no mocked widget contents).

Inputs come from the sample package (C3MAKO - Copy.zip, extracted to sample_data/C3MAKO - Copy):
  - PLM_110C3M3AK0.xlsx            -> real PLM file
  - R3 sheet of BOM_110C3M3AK0.xlsm -> exported as R3_110C3M3AK0.xlsx (real SAP CS12 output)
  - CTTT sheet of BOM_110C3M3AK0.xlsm -> split per engineer; entered through the Member screen.

Run:  python scripts/capture_e2e_guide.py
"""

from __future__ import annotations

import collections
import json
import os
import shutil
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ["QT_QPA_PLATFORM"] = "windows"
os.environ.setdefault("SSBOM_DISABLE_EXCEL_COM", "1")

import openpyxl
from PyQt6.QtCore import QTimer
from PyQt6.QtWidgets import QApplication, QCheckBox, QComboBox, QFileDialog, QMessageBox

SAMPLE = ROOT / "sample_data" / "C3MAKO - Copy"
RUN = ROOT / "sample_data" / "run"
OUT = ROOT / "docs" / "images_e2e"
MACHINE = "110C3M3AK0"
MODEL = "Virgo"
CFG = ROOT / "config" / "settings.json"
LOG: list[str] = []
W, H = 1500, 900


def shot(widget, name: str, w: int = W, h: int = H) -> None:
    from PyQt6.QtCore import QEventLoop
    widget.resize(w, h)
    widget.show()
    loop = QEventLoop()
    QTimer.singleShot(400, loop.quit)
    loop.exec()
    for _ in range(4):
        QApplication.processEvents()
    OUT.mkdir(parents=True, exist_ok=True)
    pm = widget.grab()
    pm.save(str(OUT / name), "PNG")
    print(f"[shot] {name} {pm.width()}x{pm.height()}")


def patch_dialogs(file_queue: list[str]) -> None:
    def rec(kind):
        def _f(parent, title="", text="", *a, **k):
            LOG.append(f"[{kind}] {title} | {str(text)[:300]}")
            return QMessageBox.StandardButton.Yes if kind == "question" else QMessageBox.StandardButton.Ok
        return _f

    QMessageBox.information = staticmethod(rec("info"))
    QMessageBox.warning = staticmethod(rec("warn"))
    QMessageBox.critical = staticmethod(rec("crit"))
    QMessageBox.question = staticmethod(rec("question"))

    QFileDialog.getOpenFileNames = staticmethod(lambda *a, **k: (list(file_queue), ""))
    QFileDialog.getOpenFileName = staticmethod(lambda *a, **k: ((file_queue[0] if file_queue else ""), ""))


def prepare_inputs() -> dict[str, list[tuple]]:
    """Build R3 xlsx + split CTTT by engineer from the finished BOM sample."""
    if RUN.exists():
        shutil.rmtree(RUN)
    (RUN / "_inputs").mkdir(parents=True)
    shutil.copy2(ROOT / "formnguoidung.xlsm", RUN / "formnguoidung.xlsm")
    shutil.copy2(SAMPLE / f"PLM_{MACHINE}.xlsx", RUN / "_inputs" / f"PLM_{MACHINE}.xlsx")

    wb = openpyxl.load_workbook(SAMPLE / f"BOM_{MACHINE}.xlsm", read_only=True, data_only=True)
    # R3 sheet -> xlsx (values only, same layout as SAP CS12 export)
    out = openpyxl.Workbook()
    ws_o = out.active
    ws_o.title = "Sheet1"
    for row in wb["R3"].iter_rows(values_only=True):
        ws_o.append(list(row))
    out.save(RUN / "_inputs" / f"R3_{MACHINE}.xlsx")

    per_eng: dict[str, list[tuple]] = collections.defaultdict(list)
    for r in wb["CTTT"].iter_rows(min_row=3, values_only=True):
        if r[2] and r[5]:
            per_eng[str(r[5]).strip()].append((r[0], r[1], r[2], r[3], r[4]))
    wb.close()
    # merge case-variant duplicates (e.g. Phamminhthai_mecha1 -> PhamMinhThai_mecha1)
    merged: dict[str, list[tuple]] = {}
    for k, v in per_eng.items():
        key = next((m for m in merged if m.lower() == k.lower()), k)
        merged.setdefault(key, []).extend(v)
    return merged


def main() -> None:
    cfg_backup = CFG.read_bytes() if CFG.exists() else None
    app = QApplication.instance() or QApplication(sys.argv)
    try:
        run(app)
    finally:
        if cfg_backup is not None:
            CFG.write_bytes(cfg_backup)
        Path(RUN / "dialog_log.txt").write_text("\n".join(LOG), encoding="utf-8")
        print("\n".join(LOG[-40:]))


# Account names stored in the sample CTTT are the old naming; map to current roster accounts.
ACCOUNT_MAP = {
    "AnQuyDuong_mecha1": "Duong_mecha1.2",
    "DangThiLy_mecha1": "Ly_mecha1.2",
    "HoangThiNhung_mecha1": "Nhung_mecha1.2",
    "NguyenVanHung_mecha1": "Hung_mecha1.1",
    "NguyenXuanViet_mecha1": "VietE_mecha1.2",
    "PhamHuuTu_mecha1": "Tu_mecha1.2",
    "PhamMinhThai_mecha1": "Thai_mecha1.2",
    "VuVanChien_mecha1": "Chien_mecha1.2",
    "VuvanThuong_mecha1": "Thuong_mecha1.2",
}


def tree_text(root: Path, depth: int = 3) -> str:
    lines = []
    for p in sorted(root.rglob("*")):
        rel = p.relative_to(root)
        if len(rel.parts) > depth or "_inputs" in rel.parts:
            continue
        lines.append("  " * (len(rel.parts) - 1) + ("📁 " if p.is_dir() else "📄 ") + p.name)
    return "\n".join(lines)


def run(app) -> None:
    from src.gui.app import SSBOMMainWindow
    from src.gui.styles import get_theme_manager

    get_theme_manager().set_theme("light", apply_stylesheet=True)
    per_eng = prepare_inputs()
    print("engineers:", {k: len(v) for k, v in per_eng.items()})
    files: list[str] = []
    patch_dialogs(files)

    win = SSBOMMainWindow()
    leader = win.leader_view
    s1, s2, s3, s4 = leader.step1_widget, leader.step2_widget, leader.step3_widget, leader.step4_widget
    win.resize(W, H)
    win.show()
    app.processEvents()
    if hasattr(win, "tabs"):
        win.tabs.setCurrentIndex(0)

    # ---------------- STEP 1: project + staffing ----------------
    s1.edit_base_dir.setText(str(RUN))
    s1.state.base_dir = RUN
    s1.model_combo.setCurrentText(MODEL)
    s1.stage_combo.setCurrentIndex(0)
    s1.machine_table.setRowCount(0)
    s1._add_machine_row(code=MACHINE, model_name=MODEL)
    s1._reindex_machine_stt()
    s1._set_all_staff_checked(False)
    wanted = {v for v in ACCOUNT_MAP.values()}
    for r, a in enumerate(s1.state.staff_roster):
        if a.engineer_name in wanted:
            chk = s1.staff_table.cellWidget(r, 0).findChild(QCheckBox)
            chk.setChecked(True)
    s1.chk_only_assigned_model.setChecked(True)
    s1._filter_staff_table()
    s1.sync_state_from_ui()
    shot(win, "e01_step1_setup.png")

    s1.execute_create_folders_and_packages()
    leader.refresh_kpi_cards()
    shot(win, "e02_step1_created.png")
    (RUN / "tree_after_step1.txt").write_text(tree_text(RUN), encoding="utf-8")

    # ---------------- STEP 2: sources + filter ----------------
    leader.switch_to_step(1)
    shot(win, "e03_step2_before.png")
    files[:] = [str(RUN / "_inputs" / f"PLM_{MACHINE}.xlsx"), str(RUN / "_inputs" / f"R3_{MACHINE}.xlsx")]
    s2._import_local_bom_files()
    shot(win, "e04_step2_imported.png")

    m = [x for x in s1.state.machines if x.machine_code == MACHINE][0]
    wb = openpyxl.load_workbook(m.plm_file, read_only=True)
    rows_before = wb.worksheets[0].max_row
    wb.close()
    s2.execute_bom_filtering()
    wb = openpyxl.load_workbook(m.plm_file, read_only=True)
    rows_after = wb.worksheets[0].max_row
    wb.close()
    print(f"PLM rows before filter={rows_before} after={rows_after}")
    shot(win, "e05_step2_filtered.png")
    (RUN / "tree_after_step2.txt").write_text(tree_text(RUN), encoding="utf-8")

    # ---------------- STEP 3: member work (real Member screen) ----------------
    leader.switch_to_step(2)
    s3.scan_submissions()
    print("blank templates: ok", sum(1 for x in s3.state.submissions if x.is_submitted_ok), "/",
          len(s3.state.submissions), "all_ok", s3.state.all_members_ok)
    member = win.member_view
    first = True
    for old, rows in per_eng.items():
        acct = ACCOUNT_MAP[old]
        f = RUN / MODEL / MACHINE / f"{acct}.xlsm"
        if not f.exists():
            print("MISSING member file", f)
            continue
        win.tabs.setCurrentIndex(1)
        member.open_assignment_package(f)
        member.clear_cttt_table()
        for sub, page, code, name, qty in rows:
            member.add_cttt_row(
                page=str(page or "01"), part_code=str(code), part_name=str(name or ""),
                quantity=qty or 1, sub_unit=str(sub or "").strip(),
            )
        if first:
            shot(win, "e07_member_entered.png")
        try:
            member.btn_self_check.click()
        except Exception as ex:
            print("self-check error", ex)
        if first:
            shot(win, "e08_member_selfcheck.png")
        member.submit_data()
        if first:
            shot(win, "e09_member_submitted.png")
            first = False
    win.tabs.setCurrentIndex(0)
    leader.switch_to_step(2)
    s3.scan_submissions()
    shot(win, "e10_step3_phantom_rows.png")
    # (workaround removed: empty CTTT\<unit> folders no longer block the gate)
    s3.scan_submissions()
    print("after members: ok", sum(1 for x in s3.state.submissions if x.is_submitted_ok), "/", len(s3.state.submissions),
          "all_ok", s3.state.all_members_ok)
    print("pending:", [(x.sub_unit, x.engineer_name) for x in s3.state.submissions if not x.is_submitted_ok])
    shot(win, "e11_step3_all_ok.png")
    s3.execute_consolidation()
    shot(win, "e12_step3_consolidated.png")
    print("consolidated rows:", len(s3.state.consolidated_cttt), "step3_completed:", s3.state.step3_completed)
    (RUN / "tree_after_step3.txt").write_text(tree_text(RUN), encoding="utf-8")

    # ---------------- STEP 4 ----------------
    leader.switch_to_step(3)
    shot(win, "e13_step4_before.png")
    out = s4.generate_master_bom_file()
    print("BOM file:", out, out.stat().st_size if out else None)
    shot(win, "e14_step4_generated.png")
    (RUN / "tree_after_step4.txt").write_text(tree_text(RUN), encoding="utf-8")


if __name__ == "__main__":
    main()

