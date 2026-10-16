import sys
import openpyxl
from pathlib import Path

if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

excel_path = Path(r"\\fstvn01\Data\10_Production Engineering Department(製造技術部)\02.製造技術課\PE Dept\15. FORM（BIEU MAU）-形式\Form_VBA\Form_VBA_SS_BOM\tonghop_new12052026_ma1.xlsm")
wb = openpyxl.load_workbook(str(excel_path), data_only=True)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\n=== Sheet: {sheet} ===")
    for row in list(ws.iter_rows(values_only=True))[:10]:
        if any(c is not None for c in row):
            # Print non-None
            print("  ", [str(c) for c in row if c is not None][:6])
