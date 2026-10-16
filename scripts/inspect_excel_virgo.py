import sys
import json
from pathlib import Path
import openpyxl

# Set console output encoding to utf-8
if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

excel_path = Path(r"\\fstvn01\Data\10_Production Engineering Department(製造技術部)\02.製造技術課\PE Dept\15. FORM（BIEU MAU）-形式\Form_VBA\Form_VBA_SS_BOM\tonghop_new12052026_ma1.xlsm")

wb = openpyxl.load_workbook(str(excel_path), data_only=True)
print("Sheets in workbook:", wb.sheetnames)

boloc_ws = wb["BolocBom"]
rows = list(boloc_ws.iter_rows(values_only=True))
print(f"BolocBom total rows: {len(rows)}")

# Print header and sample rows
header = rows[0]
print(f"Header: {header}")

boloc_data = []
for idx, r in enumerate(rows[1:], start=2):
    if any(c is not None for c in r):
        boloc_data.append({
            "excel_row": idx,
            "values": [str(c) if c is not None else "" for c in r[:10]]
        })

print(f"Non-empty rows count: {len(boloc_data)}")
for r in boloc_data[:20]:
    print(f"Row {r['excel_row']}: {r['values']}")

# Also save full data to json
out_json = Path("scripts/excel_bolocbom_dump.json")
with open(out_json, "w", encoding="utf-8") as f:
    json.dump({
        "header": [str(c) if c is not None else "" for c in header[:10]],
        "rows": boloc_data
    }, f, ensure_ascii=False, indent=2)

print(f"\nSaved full dump to {out_json}")
