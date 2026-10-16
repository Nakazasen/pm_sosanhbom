import openpyxl

file_path = "20261002/BolocBom_Spica.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb.active

updated_3v = 0
updated_branch = 0

for row in range(2, ws.max_row + 1):
    part_code = str(ws.cell(row=row, column=4).value or "").strip()
    notes = str(ws.cell(row=row, column=5).value or "").strip()
    
    if part_code == "3V2YJ01090":
        # Change action to Loại bỏ cả cụm (prune_node) and remove [branch: TVC1TZ0030]
        new_note = "Tạo tự động từ Cây BOM Mô phỏng (Loại bỏ cả cụm)"
        ws.cell(row=row, column=5, value=new_note)
        updated_3v += 1
        print(f"Row {row} (3V2YJ01090): Old='{notes}' -> New='{new_note}'")
    elif "[branch: TVC1TZ0030]" in notes:
        new_note = notes.replace(" [branch: TVC1TZ0030]", "").replace("[branch: TVC1TZ0030]", "").strip()
        ws.cell(row=row, column=5, value=new_note)
        updated_branch += 1
        print(f"Row {row} ({part_code}): Old='{notes}' -> New='{new_note}'")

wb.save(file_path)
print(f"Updated 3V2YJ01090 count: {updated_3v}")
print(f"Updated branch rules count: {updated_branch}")
print(f"Total updated: {updated_3v + updated_branch}")
