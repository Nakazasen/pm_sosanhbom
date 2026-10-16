import openpyxl

wb = openpyxl.load_workbook('20261002/BolocBom_Spica.xlsx')
ws = wb.active
print(f"Total rows in BolocBom_Spica.xlsx: {ws.max_row}")

tvc_rows = []
for r in range(2, ws.max_row + 1):
    vals = [cell.value for cell in ws[r]]
    model, item_name, match_mode, part_code, notes = vals[0], vals[1], vals[2], vals[3], vals[4]
    notes_str = str(notes or '')
    if 'TVC1TZ0030' in notes_str:
        tvc_rows.append((r, model, item_name, match_mode, part_code, notes_str))

print(f"Rows with TVC1TZ0030 ({len(tvc_rows)}):")
for r, model, item_name, match_mode, part_code, notes_str in tvc_rows:
    print(f"Row {r}: item={item_name!r}, part_code={part_code!r}, notes={notes_str!r}")
