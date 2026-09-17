import openpyxl

wb = openpyxl.load_workbook(r"D:\Sandbox\pm_sosanhbom\formnguoidung.xlsm", data_only=False)

with open(r"D:\Sandbox\pm_sosanhbom\.agents\teamwork_preview_explorer_survey_1\formnguoidung_formulas.txt", "w", encoding="utf-8") as out:
    for sname in wb.sheetnames:
        ws = wb[sname]
        out.write(f"\n=========================================\n")
        out.write(f"SHEET: {sname} (max_row={ws.max_row}, max_col={ws.max_column})\n")
        out.write(f"=========================================\n")
        for r in range(1, min(6, ws.max_row + 1)):
            row_items = []
            for c in range(1, min(20, ws.max_column + 1)):
                val = ws.cell(r, c).value
                if val is not None:
                    row_items.append(f"Col {c} ({ws.cell(r, c).coordinate}): {val}")
            if row_items:
                out.write(f"Row {r}:\n  " + "\n  ".join(row_items) + "\n")

wb.close()
print("formnguoidung_formulas.txt written")
